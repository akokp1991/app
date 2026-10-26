import asyncio
import calendar
from datetime import date, datetime, time, timedelta
import json
import logging
import os
import hmac
import re
import shutil
import sqlite3
import secrets
import statistics
import threading
import time as time_module
import traceback
import urllib.parse
import urllib.request
from zoneinfo import ZoneInfo

import arabic_reshaper
from bidi.algorithm import get_display
from flask import (
    Flask,
    flash,
    jsonify,
    redirect,
    render_template_string,
    request,
    send_file,
    session,
    url_for,
)
import jdatetime
from werkzeug.security import check_password_hash, generate_password_hash

import matplotlib
from matplotlib import font_manager
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from telegram import (
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    ReplyKeyboardMarkup,
    ReplyKeyboardRemove,
    Update,
)
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    ConversationHandler,
    MessageHandler,
    TypeHandler,
    ApplicationHandlerStop,
    filters,
)

# ============================================================
# تنظیمات پایه و لاگ‌برداری
# ============================================================
logging.basicConfig(format="%(asctime)s - %(name)s - %(levelname)s - %(message)s", level=logging.INFO)
logging.getLogger("httpx").setLevel(logging.WARNING)
logging.getLogger("httpcore").setLevel(logging.WARNING)
logging.getLogger("werkzeug").setLevel(logging.WARNING)
logger = logging.getLogger(__name__)
matplotlib.use("Agg")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_FILE = os.path.join(BASE_DIR, "bashmaq_reports.db")
TOKEN_FILE = os.path.join(BASE_DIR, "token.txt")
TOKEN = os.getenv("TELEGRAM_TOKEN", "").strip()
if not TOKEN and os.path.exists(TOKEN_FILE):
    with open(TOKEN_FILE, "r", encoding="utf-8") as f:
        TOKEN = f.read().strip()

ADMIN_ID = int(os.getenv("ADMIN_ID", "1439302328"))
OPEN_ACCESS = os.getenv("OPEN_ACCESS", "0").strip() == "1"   # 1 = هر کاربر تلگرام بدون تأیید مدیر دسترسی بگیرد
TRUST_PROXY = os.getenv("TRUST_PROXY", "0").strip() == "1"  # فقط پشت reverse proxy مطمئن فعال شود
OUT_DIR = os.path.join(BASE_DIR, "generated")
os.makedirs(OUT_DIR, exist_ok=True)


def _resolve_web_password():
    env_pw = os.getenv("WEB_PASSWORD", "").strip()
    if env_pw and env_pw != "123456":
        return env_pw, "env"
    pw_path = os.path.join(BASE_DIR, ".web_password")
    try:
        if os.path.exists(pw_path):
            with open(pw_path, "r", encoding="utf-8") as fh:
                saved = fh.read().strip()
            if len(saved) >= 8:
                return saved, "file"
        generated = secrets.token_urlsafe(12)
        fd = os.open(pw_path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            fh.write(generated)
        logger.warning("WEB_PASSWORD تنظیم نشده بود؛ رمز تصادفی در فایل امن %s ذخیره شد. مقدار رمز در لاگ ثبت نمی‌شود.", pw_path)
        return generated, "generated"
    except Exception:
        return secrets.token_urlsafe(12), "ephemeral"


WEB_PASS, WEB_PASS_SOURCE = _resolve_web_password()
PORT = int(os.getenv("PORT", os.getenv("SERVER_PORT", "5000")))
FIXED_FOOTER = "آکو کهنه پوشی نماینده شرکت ایران مقصد - مرز باشماق"

try:
    TEHRAN = ZoneInfo("Asia/Tehran")
except Exception:
    from datetime import timezone
    TEHRAN = timezone(timedelta(hours=3, minutes=30))

GLOBAL_BOT_APP = None
GLOBAL_LOOP = None

(DATE, COUNT, NAME, ISSUED, START, END, CANCELLED, CANCELLED_SERIAL, TRANSIT, EXPORT, CONFIRM_SAVE) = range(11)
SEARCH_DATE, RANGE_START, RANGE_END, MONTH, EDIT_DATE, EDIT_CHOICE, DELETE_DATE = 11, 12, 13, 14, 15, 16, 17
SEND_DATE, SEND_PLATFORM, SEND_USER, SEND_WA_PHONE, BRANCH_SEARCH_STATE, BROADCAST_STATE, SERIAL_SEARCH_STATE, RESTORE_DB_STATE, FAST_REPORT_STATE = 18, 19, 20, 21, 22, 23, 24, 25, 26


# ============================================================
# مدیریت هوشمند فونت استاندارد فارسی
# ============================================================
def ensure_standard_persian_font():
    path = os.path.join(BASE_DIR, "Vazirmatn-Regular.ttf")
    if os.path.exists(path) and os.path.getsize(path) > 50000:
        return path
    urls = [
        "https://cdn.jsdelivr.net/gh/rastikerdar/vazirmatn@v33.003/fonts/ttf/Vazirmatn-Regular.ttf",
        "https://raw.githubusercontent.com/rastikerdar/vazirmatn/master/fonts/ttf/Vazirmatn-Regular.ttf",
    ]
    for url in urls:
        try:
            with urllib.request.urlopen(url, timeout=15) as resp, open(path, "wb") as out:
                shutil.copyfileobj(resp, out)
            if os.path.exists(path) and os.path.getsize(path) > 50000:
                return path
        except Exception:
            pass
    for p in [os.path.join(BASE_DIR, "BNAZANIN.ttf"), r"C:\Windows\Fonts\tahoma.ttf", r"C:\Windows\Fonts\arial.ttf"]:
        if os.path.exists(p):
            return p
    return None


FONT_PATH = ensure_standard_persian_font()
CUSTOM_RESHAPER = arabic_reshaper.ArabicReshaper(configuration={"delete_harakat": True, "support_ligatures": False})

MPL_FONT_PROP, PDF_FONT_NAME = None, None
if FONT_PATH:
    try:
        font_manager.fontManager.addfont(FONT_PATH)
        MPL_FONT_PROP = font_manager.FontProperties(fname=FONT_PATH)
        plt.rcParams["font.family"] = MPL_FONT_PROP.get_name()
    except Exception:
        pass
    try:
        pdfmetrics.registerFont(TTFont("PersianFont", FONT_PATH))
        PDF_FONT_NAME = "PersianFont"
    except Exception:
        pass


def clean_text_for_render(s):
    s = str(s)
    for old_c, new_c in {"•": "", "▪": "", "▫️": "", "—": "-", "–": "-", "|": " / ", "☀️": "", "📅": "", "🏢": "", "🚛": "", "📦": "", "✅": "", "❌": ""}.items():
        s = s.replace(old_c, new_c)
    return s.strip()


def ptext(s):
    return get_display(CUSTOM_RESHAPER.reshape(clean_text_for_render(s)))


# ============================================================
# توابع کمکی، تاریخ و دسترسی‌ها
# ============================================================
def tr_num(s):
    for p, a, e in zip("۰۱۲۳۴۵۶۷۸۹", "٠١٢٣٤٥٦٧٨٩", "0123456789"):
        s = str(s).replace(p, e).replace(a, e)
    return s.strip()


def normalize_phone(ps):
    p = tr_num(ps).replace("+", "").replace(" ", "").replace("-", "")
    if p.startswith("00"):
        p = p[2:]
    elif p.startswith("09") and len(p) == 11:
        p = "98" + p[1:]
    elif p.startswith("9") and len(p) == 10:
        p = "98" + p
    return p if p.isdigit() and len(p) >= 10 else None


def parse_smart_date(s):
    pts = [int(x) for x in tr_num(s).replace("-", "/").replace(".", "/").split("/") if x.isdigit()]
    if len(pts) != 3:
        raise ValueError("Invalid date")
    d, m, y = (pts[0], pts[1], pts[2]) if pts[0] <= 1000 else (pts[2], pts[1], pts[0])
    if 1300 <= y <= 1500:
        return jdatetime.date(y, m, d).togregorian()
    elif 1900 <= y <= 2200:
        return date(y, m, d)
    raise ValueError("Year out of range")


def norm_date(s):
    return parse_smart_date(s).strftime("%d/%m/%Y")


def shamsi(d):
    return jdatetime.date.fromgregorian(date=d).strftime("%d/%m/%Y")


def iran_today():
    return datetime.now(TEHRAN).date()


def safe_int(s):
    try:
        return int(tr_num(s)) if int(tr_num(s)) >= 0 else None
    except Exception:
        return None


def safe_float(s):
    try:
        v = tr_num(s).replace(",", "")
        return float(v) if v else None
    except Exception:
        return None


def total_issued(data):
    return sum(int(b.get("issued", 0)) for b in data.get("branches", []))


def total_cancelled(data):
    return sum(int(b.get("cancelled", 0)) for b in data.get("branches", []))


def get_branch_ranges_text(b):
    ranges = b.get("ranges", [{"start": b.get("start", ""), "end": b.get("end", "")}])
    if not ranges:
        return "0"
    return " و ".join([f"{r['start']} الی {r['end']}" for r in ranges if r.get("start") and r.get("end")])


# ============================================================
# دیتابیس و توابع هسته (شامل جدول حسابداری ماهانه)
# ============================================================
def db():
    conn = sqlite3.connect(DB_FILE, timeout=20.0)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA busy_timeout=20000")
    conn.execute("PRAGMA synchronous=FULL")
    return conn


def init_db():
    conn = db()
    conn.execute("CREATE TABLE IF NOT EXISTS users (user_id INTEGER PRIMARY KEY, username TEXT, first_name TEXT, allowed INTEGER DEFAULT 1, is_admin INTEGER DEFAULT 0, created_at TEXT NOT NULL DEFAULT '')")
    conn.execute("CREATE TABLE IF NOT EXISTS admins (user_id INTEGER PRIMARY KEY)")
    conn.execute("CREATE TABLE IF NOT EXISTS audit_logs (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, action TEXT NOT NULL, timestamp TEXT NOT NULL)")
    conn.execute("CREATE TABLE IF NOT EXISTS reports (date TEXT PRIMARY KEY, data TEXT NOT NULL, created_at TEXT NOT NULL)")
    conn.execute("CREATE TABLE IF NOT EXISTS financial_sheets (month_key TEXT PRIMARY KEY, data TEXT NOT NULL, updated_at TEXT NOT NULL)")
    conn.execute("CREATE TABLE IF NOT EXISTS branch_employees (user_id INTEGER PRIMARY KEY, employee_name TEXT NOT NULL, branch_name TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL)")
    conn.execute("CREATE TABLE IF NOT EXISTS pending_branch_reports (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, report_date TEXT NOT NULL, shamsi_date TEXT NOT NULL, branch_name TEXT NOT NULL, issued INTEGER NOT NULL, serial_start TEXT NOT NULL, serial_end TEXT NOT NULL, cancelled INTEGER NOT NULL DEFAULT 0, cancelled_serial TEXT NOT NULL DEFAULT '0', status TEXT NOT NULL DEFAULT 'pending', created_at TEXT NOT NULL, reviewed_at TEXT)")
    conn.execute("CREATE TABLE IF NOT EXISTS closed_workdays (report_date TEXT PRIMARY KEY, shamsi_date TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'closed', transit INTEGER NOT NULL DEFAULT 0, export INTEGER NOT NULL DEFAULT 0, closed_by INTEGER NOT NULL, closed_at TEXT NOT NULL, close_note TEXT NOT NULL DEFAULT '')")
    conn.execute("CREATE TABLE IF NOT EXISTS report_change_log (id INTEGER PRIMARY KEY AUTOINCREMENT, report_date TEXT NOT NULL, action TEXT NOT NULL, reason TEXT NOT NULL, user_id INTEGER NOT NULL, created_at TEXT NOT NULL)")

    conn.execute("CREATE TABLE IF NOT EXISTS access_requests (user_id INTEGER PRIMARY KEY, status TEXT NOT NULL DEFAULT 'pending', requested_at TEXT NOT NULL)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_pending_status_date ON pending_branch_reports(status, report_date)")
    _emp_cols = {r[1] for r in conn.execute("PRAGMA table_info(branch_employees)")}
    if "can_submit" not in _emp_cols:
        conn.execute("ALTER TABLE branch_employees ADD COLUMN can_submit INTEGER NOT NULL DEFAULT 1")
    if "note" not in _emp_cols:
        conn.execute("ALTER TABLE branch_employees ADD COLUMN note TEXT NOT NULL DEFAULT ''")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_audit_user ON audit_logs(user_id)")
    conn.execute("INSERT OR IGNORE INTO admins(user_id) VALUES (?)", (ADMIN_ID,))
    conn.execute("UPDATE users SET is_admin=1, allowed=1 WHERE user_id=?", (ADMIN_ID,))
    conn.commit()
    conn.close()
    invalidate_reports_cache()


def log_action(uid, act):
    try:
        conn = db()
        conn.execute("INSERT INTO audit_logs(user_id, action, timestamp) VALUES (?, ?, ?)", (uid, act, datetime.now(TEHRAN).isoformat(timespec="seconds")))
        conn.commit()
        conn.close()
    except Exception:
        pass


def get_recent_logs(l=35):
    conn = db()
    rows = conn.execute("SELECT user_id, action, timestamp FROM audit_logs ORDER BY id DESC LIMIT ?", (l,)).fetchall()
    conn.close()
    return rows


def save_user(u):
    if not u:
        return
    conn = db()
    conn.execute(
        "INSERT INTO users(user_id, username, first_name, allowed, is_admin, created_at) VALUES (?, ?, ?, ?, ?, ?) ON CONFLICT(user_id) DO UPDATE SET username=excluded.username, first_name=excluded.first_name",
        (u.id, u.username or "", u.first_name or "", 1 if (OPEN_ACCESS or is_admin_id(u.id)) else 0, 1 if is_admin_id(u.id) else 0, datetime.now(TEHRAN).isoformat(timespec="seconds")),
    )
    conn.commit()
    conn.close()


def is_admin_id(uid):
    conn = db()
    r = conn.execute("SELECT 1 FROM admins WHERE user_id=?", (uid,)).fetchone()
    conn.close()
    return bool(r) or uid == ADMIN_ID


def is_allowed(uid):
    if is_admin_id(uid):
        return True
    conn = db()
    r = conn.execute("SELECT allowed FROM users WHERE user_id=?", (uid,)).fetchone()
    conn.close()
    if r is not None and r[0]:
        return True
    if get_employee(uid):
        return True
    return OPEN_ACCESS if r is None else False


def is_admin(update):
    return bool(update.effective_user and is_admin_id(update.effective_user.id))


def can_use(update):
    return bool(update.effective_user and is_allowed(update.effective_user.id))


def add_admin(uid):
    conn = db()
    conn.execute("INSERT OR IGNORE INTO admins(user_id) VALUES (?)", (uid,))
    conn.execute("INSERT OR IGNORE INTO users(user_id, allowed, is_admin) VALUES (?,1,1)", (uid,))
    conn.execute("UPDATE users SET is_admin=1, allowed=1 WHERE user_id=?", (uid,))
    conn.commit()
    conn.close()


def remove_admin(uid):
    if uid == ADMIN_ID:
        return False
    conn = db()
    c = conn.execute("DELETE FROM admins WHERE user_id=?", (uid,))
    conn.execute("UPDATE users SET is_admin=0 WHERE user_id=?", (uid,))
    conn.commit()
    conn.close()
    return c.rowcount > 0


def set_allowed(uid, a):
    conn = db()
    conn.execute("INSERT OR IGNORE INTO users(user_id, allowed) VALUES (?,?)", (uid, 1 if a else 0))
    conn.execute("UPDATE users SET allowed=? WHERE user_id=?", (1 if a else 0, uid))
    conn.commit()
    conn.close()


def users_rows():
    conn = db()
    rows = conn.execute("SELECT user_id, username, first_name, allowed, is_admin FROM users ORDER BY user_id").fetchall()
    conn.close()
    return rows


def save_report(data):
    clean_b = []
    for b in data["branches"]:
        ranges = b.get("ranges", [])
        if not ranges and b.get("start") and b.get("end"):
            ranges = [{"start": b["start"], "end": b["end"]}]
        clean_b.append({
            "name": b["name"],
            "issued": b["issued"],
            "ranges": ranges,
            "start": ranges[0]["start"] if ranges else "0",
            "end": ranges[-1]["end"] if ranges else "0",
            "range_text": " و ".join([f"{r['start']} الی {r['end']}" for r in ranges]) if ranges else "0",
            "cancelled": b.get("cancelled", 0),
            "cancelled_serial": b.get("cancelled_serial", "0"),
        })
    cd = {
        "date": data["date"],
        "shamsi_date": data["shamsi_date"],
        "branches": clean_b,
        "transit": int(data.get("transit", 0)),
        "export": int(data.get("export", 0)),
        "created_by": data.get("created_by", ADMIN_ID),
    }
    conn = db()
    conn.execute(
        "INSERT INTO reports(date, data, created_at) VALUES (?, ?, ?) "
        "ON CONFLICT(date) DO UPDATE SET data=excluded.data, created_at=excluded.created_at",
        (cd["date"], json.dumps(cd, ensure_ascii=False), datetime.now(TEHRAN).isoformat(timespec="seconds")),
    )
    conn.commit()
    conn.close()
    invalidate_reports_cache()
    return cd


def load_report(dt):
    conn = db()
    r = conn.execute("SELECT data FROM reports WHERE date=?", (dt,)).fetchone()
    conn.close()
    return json.loads(r[0]) if r else None


def delete_report(dt):
    conn = db()
    c = conn.execute("DELETE FROM reports WHERE date=?", (dt,))
    conn.commit()
    conn.close()
    invalidate_reports_cache()
    return c.rowcount > 0


def all_reports():
    """لیست مرتب‌شده‌ی گزارش‌ها؛ با کش (باطل‌سازی هنگام ذخیره/حذف/بازیابی). رکورد خراب نادیده گرفته می‌شود."""
    with _REPORTS_LOCK:
        raw = _REPORTS_CACHE["raw"]
        if raw is None:
            conn = db()
            try:
                rows = conn.execute("SELECT data FROM reports").fetchall()
            finally:
                conn.close()
            tmp = []
            for (txt,) in rows:
                try:
                    it = json.loads(txt)
                    tmp.append((parse_smart_date(it.get("date", "01/01/2000")), txt))
                except Exception:
                    continue
            tmp.sort(key=lambda x: x[0])
            raw = [t for _, t in tmp]
            _REPORTS_CACHE["raw"] = raw
    return [json.loads(t) for t in raw]


def list_dates():
    items = all_reports()
    items.reverse()
    return [f"{x['date']} (شمسی: {x.get('shamsi_date', '-')})" for x in items]


def get_known_branches():
    brs = ["تیرپارک", "پارسیان"]
    for r in all_reports():
        for b in r.get("branches", []):
            nm = b.get("name", "").strip()
            if nm and nm not in brs and nm not in {"aa", "ff", "b", "B", "ت", "بدون نام"} and len(nm) >= 2:
                brs.append(nm)
    return brs[:6]


def get_last_end_serial(branch_name, before_date_str=None):
    rep = all_reports()
    td = parse_smart_date(before_date_str) if before_date_str else None
    le, ld = None, None
    for r in rep:
        try:
            if td and parse_smart_date(r["date"]) >= td:
                continue
        except Exception:
            continue
        for b in r.get("branches", []):
            if b.get("name", "").strip() == branch_name.strip():
                ranges = b.get("ranges", [])
                es = tr_num(ranges[-1]["end"]) if ranges else tr_num(b.get("end", "0"))
                if es.isdigit() and int(es) > 0:
                    le, ld = int(es), r.get("shamsi_date", r["date"])
    return le, ld


def find_missing_dates(days_back=30):
    td = iran_today()
    ed = {r.get("date") for r in all_reports()}
    m = []
    for i in range(days_back, -1, -1):
        d = td - timedelta(days=i)
        ds = d.strftime("%d/%m/%Y")
        if ds not in ed:
            m.append({"gregorian": ds, "shamsi": shamsi(d)})
    return m


def save_financial_sheet(month_key, sheet_data):
    conn = db()
    conn.execute(
        "INSERT OR REPLACE INTO financial_sheets(month_key, data, updated_at) VALUES (?, ?, ?)",
        (month_key, json.dumps(sheet_data, ensure_ascii=False), datetime.now(TEHRAN).isoformat(timespec="seconds")),
    )
    conn.commit()
    conn.close()


def load_financial_sheet(month_key):
    conn = db()
    row = conn.execute("SELECT data FROM financial_sheets WHERE month_key=?", (month_key,)).fetchone()
    conn.close()
    return json.loads(row[0]) if row else None


def build_month_accounting_rows(year_int, month_int, sync_from_reports=False):
    month_key = f"{year_int:04d}-{month_int:02d}"
    saved = load_financial_sheet(month_key) or {}
    saved_rows = {r.get("g_date"): r for r in saved.get("rows", [])}
    dinar_fee = int(saved.get("dinar_rate", 15000))

    _, num_days = calendar.monthrange(year_int, month_int)
    rows = []
    for day in range(1, num_days + 1):
        dt = date(year_int, month_int, day)
        g_iso = dt.strftime("%Y-%m-%d")
        g_slash = dt.strftime("%d/%m/%Y")
        jd = jdatetime.date.fromgregorian(date=dt)
        s_short = f"{jd.year - 1000:03d}-{jd.month:02d}-{jd.day:02d}"

        existing_row = saved_rows.get(g_iso, {})
        amar_val = existing_row.get("amar", "")

        if sync_from_reports or amar_val == "":
            rep = load_report(g_slash)
            if rep:
                amar_val = total_issued(rep)

        rows.append({
            "s_date": s_short,
            "g_date": g_iso,
            "amar": amar_val if amar_val != "" else "",
            "iraq_hawala": existing_row.get("iraq_hawala", ""),
            "torshizi": existing_row.get("torshizi", ""),
            "naser": existing_row.get("naser", ""),
            "ex_rate": existing_row.get("ex_rate", ""),
            "iran_maqsad": existing_row.get("iran_maqsad", ""),
            "sabokbar": existing_row.get("sabokbar", ""),
            "motafareqeh": existing_row.get("motafareqeh", ""),
            "desc": existing_row.get("desc", ""),
        })
    return {"month_key": month_key, "dinar_rate": dinar_fee, "rows": rows}


def create_accounting_pdf_file(year_int, month_int):
    """گزارش مالی ماهانه با طراحی رسمی و حرفه‌ای، A4 افقی و جدول رنگی."""
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT, TA_LEFT
    from reportlab.lib.units import mm
    from xml.sax.saxutils import escape

    sheet = build_month_accounting_rows(year_int, month_int)
    rows = sheet.get('rows', [])
    rate = float(sheet.get('dinar_rate', 15000) or 15000)

    def num(v):
        try:
            return float(str(v or 0).replace(',', '').replace('٬', '').replace('٫', '.'))
        except Exception:
            return 0.0

    total_count = sum(num(r.get('amar')) for r in rows)
    total_dinar = sum(num(r.get('amar')) * rate for r in rows)
    iraq = sum(num(r.get('iraq_hawala')) for r in rows)
    torshizi = sum(num(r.get('torshizi')) for r in rows)
    naser = sum(num(r.get('naser')) for r in rows)
    deductions = iraq + torshizi + naser
    net = total_dinar - deductions
    usd = sum((num(r.get('amar')) * rate / num(r.get('ex_rate')) if num(r.get('ex_rate')) else 0)
              + num(r.get('iran_maqsad')) + num(r.get('sabokbar')) + num(r.get('motafareqeh')) for r in rows)

    month_name = jdatetime.date.fromgregorian(date=date(year_int, month_int, 1)).strftime('%B %Y')
    path = os.path.join(OUT_DIR, f"hesabdari_{year_int}_{month_int:02d}.pdf")
    navy = colors.HexColor('#142D4E')
    blue = colors.HexColor('#1D6FA5')
    pale = colors.HexColor('#F2F6FB')
    line = colors.HexColor('#B8C6D6')
    teal = colors.HexColor('#DDF4F0')
    red = colors.HexColor('#F3DADA')
    gold = colors.HexColor('#F6E7BE')
    green = colors.HexColor('#E1F3E5')

    doc = SimpleDocTemplate(path, pagesize=landscape(A4), rightMargin=5*mm, leftMargin=5*mm,
                            topMargin=5*mm, bottomMargin=5*mm, title=f'صورت‌حساب مالی ماهانه {month_name}',
                            author='سامانه مدیریت مالی مرز باشماق', subject='گزارش مالی ماهانه')
    styles = getSampleStyleSheet()
    font = PDF_FONT_NAME or 'Helvetica'
    title_style = ParagraphStyle('AccTitle', parent=styles['Title'], fontName=font,
        fontSize=15, leading=18, textColor=colors.white, alignment=TA_RIGHT, spaceAfter=0)
    subtitle_style = ParagraphStyle('AccSubtitle', parent=styles['Normal'], fontName=font,
        fontSize=8, leading=10, textColor=colors.HexColor('#E3EDF7'), alignment=TA_RIGHT)
    cell = ParagraphStyle('AccountingCell', parent=styles['Normal'], fontName=font,
        fontSize=6.6, leading=7.4, alignment=TA_CENTER, textColor=colors.HexColor('#243447'),
        spaceAfter=0, spaceBefore=0)
    head_cell = ParagraphStyle('AccountingHead', parent=cell, fontSize=6.5, leading=7.1,
        textColor=colors.white)
    small = ParagraphStyle('AccountingSmall', parent=cell, fontSize=6.2, leading=7,
        textColor=colors.HexColor('#5B6777'))
    def P(v, style=cell):
        txt = str(v if v is not None else '')
        return Paragraph(ptext(txt) if PDF_FONT_NAME else escape(txt), style)

    # سربرگ رسمی با عنوان، بازه گزارش و نرخ محاسبات
    header = Table([
        [P('صورت‌حساب مالی ماهانه | مرز باشماق', title_style)],
        [P(f'گزارش دوره: {month_name}  •  نرخ محاسبات: {rate:,.0f} دینار  •  تاریخ تهیه: {date.today().isoformat()}', subtitle_style)]
    ], colWidths=[277*mm], rowHeights=[10*mm, 6*mm])
    header.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,-1),navy), ('BOX',(0,0),(-1,-1),0.8,navy),
        ('LEFTPADDING',(0,0),(-1,-1),5*mm), ('RIGHTPADDING',(0,0),(-1,-1),5*mm),
        ('TOPPADDING',(0,0),(-1,-1),1.1*mm), ('BOTTOMPADDING',(0,0),(-1,-1),1.1*mm),
        ('VALIGN',(0,0),(-1,-1),'MIDDLE')]))

    headers = ['مبلغ دلار', 'مبلغ تحویلی به ناصرمقصد', 'مبلغ تحویلی به آقای ترتیفی',
               'حواله عراق', 'مبلغ دینار', 'آمار', 'تاریخ میلادی', 'تاریخ شمسی']
    data = [[P(h, head_cell) for h in headers]]
    for r in rows:
        count = num(r.get('amar'))
        dinar = count * rate
        ex = num(r.get('ex_rate'))
        dollars = (dinar / ex if ex else 0) + num(r.get('iran_maqsad')) + num(r.get('sabokbar')) + num(r.get('motafareqeh'))
        data.append([P(f'${dollars:,.0f}'), P(f'{num(r.get("naser")):,.0f}' if num(r.get('naser')) else '—'),
                     P(f'{num(r.get("torshizi")):,.0f}' if num(r.get('torshizi')) else '—'),
                     P(f'{num(r.get("iraq_hawala")):,.0f}' if num(r.get('iraq_hawala')) else '—'),
                     P(f'{dinar:,.0f}'), P(f'{count:,.0f}' if r.get('amar') not in ('', None) else '—'),
                     P(r.get('g_date', '')), P(r.get('s_date', ''))])
    data.append([P(f'${usd:,.0f}'), P(f'{naser:,.0f}'), P(f'{torshizi:,.0f}'), P(f'{iraq:,.0f}'),
                 P(f'{total_dinar:,.0f}'), P(f'{total_count:,.0f}'), P('جمع ستون‌ها'), P('')])
    widths = [25*mm, 39*mm, 39*mm, 38*mm, 31*mm, 31*mm, 27*mm, 27*mm]
    tbl = Table(data, colWidths=widths, repeatRows=1, hAlign='CENTER')
    table_commands = [
        ('BACKGROUND',(0,0),(-1,0),blue), ('TEXTCOLOR',(0,0),(-1,0),colors.white),
        ('GRID',(0,0),(-1,-1),0.3,line), ('BOX',(0,0),(-1,-1),0.8,navy),
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'), ('ALIGN',(0,0),(-1,-1),'CENTER'),
        ('FONTNAME',(0,0),(-1,-1),font), ('TOPPADDING',(0,0),(-1,-1),1.45),
        ('BOTTOMPADDING',(0,0),(-1,-1),1.45),
        ('BACKGROUND',(4,1),(4,-2),colors.HexColor('#F7FAFD')),
        ('BACKGROUND',(0,1),(0,-2),colors.HexColor('#FBFCFE')),
        ('BACKGROUND',(0,-1),(-1,-1),pale), ('LINEABOVE',(0,-1),(-1,-1),1.1,navy),
        ('FONTNAME',(0,-1),(-1,-1),font), ('TEXTCOLOR',(0,-1),(-1,-1),navy),
        ('BACKGROUND',(4,-1),(4,-1),red), ('BACKGROUND',(5,-1),(5,-1),colors.HexColor('#BDE8F8')),
        ('BACKGROUND',(0,-1),(0,-1),red), ('FONTNAME',(4,-1),(5,-1),font),
    ]
    for ridx in range(1, len(data)-1):
        if ridx % 2 == 0:
            table_commands.append(('BACKGROUND',(0,ridx),(-1,ridx),colors.HexColor('#F8FAFC')))
    tbl.setStyle(TableStyle(table_commands))

    # خلاصه مالی به شکل کارت‌های رنگی و ردیف تفصیلی کسورات
    card_style = ParagraphStyle('AccCardLabel', parent=cell, fontSize=7, leading=8,
                                textColor=colors.HexColor('#506174'))
    value_style = ParagraphStyle('AccCardValue', parent=cell, fontSize=10, leading=12,
                                 textColor=navy)
    net_style = ParagraphStyle('AccNetValue', parent=value_style, fontSize=12,
                               textColor=colors.HexColor('#1769C2'))
    cards = Table([
        [P('جمع کل دینار', card_style), P('جمع کسورات', card_style), P('مبلغ نهایی', card_style), P('جمع دلار', card_style)],
        [P(f'{total_dinar:,.0f}', value_style), P(f'{deductions:,.0f}', value_style),
         P(f'{net:,.0f}', net_style), P(f'${usd:,.0f}', value_style)],
        [P(f'تعداد کل: {total_count:,.0f}', small), P(f'حواله عراق: {iraq:,.0f}', small),
         P(f'ناصرمقصد: {naser:,.0f}', small), P(f'آقای ترتیفی: {torshizi:,.0f}', small)]
    ], colWidths=[69.25*mm]*4, rowHeights=[6*mm, 9*mm, 6*mm])
    cards.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,-1),pale), ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E5EDF6')),
        ('BACKGROUND',(1,1),(1,1),gold), ('BACKGROUND',(2,1),(2,1),teal),
        ('BOX',(0,0),(-1,-1),0.65,line), ('INNERGRID',(0,0),(-1,-1),0.35,line),
        ('ALIGN',(0,0),(-1,-1),'CENTER'), ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('FONTNAME',(0,0),(-1,-1),font), ('LEFTPADDING',(0,0),(-1,-1),2),
        ('RIGHTPADDING',(0,0),(-1,-1),2), ('TOPPADDING',(0,0),(-1,-1),1),
        ('BOTTOMPADDING',(0,0),(-1,-1),1)]))

    signature = Table([
        [P('محل مهر و امضای مسئول', card_style), P('سید آکو کهنه‌پوشی', cell)],
        ['', P('امضا و تأیید نماینده مرز باشماق', small)],
        ['', '']
    ], colWidths=[65*mm, 65*mm], rowHeights=[5*mm, 5*mm, 5*mm])
    signature.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(0,-1),colors.HexColor('#F8FAFC')),
        ('BOX',(0,0),(-1,-1),0.55,line), ('LINEBELOW',(1,1),(1,1),0.7,colors.HexColor('#475569')),
        ('ALIGN',(0,0),(-1,-1),'CENTER'), ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('FONTNAME',(0,0),(-1,-1),font), ('LEFTPADDING',(0,0),(-1,-1),2),
        ('RIGHTPADDING',(0,0),(-1,-1),2), ('TOPPADDING',(0,0),(-1,-1),1),
        ('BOTTOMPADDING',(0,0),(-1,-1),1)]))

    # مهم: کارت‌ها عرض کامل صفحه (277mm) دارند؛ قرار دادن آن‌ها در یک ستون
    # نصف‌عرض کنار امضا باعث سرریز، هم‌پوشانی و ناخوانا شدن PDF می‌شد.
    # بنابراین خلاصه مالی و بخش امضا در دو ردیف مستقل قرار می‌گیرند.
    signature.hAlign = 'CENTER'

    def page_decor(c, d):
        w, h = landscape(A4)
        c.saveState()
        c.setStrokeColor(colors.HexColor('#D5DEE8'))
        c.setLineWidth(0.6)
        c.roundRect(3*mm, 3*mm, w-6*mm, h-6*mm, 2*mm, stroke=1, fill=0)
        c.setFont(font, 6)
        footer = f'سامانه مدیریت مالی مرز باشماق  •  صورت‌حساب {year_int:04d}/{month_int:02d}  •  صفحه {d.page}'
        c.setFillColor(colors.HexColor('#64748B'))
        c.drawCentredString(w/2, 4.5*mm, ptext(footer) if PDF_FONT_NAME else footer)
        c.restoreState()

    story = [header, Spacer(1, 2*mm), tbl, Spacer(1, 4*mm), cards, Spacer(1, 3*mm), signature]
    doc.build(story, onFirstPage=page_decor, onLaterPages=page_decor)
    return path

def create_accounting_excel_file(year_int, month_int):
    sheet_info = build_month_accounting_rows(year_int, month_int)
    dinar_rate = int(sheet_info.get("dinar_rate", 15000))
    rows_data = sheet_info["rows"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{month_int:02d}-{year_int}"
    ws.sheet_view.rightToLeft = True

    font_header = Font(name="Tahoma", size=10, bold=True, color="000000")
    font_data = Font(name="Tahoma", size=10, bold=False, color="000000")
    font_bold = Font(name="Tahoma", size=10, bold=True, color="000000")

    fill_header = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_grey = PatternFill(start_color="C5B3B3", end_color="C5B3B3", fill_type="solid")
    fill_pink_dark = PatternFill(start_color="E69999", end_color="E69999", fill_type="solid")
    fill_pink_light = PatternFill(start_color="F2CACA", end_color="F2CACA", fill_type="solid")
    fill_blue = PatternFill(start_color="5BC0EB", end_color="5BC0EB", fill_type="solid")
    fill_green = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")

    thin_side = Side(style="thin", color="595959")
    border_row = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
    center_align = Alignment(horizontal="center", vertical="center")

    headers = [
        "تاریخ شمسی",
        "تاریخ میلادی",
        "آمار",
        "مبلغ دینار",
        "حواله عراق",
        "مبلغ تحویلی به آقای ترتیفی",
        "مبلغ تحویلی به ناصردینار",
        "مبلغ دلار",
        "نرخ تبدیل",
        "ایران مقصد",
        "سبکبار",
        "متفرقه",
        "توضیحات",
    ]
    ws.append(headers)
    for col_idx in range(1, 14):
        c = ws.cell(row=1, column=col_idx)
        c.font = font_header
        c.alignment = center_align
        c.border = border_row
        if col_idx == 9:
            c.fill = fill_grey
        elif col_idx == 10:
            c.fill = fill_pink_dark
        elif col_idx in (11, 12):
            c.fill = fill_pink_light
        else:
            c.fill = fill_header

    for idx, r in enumerate(rows_data, start=2):
        amar_v = safe_int(r.get("amar", ""))
        hawala_v = safe_int(r.get("iraq_hawala", ""))
        tor_v = safe_int(r.get("torshizi", ""))
        nas_v = safe_int(r.get("naser", ""))
        ex_v = safe_float(r.get("ex_rate", ""))
        im_v = safe_float(r.get("iran_maqsad", ""))
        sb_v = safe_float(r.get("sabokbar", ""))
        mt_v = safe_float(r.get("motafareqeh", ""))

        ws.cell(row=idx, column=1, value=r["s_date"])
        ws.cell(row=idx, column=2, value=r["g_date"])
        ws.cell(row=idx, column=3, value=amar_v if amar_v is not None else "")
        ws.cell(row=idx, column=4, value=f"=IF(ISNUMBER(C{idx}), C{idx}*{dinar_rate}, 0)")
        ws.cell(row=idx, column=5, value=hawala_v if hawala_v is not None else "")
        ws.cell(row=idx, column=6, value=tor_v if tor_v is not None else "")
        ws.cell(row=idx, column=7, value=nas_v if nas_v is not None else "")
        ws.cell(row=idx, column=8, value=f"=IF(I{idx}>0, D{idx}/I{idx}, 0)+SUM(J{idx}:L{idx})")
        ws.cell(row=idx, column=9, value=ex_v if ex_v is not None else "")
        ws.cell(row=idx, column=10, value=im_v if im_v is not None else "")
        ws.cell(row=idx, column=11, value=sb_v if sb_v is not None else "")
        ws.cell(row=idx, column=12, value=mt_v if mt_v is not None else "")
        ws.cell(row=idx, column=13, value=r.get("desc", ""))

        for col_idx in range(1, 14):
            cell = ws.cell(row=idx, column=col_idx)
            cell.font = font_data
            cell.alignment = center_align
            cell.border = border_row
            if col_idx == 4:
                cell.number_format = "#,##0"
            elif col_idx == 8:
                cell.number_format = "$#,##0"
                cell.font = font_bold
            elif col_idx == 9:
                cell.fill = fill_grey
            elif col_idx == 10:
                cell.fill = fill_pink_dark
            elif col_idx in (11, 12):
                cell.fill = fill_pink_light

    last_r = len(rows_data) + 1
    sum_r = last_r + 1

    ws.cell(row=sum_r, column=3, value=f"=SUM(C2:C{last_r})").fill = fill_blue
    ws.cell(row=sum_r, column=4, value=f"=SUM(D2:D{last_r})").fill = fill_pink_dark
    ws.cell(row=sum_r, column=5, value=f"=SUM(E2:E{last_r})")
    ws.cell(row=sum_r, column=6, value=f"=SUM(F2:F{last_r})")
    ws.cell(row=sum_r, column=7, value=f"=SUM(G2:G{last_r})")
    ws.cell(row=sum_r, column=8, value=f"=SUM(H2:H{last_r})")
    ws.cell(row=sum_r, column=10, value=f"=SUM(J2:J{last_r})")
    ws.cell(row=sum_r, column=11, value=f"=SUM(K2:K{last_r})")
    ws.cell(row=sum_r, column=12, value=f"=SUM(L2:L{last_r})")

    for col_idx in range(1, 14):
        c = ws.cell(row=sum_r, column=col_idx)
        c.font = font_bold
        c.alignment = center_align
        c.border = border_row
        if col_idx in (4, 5, 6, 7):
            c.number_format = "#,##0"

    b1 = sum_r + 1
    ws.cell(row=b1, column=6, value="جمع کل").font = font_bold
    ws.cell(row=b1, column=8, value=f"=D{sum_r}").fill = fill_pink_dark
    ws.cell(row=b1, column=8).number_format = "#,##0"
    ws.cell(row=b1, column=10, value="جمع آمار").fill = fill_blue
    ws.cell(row=b1, column=11, value=f"=C{sum_r}").fill = fill_blue

    b2 = sum_r + 2
    ws.cell(row=b2, column=6, value="جمع کسورات").font = font_bold
    ws.cell(row=b2, column=8, value=f"=SUM(E{sum_r}:G{sum_r})").number_format = "#,##0"
    ws.cell(row=b2, column=10, value="سبکبار").fill = fill_pink_light
    ws.cell(row=b2, column=11, value=f"=K{sum_r}").fill = fill_pink_light
    ws.cell(row=b2, column=11).number_format = "$#,##0"

    b3 = sum_r + 3
    ws.cell(row=b3, column=2, value="سید آکو کهنه پوشی").font = font_bold
    ws.cell(row=b3, column=6, value="مبلغ نهایی").font = font_bold
    ws.cell(row=b3, column=8, value=f"=H{b1}-H{b2}").number_format = "#,##0"
    ws.cell(row=b3, column=10, value="ایران مقصد").fill = fill_pink_light
    ws.cell(row=b3, column=11, value=f"=J{sum_r}").fill = fill_pink_light
    ws.cell(row=b3, column=11).number_format = "$#,##0"

    b4 = sum_r + 4
    ws.cell(row=b4, column=10, value="متفرقه").fill = fill_pink_light
    ws.cell(row=b4, column=11, value=f"=L{sum_r}").fill = fill_pink_light
    ws.cell(row=b4, column=11).number_format = "$#,##0"

    b5 = sum_r + 5
    ws.cell(row=b5, column=10, value="جمع کل").fill = fill_green
    ws.cell(row=b5, column=11, value=f"=SUM(K{b2}:K{b4})").fill = fill_green
    ws.cell(row=b5, column=11).number_format = "$#,##0"

    for r_i in range(b1, b5 + 1):
        for c_i in range(1, 14):
            cell = ws.cell(row=r_i, column=c_i)
            cell.alignment = center_align
            if cell.value is not None:
                cell.border = border_row
                cell.font = font_bold

    widths = [14, 14, 10, 16, 14, 22, 20, 14, 14, 14, 14, 14, 24]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    path = os.path.join(OUT_DIR, f"hesabdari_{year_int}_{month_int:02d}.xlsx")
    wb.save(path)
    return path


# ============================================================
# منوهای تلگرام و گزارش‌سازهای گرافیکی و متنی
# ============================================================
def admin_menu():
    return ReplyKeyboardMarkup(
        [
            ["📝 ثبت گزارش جدید", "⚡ ثبت فوق‌سریع"],
            ["🔎 جستجوی گزارش", "📊 گزارش بازه‌ای"],
            ["📅 گزارش امروز", "📈 آمار ماهانه"],
            ["🔢 جستجوی سریال", "📤 ارسال گزارش"],
            ["💰 صورت‌حساب مالی"],
            ["⚙️ مدیریت", "❌ لغو"],
        ],
        resize_keyboard=True,
    )


def user_menu():
    return ReplyKeyboardMarkup(
        [
            ["🔎 جستجوی گزارش", "📊 گزارش بازه‌ای"],
            ["📅 گزارش امروز", "📈 آمار ماهانه"],
            ["🔢 جستجوی سریال", "💰 صورت‌حساب مالی"],
            ["❌ لغو"],
        ],
        resize_keyboard=True,
    )


def employee_menu():
    return ReplyKeyboardMarkup(
        [
            ["📝 ثبت آمار شعبه", "📋 وضعیت گزارش‌های من"],
            ["📊 وضعیت شعبه من", "❓ راهنمای کارمند"],
            ["❌ لغو ثبت آمار"],
        ], resize_keyboard=True, is_persistent=True
    )


def management_menu():
    return ReplyKeyboardMarkup(
        [
            ["✏️ ویرایش گزارش", "🗑 حذف گزارش"],
            ["📋 لیست گزارش‌ها", "📅 روزهای جامانده"],
            ["🔍 آمار شعبه", "👥 مدیریت کاربران"],
            ["📊 وضعیت دو شعبه", "📢 پیام همگانی"],
            ["📊 داشبورد سیستم", "🔐 کنترل روز کاری"],
            ["💾 پشتیبان‌گیری", "📥 بازیابی دیتابیس"],
            ["🧠 تحلیل هوشمند", "🩺 سلامت سیستم"],
            ["👷 مدیریت کارمندان"],
            ["📜 لاگ مدیران", "🔙 بازگشت"],
        ],
        resize_keyboard=True,
    )


def get_menu(u):
    if is_admin(u):
        return admin_menu()
    try:
        uid = u.effective_user.id if hasattr(u, "effective_user") else u.id
        if get_employee(uid):
            return employee_menu()
    except Exception:
        pass
    return user_menu()


def create_branch_report(d):
    cb = len(d.get("branches", []))
    r = f"باسلام\n\n({ {1:'یک', 2:'دو', 3:'سه', 4:'چهار'}.get(cb, str(cb)) } شعبه مرز باشماق)\n\nدر روز {d['date']}\nتاریخ شمسی: {d['shamsi_date']}\n\n"
    for i, b in enumerate(d["branches"], 1):
        r += f"شعبه {i} ({b['name']}) مرز باشماق\n\nتعداد صادره {b['issued']} فقره مانیفست از سریال {get_branch_ranges_text(b)} درمرز باشماق صادر گردید\nتعداد ابطالی {b.get('cancelled', 0)} فقره\nسریال ابطالی: {b.get('cancelled_serial', '0')}\n\n"
    r += f"جمع کل صادرها : {total_issued(d)} فقره"
    return r


def create_trade_report(d):
    return f"شعبه مرز باشماق\n\n{d['date']}\n{d['shamsi_date']}\n\nتعداد صادره\nترانزیت : {d['transit']}\nصادرات : {d['export']}"


def make_whatsapp_url(d, phone=None):
    return f"https://wa.me/{phone if phone else ''}?text={urllib.parse.quote(create_branch_report(d) + str(chr(10)*2) + '────────────────' + str(chr(10)*2) + create_trade_report(d))}"


def whatsapp_share_kb(d, phone=None):
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("🟢 ارسال گزارش کامل به واتساپ", url=make_whatsapp_url(d, phone))],
        [
            InlineKeyboardButton("🟢 فقط گزارش شعب", url=f"https://wa.me/{phone if phone else ''}?text={urllib.parse.quote(create_branch_report(d))}"),
            InlineKeyboardButton("🟢 فقط ترانزیت/صادرات", url=f"https://wa.me/{phone if phone else ''}?text={urllib.parse.quote(create_trade_report(d))}"),
        ],
    ])


def whatsapp_range_kb(s, e):
    return InlineKeyboardMarkup([[InlineKeyboardButton("🟢 اشتراک‌گذاری این آمار در واتساپ", url=f"https://wa.me/?text={urllib.parse.quote(range_text(s, e))}")]])


def range_summary(s, e):
    rows = [r for r in all_reports() if s <= parse_smart_date(r["date"]) <= e]
    rows.sort(key=lambda x: parse_smart_date(x["date"]))
    iss = sum(total_issued(x) for x in rows)
    tr = sum(int(x.get("transit", 0)) for x in rows)
    ex = sum(int(x.get("export", 0)) for x in rows)
    bt = {}
    for x in rows:
        for b in x.get("branches", []):
            bn = b.get("name", "بدون نام")
            bt[bn] = bt.get(bn, 0) + int(b.get("issued", 0))
    return rows, iss, tr, ex, sum(total_cancelled(x) for x in rows), bt


def range_text(s, e):
    rows, iss, tr, ex, canc, bt = range_summary(s, e)
    res = f"📊 جمع گزارش‌های بازه‌ای مرز باشماق\n\n📅 میلادی: {s.strftime('%d/%m/%Y')} تا {e.strftime('%d/%m/%Y')}\n☀️ شمسی: {shamsi(s)} تا {shamsi(e)}\n\n• روزهای ثبت‌شده: {len(rows)} روز | میانگین روزانه: {round(iss/len(rows), 1) if rows else 0} فقره\n\n✅ کل صادره: {iss} فقره | ❌ کل ابطالی: {canc}\n🚛 ترانزیت: {tr} | 📦 صادرات: {ex}"
    if bt:
        res += "\n\n🏢 تفکیک شعبه‌ها:\n" + "\n".join(f"▫️ {n}: {v} فقره" for n, v in bt.items())
    if not rows:
        res += "\n\n❌ هیچ گزارشی در این بازه یافت نشد."
    return res


def create_excel_report(s, e):
    rows, iss, tr, ex, canc, bt = range_summary(s, e)
    wb = openpyxl.Workbook()
    hf = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
    hf_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    cnt = Alignment(horizontal="center", vertical="center")
    brd = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    w1 = wb.active
    w1.title, w1.sheet_view.rightToLeft = "خلاصه روزانه", True
    w1.append(["ردیف", "تاریخ میلادی", "تاریخ شمسی", "ترانزیت", "صادرات", "مجموع ابطالی", "مجموع صادره"])
    for i, it in enumerate(rows, 1):
        w1.append([i, it.get("date"), it.get("shamsi_date"), int(it.get("transit", 0)), int(it.get("export", 0)), total_cancelled(it), total_issued(it)])
    w1.append(["جمع کل", "-", "-", tr, ex, canc, iss])
    w2 = wb.create_sheet(title="ریز عملکرد شعب")
    w2.sheet_view.rightToLeft = True
    w2.append(["تاریخ میلادی", "تاریخ شمسی", "نام شعبه", "تعداد صادره", "بازه سریال‌ها", "تعداد ابطالی", "سریال ابطالی"])
    for it in rows:
        for b in it.get("branches", []):
            w2.append([it.get("date"), it.get("shamsi_date"), b.get("name"), int(b.get("issued", 0)), get_branch_ranges_text(b), int(b.get("cancelled", 0)), b.get("cancelled_serial")])
    w3 = wb.create_sheet(title="تجمیع شعب")
    w3.sheet_view.rightToLeft = True
    w3.append(["نام شعبه", "مجموع کل صادره در بازه"])
    for n, v in bt.items():
        w3.append([n, v])
    for w in [w1, w2, w3]:
        for r in w.iter_rows():
            for c in r:
                c.alignment, c.border = cnt, brd
                if c.row == 1:
                    c.font, c.fill = hf, hf_fill
                elif w == w1 and c.row == w1.max_row:
                    c.font, c.fill = Font(name="Tahoma", size=11, bold=True), PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
                else:
                    c.font = Font(name="Tahoma", size=10)
        for cl in w.columns:
            w.column_dimensions[cl[0].column_letter].width = 18
    w2.column_dimensions["E"].width = 35
    path = os.path.join(OUT_DIR, f"bashmaq_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    wb.save(path)
    return path


def create_chart_report(s, e):
    rows, _, _, _, _, _ = range_summary(s, e)
    if not rows:
        return None
    dt = [r.get("shamsi_date", r["date"]).replace("/", "-") for r in rows]
    tot = [total_issued(r) for r in rows]
    trn = [int(r.get("transit", 0)) for r in rows]
    exp = [int(r.get("export", 0)) for r in rows]
    fig, ax = plt.subplots(figsize=(11, 5.5))
    ax.plot(dt, tot, marker="o", color="#4f46e5", linewidth=3, label=ptext("کل صادره"))
    ax.plot(dt, exp, marker="s", linestyle="--", color="#10b981", linewidth=2, label=ptext("صادرات"))
    ax.plot(dt, trn, marker="^", linestyle=":", color="#f97316", linewidth=2, label=ptext("ترانزیت"))
    p = MPL_FONT_PROP if MPL_FONT_PROP else None
    ax.set_xlabel(ptext("تاریخ شمسی"), fontsize=11, fontproperties=p)
    ax.set_title(ptext("نمودار روند صدور مانیفست، صادرات و ترانزیت"), fontsize=13, fontproperties=p)
    ax.legend(prop=p)
    plt.xticks(rotation=45, ha="right")
    ax.grid(True, linestyle="--", alpha=0.5)
    plt.tight_layout()
    path = os.path.join(OUT_DIR, f"chart_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.png")
    plt.savefig(path, dpi=300)
    plt.close(fig)
    return path


def draw_rtl_text(draw, rx, y, txt, fnt, fill="black", dc=None):
    if dc:
        draw.ellipse([rx - 14, y + 14, rx, y + 28], fill=dc)
        rx -= 26
    d = ptext(txt)
    bb = draw.textbbox((0, 0), d, font=fnt)
    draw.text((rx - (bb[2] - bb[0]), y), d, font=fnt, fill=fill)


def create_jpg_report(d, pref="gozaresh"):
    if not FONT_PATH:
        raise RuntimeError("Font not found")
    ft = ImageFont.truetype(FONT_PATH, 44)
    fs = ImageFont.truetype(FONT_PATH, 30)
    fc = ImageFont.truetype(FONT_PATH, 34)
    fn = ImageFont.truetype(FONT_PATH, 28)
    fb = ImageFont.truetype(FONT_PATH, 32)
    ff = ImageFont.truetype(FONT_PATH, 26)
    brs = d.get("branches", [])
    w, pad, ch, hh, sh, fh = 1400, 60, 235, 225, 210, 100
    h = hh + (len(brs) * (ch + 28)) + sh + fh + 75
    img = Image.new("RGB", (w, h), "#F4F6F9")
    draw = ImageDraw.Draw(img)
    draw.rectangle([0, 0, w, hh], fill="#1e1b4b")
    draw.text((w // 2 - draw.textbbox((0, 0), ptext("گزارش رسمی مانیفست مرز باشماق"), font=ft)[2] // 2, 32), ptext("گزارش رسمی مانیفست مرز باشماق"), font=ft, fill="#FFFFFF")
    draw.rounded_rectangle([pad, 120, w - pad, 192], radius=12, fill="#3730a3")
    draw_rtl_text(draw, w - pad - 35, 138, f"تاریخ شمسی :  {str(d.get('shamsi_date', '-')).replace('/','-')}", fs, fill="#a5b4fc")
    draw_rtl_text(draw, w // 2 - 20, 138, f"تاریخ میلادی :  {str(d.get('date', '-')).replace('/','-')}", fs, fill="#FFFFFF")
    y = hh + 30
    for i, b in enumerate(brs, 1):
        draw.rounded_rectangle([pad, y, w - pad, y + ch], radius=15, fill="#FFFFFF", outline="#e2e8f0", width=3)
        draw.rounded_rectangle([w - pad - 14, y, w - pad, y + ch], radius=6, fill="#4f46e5")
        draw_rtl_text(draw, w - pad - 35, y + 20, f"شعبه {i}  -  {b.get('name', '-')}", fc, fill="#1e1b4b")
        draw.line([pad + 30, y + 75, w - pad - 30, y + 75], fill="#e2e8f0", width=2)
        draw_rtl_text(draw, w - pad - 35, y + 95, f"تعداد صادره :  {b.get('issued', 0)} فقره", fb, fill="#15803d", dc="#15803d")
        rs = get_branch_ranges_text(b)
        fr = fn
        if len(rs) > 40:
            fr = ImageFont.truetype(FONT_PATH, 24)
        if len(rs) > 60:
            fr = ImageFont.truetype(FONT_PATH, 20)
        draw_rtl_text(draw, w - pad - 35, y + 155, f"بازه سریال :  {rs}", fr, fill="#334155", dc="#4f46e5")
        cc = int(b.get("cancelled", 0))
        draw_rtl_text(draw, w // 2 - 30, y + 95, f"ابطالی :  {cc} فقره", fn, fill="#b91c1c" if cc > 0 else "#64748b", dc="#b91c1c" if cc > 0 else "#64748b")
        draw_rtl_text(draw, w // 2 - 30, y + 155, f"سریال ابطالی :  {b.get('cancelled_serial', '0')}", fn, fill="#b91c1c" if cc > 0 else "#64748b", dc="#b91c1c" if cc > 0 else "#64748b")
        y += ch + 28
    draw.rounded_rectangle([pad, y, w - pad, y + sh], radius=15, fill="#eef2ff", outline="#4f46e5", width=3)
    draw_rtl_text(draw, w - pad - 35, y + 28, f"جمع کل صادره :  {total_issued(d)} فقره", fb, fill="#1e1b4b", dc="#1e1b4b")
    draw_rtl_text(draw, w // 2 - 30, y + 28, f"کل ابطالی :  {total_cancelled(d)} فقره", fb, fill="#b91c1c", dc="#b91c1c")
    draw.line([pad + 30, y + 98, w - pad - 30, y + 98], fill="#c7d2fe", width=2)
    draw_rtl_text(draw, w - pad - 35, y + 122, f"ترانزیت :  {d.get('transit', 0)} دستگاه", fc, fill="#15803d", dc="#15803d")
    draw_rtl_text(draw, w // 2 - 30, y + 122, f"صادرات :  {d.get('export', 0)} دستگاه", fc, fill="#15803d", dc="#15803d")
    y += sh + 22
    draw.rounded_rectangle([pad, y, w - pad, y + 68], radius=10, fill="#0f172a")
    draw.text((w // 2 - draw.textbbox((0, 0), ptext(FIXED_FOOTER), font=ff)[2] // 2, y + 18), ptext(FIXED_FOOTER), font=ff, fill="#FFFFFF")
    path = os.path.join(OUT_DIR, f"{pref}_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.jpg")
    img.save(path, "JPEG", quality=95)
    return path


def create_pdf_report(d, pref="gozaresh"):
    path = os.path.join(OUT_DIR, f"{pref}_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.pdf")
    c, w, h = canvas.Canvas(path, pagesize=A4), A4[0], A4[1]
    fnt = PDF_FONT_NAME or "Helvetica"
    c.setFillColor(colors.HexColor("#1e1b4b"))
    c.rect(0, h - 95, w, 95, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont(fnt, 18)
    c.drawCentredString(w / 2, h - 42, ptext("گزارش رسمی مانیفست مرز باشماق") if PDF_FONT_NAME else "Report")
    c.setFont(fnt, 12)
    c.drawCentredString(w / 2, h - 75, ptext(f"تاریخ شمسی : {str(d.get('shamsi_date', '-')).replace('/','-')}   -   میلادی : {str(d.get('date', '-')).replace('/','-')}") if PDF_FONT_NAME else "")
    y = h - 130
    for i, b in enumerate(d.get("branches", []), 1):
        if y < 160:
            c.showPage()
            y = h - 60
        c.setFillColor(colors.HexColor("#f8fafc"))
        c.setStrokeColor(colors.HexColor("#cbd5e1"))
        c.roundRect(35, y - 95, w - 70, 95, 8, fill=1, stroke=1)
        c.setFillColor(colors.HexColor("#1e1b4b"))
        c.setFont(fnt, 13)
        c.drawRightString(w - 50, y - 25, ptext(f"شعبه {i}  -  {b.get('name', '-')} مرز باشماق") if PDF_FONT_NAME else "")
        c.setFillColor(colors.HexColor("#0f172a"))
        c.setFont(fnt, 10)
        c.drawRightString(w - 50, y - 52, ptext(f"تعداد صادره : {b.get('issued', 0)} فقره   -   بازه سریال : {get_branch_ranges_text(b)}") if PDF_FONT_NAME else "")
        c.drawRightString(w - 50, y - 76, ptext(f"تعداد ابطالی : {b.get('cancelled', 0)} فقره   -   سریال ابطالی : {b.get('cancelled_serial', '0')}") if PDF_FONT_NAME else "")
        y -= 115
    if y < 140:
        c.showPage()
        y = h - 60
    c.setFillColor(colors.HexColor("#eef2ff"))
    c.setStrokeColor(colors.HexColor("#4f46e5"))
    c.roundRect(35, y - 80, w - 70, 80, 8, fill=1, stroke=1)
    c.setFillColor(colors.HexColor("#312e81"))
    c.setFont(fnt, 12)
    c.drawRightString(w - 50, y - 30, ptext(f"جمع کل صادره : {total_issued(d)} فقره     -     مجموع ابطالی : {total_cancelled(d)} فقره") if PDF_FONT_NAME else "")
    c.drawRightString(w - 50, y - 58, ptext(f"آمار ترانزیت : {d.get('transit', 0)} دستگاه     -     آمار صادرات : {d.get('export', 0)} دستگاه") if PDF_FONT_NAME else "")
    c.setFillColor(colors.HexColor("#334155"))
    c.setFont(fnt, 10)
    c.drawCentredString(w / 2, 35, ptext(FIXED_FOOTER) if PDF_FONT_NAME else FIXED_FOOTER)
    c.save()
    return path


def create_range_jpg(s, e):
    _, iss, tr, ex, _, bt = range_summary(s, e)
    bl = [{"name": n, "issued": v, "ranges": [{"start": "-", "end": "-"}], "cancelled": 0, "cancelled_serial": "-"} for n, v in bt.items()]
    return create_jpg_report({"date": f"{s.strftime('%d-%m-%Y')} تا {e.strftime('%d-%m-%Y')}", "shamsi_date": f"{shamsi(s).replace('/','-')} تا {shamsi(e).replace('/','-')}", "branches": bl, "transit": tr, "export": ex}, "jam_bazeh")


def create_range_pdf(s, e):
    rows, iss, tr, ex, canc, bt = range_summary(s, e)
    path = os.path.join(OUT_DIR, f"range_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")
    c, w, h = canvas.Canvas(path, pagesize=A4), A4[0], A4[1]
    fnt = PDF_FONT_NAME or "Helvetica"
    c.setFillColor(colors.HexColor("#1e1b4b"))
    c.rect(0, h - 90, w, 90, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont(fnt, 17)
    c.drawCentredString(w / 2, h - 38, ptext("گزارش تجمیعی بازه‌ای مرز باشماق") if PDF_FONT_NAME else "")
    c.setFont(fnt, 11)
    c.drawCentredString(w / 2, h - 68, ptext(f"از شمسی {shamsi(s).replace('/','-')} تا {shamsi(e).replace('/','-')}") if PDF_FONT_NAME else "")
    y = h - 120
    c.setFillColor(colors.HexColor("#f0fdfa"))
    c.setStrokeColor(colors.HexColor("#0d9488"))
    c.roundRect(35, y - 65, w - 70, 65, 8, fill=1, stroke=1)
    c.setFillColor(colors.HexColor("#0f172a"))
    c.setFont(fnt, 11)
    c.drawRightString(w - 50, y - 24, ptext(f"روزهای کاری : {len(rows)} روز") if PDF_FONT_NAME else "")
    c.drawRightString(w - 50, y - 48, ptext(f"صادره : {iss}  |  ابطالی : {canc}  |  ترانزیت : {tr}  |  صادرات : {ex}") if PDF_FONT_NAME else "")
    y -= 90
    c.setFillColor(colors.HexColor("#1e1b4b"))
    c.setFont(fnt, 13)
    c.drawRightString(w - 40, y, ptext("عملکرد تفکیکی شعب:") if PDF_FONT_NAME else "")
    y -= 10
    c.setStrokeColor(colors.HexColor("#cbd5e1"))
    c.setLineWidth(1)
    c.line(40, y, w - 40, y)
    y -= 22
    c.setFillColor(colors.HexColor("#1e293b"))
    c.setFont(fnt, 11)
    for n, v in bt.items():
        c.drawRightString(w - 50, y, ptext(f"شعبه {n} : {v} فقره") if PDF_FONT_NAME else "")
        y -= 20
    y -= 10
    c.setFillColor(colors.HexColor("#1e1b4b"))
    c.setFont(fnt, 13)
    c.drawRightString(w - 40, y, ptext("ریز گزارش روزها:") if PDF_FONT_NAME else "")
    y -= 10
    c.line(40, y, w - 40, y)
    y -= 22
    c.setFillColor(colors.HexColor("#0f172a"))
    c.setFont(fnt, 10)
    for it in rows:
        if y < 60:
            c.showPage()
            c.setFont(fnt, 10)
            y = h - 50
        c.drawRightString(w - 45, y, ptext(f"تاریخ {it.get('shamsi_date', '-').replace('/','-')}  /  صادره: {total_issued(it)}  |  ترانزیت: {it.get('transit', 0)}  |  صادرات: {it.get('export', 0)}") if PDF_FONT_NAME else "")
        y -= 19
    c.save()
    return path


# ============================================================
# بخش اول: وب‌سایت فوق‌پیشرفته رویال (مجهز به جدول حسابداری ماهانه)
# ============================================================
flask_app = Flask(__name__)
def _load_flask_secret():
    """کلید پایدار نشست؛ نشست‌ها بعد از راه‌اندازی مجدد بی‌دلیل از بین نمی‌روند."""
    env_key = os.getenv("FLASK_SECRET_KEY", "").strip()
    if env_key:
        return env_key
    key_path = os.path.join(BASE_DIR, ".flask_secret_key")
    try:
        if os.path.exists(key_path):
            with open(key_path, "r", encoding="utf-8") as key_file:
                key = key_file.read().strip()
                if len(key) >= 32:
                    return key
        key = secrets.token_hex(32)
        fd = os.open(key_path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
        with os.fdopen(fd, "w", encoding="utf-8") as key_file:
            key_file.write(key)
        return key
    except Exception:
        logger.warning("Could not persist Flask secret key; using an ephemeral key")
        return secrets.token_hex(32)

flask_app.secret_key = _load_flask_secret()
def _session_hours():
    """مدت اعتبار نشست ورود به پنل (ساعت)؛ با متغیر SESSION_HOURS قابل‌تنظیم است (۱ تا ۲۱۶۰، پیش‌فرض ۸)."""
    try:
        return max(1, min(2160, int(os.getenv("SESSION_HOURS", "8"))))
    except ValueError:
        return 8


flask_app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.getenv("COOKIE_SECURE", "0").strip() == "1",
    PERMANENT_SESSION_LIFETIME=timedelta(hours=_session_hours()),
    MAX_CONTENT_LENGTH=max(1, int(os.getenv("MAX_UPLOAD_MB", "64"))) * 1024 * 1024,
)
app = flask_app
_LOGIN_ATTEMPTS = {}

WEB_TEMPLATE = """
<!DOCTYPE html>
<html lang="fa" dir="rtl" data-bs-theme="light">
<head>
    <meta name="csrf-token" content="{{ csrf_token() }}">
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>سامانه هوشمند مرز باشماق | نسخه رویال ۲۰۲۶</title>
    <link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/css/bootstrap.rtl.min.css">
    <link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/bootstrap-icons@1.11.3/font/bootstrap-icons.min.css">
    <link href="https://cdn.jsdelivr.net/gh/rastikerdar/vazirmatn@v33.003/Vazirmatn-font-face.css" rel="stylesheet" />
    <link rel="stylesheet" href="https://unpkg.com/@majidh1/jalalidatepicker/dist/jalalidatepicker.min.css">
    <script src="https://unpkg.com/@majidh1/jalalidatepicker/dist/jalalidatepicker.min.js"></script>
    <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
    <style>
        :root {
            --bg-body: #f1f5f9;
            --card-bg: rgba(255, 255, 255, 0.94);
            --card-border: rgba(203, 213, 225, 0.7);
            --primary-grad: linear-gradient(135deg, #312e81 0%, #4f46e5 50%, #7c3aed 100%);
            --success-grad: linear-gradient(135deg, #065f46 0%, #059669 50%, #10b981 100%);
            --info-grad: linear-gradient(135deg, #0c4a6e 0%, #0284c7 50%, #06b6d4 100%);
            --warn-grad: linear-gradient(135deg, #9a3412 0%, #ea580c 50%, #f97316 100%);
            --gold-grad: linear-gradient(135deg, #78350f 0%, #d97706 50%, #fbbf24 100%);
            --danger-grad: linear-gradient(135deg, #7f1d1d 0%, #dc2626 50%, #ef4444 100%);
        }
        [data-bs-theme="dark"] {
            --bg-body: #090d16;
            --card-bg: rgba(17, 24, 39, 0.92);
            --card-border: rgba(55, 65, 81, 0.8);
        }
        body { font-family: 'Vazirmatn', Tahoma, sans-serif; background-color: var(--bg-body); min-height: 100vh; display: flex; flex-direction: column; }
        .navbar-glass { background: var(--primary-grad); box-shadow: 0 8px 32px rgba(49, 46, 129, 0.28); }
        .glass-card { background: var(--card-bg); backdrop-filter: blur(16px); border: 1px solid var(--card-border); border-radius: 22px; box-shadow: 0 12px 35px rgba(0, 0, 0, 0.05); }
        .stat-box { border-radius: 20px; color: #fff; padding: 1.2rem; position: relative; overflow: hidden; box-shadow: 0 10px 25px rgba(0,0,0,0.1); transition: transform 0.25s ease; }
        .stat-box:hover { transform: translateY(-4px); }
        .stat-box .stat-icon { position: absolute; left: 12px; bottom: -12px; font-size: 4rem; opacity: 0.18; }
        .bg-grad-primary { background: var(--primary-grad); }
        .bg-grad-success { background: var(--success-grad); }
        .bg-grad-info { background: var(--info-grad); }
        .bg-grad-warn { background: var(--warn-grad); }
        .bg-grad-gold { background: var(--gold-grad); }
        .bg-grad-danger { background: var(--danger-grad); }
        .branch-box { background: rgba(79, 70, 229, 0.04); border: 1.5px solid rgba(79, 70, 229, 0.18); border-radius: 18px; padding: 16px; margin-bottom: 14px; }
        .range-row-box { background: rgba(255,255,255,0.7); border: 1px dashed rgba(99, 102, 241, 0.4); border-radius: 12px; padding: 10px; margin-bottom: 8px; }
        [data-bs-theme="dark"] .range-row-box { background: rgba(30, 41, 59, 0.7); }
        .nav-pills .nav-link { border-radius: 14px; font-weight: 700; padding: 10px 16px; color: inherit; }
        .nav-pills .nav-link.active { background: var(--primary-grad); color: #fff !important; box-shadow: 0 6px 18px rgba(79, 70, 229, 0.35); }
        .table thead th { background: #1e1b4b; color: #fff; font-weight: 600; border: none; padding: 10px; }
        .btn-wa { background: linear-gradient(135deg, #16a34a 0%, #22c55e 100%); color: white !important; border: none; }
        .magic-box { background: linear-gradient(135deg, rgba(79,70,229,0.1) 0%, rgba(124,58,237,0.15) 100%); border: 1.5px dashed #6366f1; border-radius: 16px; padding: 12px 15px; }
        .acc-table th, .acc-table td { border: 1px solid #64748b !important; padding: 4px 6px !important; vertical-align: middle; font-size: 0.82rem; }
        .acc-table thead th { background: #f8fafc !important; color: #0f172a !important; font-weight: 800; border-bottom: 2px solid #0f172a !important; }
        .acc-col-grey { background-color: #c5b3b3 !important; color: #000 !important; }
        .acc-col-pink-dark { background-color: #e69999 !important; color: #000 !important; }
        .acc-col-pink-light { background-color: #f2caca !important; color: #000 !important; }
        .acc-cell-blue { background-color: #5bc0eb !important; color: #000 !important; font-weight: 800; }
        .acc-cell-green { background-color: #70ad47 !important; color: #000 !important; font-weight: 800; }
        .acc-inp { width: 100%; border: none; background: transparent; text-align: center; font-weight: 700; color: inherit; outline: none; }
        .acc-inp:focus { background: rgba(255, 255, 255, 0.85); box-shadow: inset 0 0 0 2px #2563eb; border-radius: 4px; color: #000; }

        /* پوسته تازه و واکنش‌گرا */
        :root { color-scheme: light; --focus-ring: rgba(99,102,241,.42); }
        [data-bs-theme="dark"] { color-scheme: dark; }
        html { scroll-behavior: smooth; scroll-padding-top: 90px; }
        body { background-image: radial-gradient(ellipse at 8% 0%, rgba(99,102,241,.09), transparent 32%), radial-gradient(ellipse at 92% 14%, rgba(6,182,212,.07), transparent 28%); }
        .navbar-glass { position: sticky; top: 0; z-index: 1030; backdrop-filter: blur(18px) saturate(145%); border-bottom: 1px solid rgba(255,255,255,.18); }
        .navbar-brand { letter-spacing: -.25px; }
        .glass-card { transition: transform .22s ease, box-shadow .22s ease, border-color .22s ease; }
        .glass-card:hover { box-shadow: 0 18px 42px rgba(15,23,42,.09); border-color: rgba(99,102,241,.28); }
        .stat-box { min-height: 126px; isolation: isolate; }
        .stat-box::after { content: ""; position: absolute; width: 145px; height: 145px; left: -48px; top: -72px; border: 1px solid rgba(255,255,255,.2); border-radius: 50%; z-index: -1; }
        .btn { transition: transform .16s ease, box-shadow .16s ease, filter .16s ease; }
        .btn:hover { filter: brightness(1.035); }
        .btn:active { transform: translateY(1px) scale(.99); }
        :focus-visible { outline: 3px solid var(--focus-ring) !important; outline-offset: 2px; }
        .form-control, .form-select { border-radius: 12px; min-height: 42px; }
        .form-control:focus, .form-select:focus { border-color: #818cf8; box-shadow: 0 0 0 .22rem var(--focus-ring); }
        .table-responsive { border-radius: 14px; scrollbar-width: thin; }
        .table { --bs-table-bg: transparent; }
        .table tbody tr { transition: background-color .15s ease; }
        .table tbody tr:hover > * { --bs-table-bg: rgba(99,102,241,.055); }
        .glass-card h1, .glass-card h2, .glass-card h3, .glass-card h4 { letter-spacing: -.35px; }
        .section-heading { display:flex; align-items:center; gap:.65rem; font-weight:800; }
        .section-heading::before { content:""; width:5px; height:1.35em; border-radius:9px; background:linear-gradient(180deg,#6366f1,#06b6d4); }
        .live-pill { display:inline-flex; align-items:center; gap:.45rem; padding:.42rem .75rem; border:1px solid rgba(255,255,255,.28); background:rgba(255,255,255,.12); color:#fff; border-radius:999px; font-size:.78rem; white-space:nowrap; }
        .live-dot { width:7px; height:7px; border-radius:50%; background:#4ade80; box-shadow:0 0 0 4px rgba(74,222,128,.15); }
        .app-toast { position:fixed; z-index:2000; bottom:20px; left:20px; max-width:min(420px,calc(100vw - 40px)); padding:12px 16px; border-radius:14px; background:#111827; color:white; box-shadow:0 12px 35px rgba(0,0,0,.22); opacity:0; transform:translateY(12px); pointer-events:none; transition:all .22s ease; }
        .app-toast.show { opacity:1; transform:translateY(0); }
        [data-bs-theme="dark"] body { background-image: radial-gradient(ellipse at 8% 0%, rgba(99,102,241,.13), transparent 32%), radial-gradient(ellipse at 92% 14%, rgba(6,182,212,.09), transparent 28%); }
        [data-bs-theme="dark"] .text-muted { color:#aab5c5 !important; }
        [data-bs-theme="dark"] .form-control, [data-bs-theme="dark"] .form-select { background-color:#111827; border-color:#374151; color:#e5e7eb; }
        [data-bs-theme="dark"] .modal-content { background:#111827; color:#e5e7eb; border:1px solid #374151; }
        [data-bs-theme="dark"] .acc-table thead th { background:#1f2937 !important; color:#f9fafb !important; }
        @media (max-width: 768px) {
          .container-fluid.px-4 { padding-left: .8rem !important; padding-right: .8rem !important; }
          .navbar .container-fluid { padding-left:.8rem !important; padding-right:.8rem !important; }
          .navbar-brand { font-size:.94rem !important; }
          .navbar-brand .badge { display:none; }
          .live-pill { display:none; }
          .glass-card { border-radius:16px; }
          .stat-box { padding:1rem; border-radius:16px; }
          .nav-pills { gap:.35rem; }
          .nav-pills .nav-link { padding:8px 11px; font-size:.88rem; }
          .acc-table th, .acc-table td { white-space:nowrap; font-size:.75rem; }
        }
        @media (prefers-reduced-motion: reduce) {
          *, *::before, *::after { scroll-behavior:auto !important; animation-duration:.01ms !important; transition-duration:.01ms !important; }
        }
        .smart-toolbar {display:flex;gap:.65rem;align-items:center;flex-wrap:wrap;padding:1rem;border-radius:18px;background:linear-gradient(120deg,rgba(79,70,229,.10),rgba(6,182,212,.08));border:1px solid var(--card-border)}
        .smart-toolbar .form-control {min-width:220px;flex:1}
        .insight-item {padding:.8rem 1rem;border-radius:14px;background:rgba(99,102,241,.07);border:1px solid rgba(99,102,241,.13);margin-bottom:.6rem}
        .status-chip {display:inline-flex;align-items:center;gap:.35rem;border-radius:99px;padding:.3rem .65rem;font-size:.78rem;font-weight:800;background:rgba(16,185,129,.12);color:#059669}
        .global-search-results {position:relative;z-index:10;max-height:320px;overflow:auto}
        [data-bs-theme="dark"] .insight-item {background:rgba(99,102,241,.12)}
        @media print {.navbar,.nav-pills,.btn,.smart-toolbar{display:none!important}.glass-card,.stat-box{break-inside:avoid}}

        /* پوسته نهایی: سرمه‌ای + فیروزه‌ای + سبز زمردی */
        :root {
            --bg-body: #eef4f7;
            --card-bg: rgba(255,255,255,.97);
            --card-border: rgba(148,163,184,.30);
            --primary-grad: linear-gradient(135deg,#102a43 0%,#145374 52%,#0f766e 100%);
            --success-grad: linear-gradient(135deg,#065f46 0%,#059669 100%);
            --info-grad: linear-gradient(135deg,#164e63 0%,#0891b2 100%);
            --warn-grad: linear-gradient(135deg,#92400e 0%,#d97706 100%);
            --gold-grad: linear-gradient(135deg,#854d0e 0%,#ca8a04 100%);
            --danger-grad: linear-gradient(135deg,#7f1d1d 0%,#dc2626 100%);
            --focus-ring: rgba(13,148,136,.28);
        }
        [data-bs-theme="dark"] {
            --bg-body: #0b1420;
            --card-bg: rgba(17,30,45,.96);
            --card-border: rgba(100,116,139,.35);
            --focus-ring: rgba(45,212,191,.28);
        }
        body {
            background-image: radial-gradient(ellipse at 0% 0%,rgba(20,83,116,.10),transparent 34%),
                              radial-gradient(ellipse at 100% 8%,rgba(13,148,136,.10),transparent 30%);
            color: var(--bs-body-color, #172b3a);
        }
        .navbar-glass { background: linear-gradient(115deg,#102a43 0%,#145374 58%,#0f766e 100%); }
        .glass-card { border-radius: 20px; box-shadow: 0 10px 28px rgba(16,42,67,.065); }
        .stat-box { border-radius: 18px; box-shadow: 0 12px 24px rgba(16,42,67,.13); }
        .nav-pills .nav-link.active { background: linear-gradient(120deg,#145374,#0f766e); box-shadow: 0 5px 14px rgba(15,118,110,.22); }
        .table thead th { background: #163b52; }
        .btn-primary { --bs-btn-bg:#0f766e; --bs-btn-border-color:#0f766e; --bs-btn-hover-bg:#115e59; --bs-btn-hover-border-color:#115e59; --bs-btn-active-bg:#134e4a; }
        .btn-outline-primary { --bs-btn-color:#0f766e; --bs-btn-border-color:#0f766e; --bs-btn-hover-bg:#0f766e; --bs-btn-hover-border-color:#0f766e; }
        .form-control:focus, .form-select:focus { border-color:#0d9488; box-shadow:0 0 0 .22rem var(--focus-ring); }
        .branch-box { background:rgba(13,148,136,.055); border-color:rgba(13,148,136,.22); }
        .magic-box { background:linear-gradient(135deg,rgba(20,83,116,.07),rgba(13,148,136,.11)); border-color:#0d9488; }
        .section-heading::before { background:linear-gradient(180deg,#145374,#2dd4bf); }
        .table tbody tr:hover > * { --bs-table-bg:rgba(13,148,136,.055); }
        .badge { letter-spacing:.05px; }
        @media (max-width: 768px) {
            .glass-card { border-radius:15px; }
            .stat-box { min-height:105px; padding:1rem; }
            .nav-pills { gap:.35rem; }
            .nav-pills .nav-link { padding:8px 11px; font-size:.9rem; }
            .table { font-size:.88rem; }
        }
        @media (prefers-reduced-motion: reduce) {
            *, *::before, *::after { scroll-behavior:auto !important; transition:none !important; animation:none !important; }
        }
    </style>
</head>
<body>
<nav class="navbar navbar-expand-lg navbar-dark navbar-glass sticky-top mb-4 py-3">
    <div class="container-fluid px-4">
        <a class="navbar-brand fw-bold d-flex align-items-center gap-2 fs-5" href="/">
            <i class="bi bi-shield-check text-warning fs-3"></i>
            <span>سامانه هوشمند مرز باشماق</span>
            <span class="badge bg-warning text-dark rounded-pill" style="font-size: 0.72rem;">VIP 2026</span>
        </a>
        <div class="d-flex align-items-center flex-wrap gap-2">
            <span class="live-pill" title="وضعیت رابط کاربری"><span class="live-dot"></span><span id="liveClock">سامانه آماده</span></span>
            <button class="btn btn-outline-light btn-sm rounded-pill px-3" onclick="toggleTheme()" aria-label="تغییر حالت روشن و تاریک">
                <i class="bi bi-moon-stars-fill"></i> تغییر تم
            </button>
            {% if logged_in %}
            <a href="/backup_db" class="btn btn-warning btn-sm rounded-pill fw-bold px-3">
                <i class="bi bi-cloud-arrow-down-fill"></i> بکاپ دیتابیس
            </a>
            <form method="POST" action="/logout" class="d-inline">
                <button type="submit" class="btn btn-danger btn-sm rounded-pill px-3"><i class="bi bi-box-arrow-left"></i> خروج</button>
            </form>
            <button type="button" class="btn btn-outline-primary btn-sm rounded-pill px-3" onclick="openAccountModal()"><i class="bi bi-person-gear"></i> حساب کاربری{% if session.get('web_user') %} ({{ session.get('web_user') }}){% endif %}</button>
            <div class="modal fade" id="accountModal" tabindex="-1"><div class="modal-dialog modal-dialog-centered"><div class="modal-content">
              <form method="POST" action="/account/update" autocomplete="off">
                <div class="modal-header"><h5 class="modal-title"><i class="bi bi-shield-lock"></i> تغییر نام کاربری و رمز پنل</h5><button type="button" class="btn-close" data-bs-dismiss="modal"></button></div>
                <div class="modal-body">
                  <div class="alert alert-info small">هر کدام از دو بخش را که خواستید تغییر دهید؛ خالی بگذارید تا بدون تغییر بماند. با ذخیره، سایر دستگاه‌ها از پنل خارج می‌شوند.</div>
                  <label class="form-label">نام کاربری جدید <span class="text-muted small">(اختیاری، ۳ تا ۳۲ حرف انگلیسی/عدد)</span></label>
                  <input class="form-control mb-3" name="new_username" maxlength="32" placeholder="{{ session.get('web_user', '') }}" autocomplete="off">
                  <label class="form-label">رمز عبور جدید <span class="text-muted small">(اختیاری، حداقل ۸ کاراکتر)</span></label>
                  <input class="form-control mb-2" type="password" name="new_password" maxlength="128" autocomplete="new-password">
                  <input class="form-control mb-3" type="password" name="confirm_password" maxlength="128" placeholder="تکرار رمز جدید" autocomplete="new-password">
                  <hr><label class="form-label fw-bold">رمز فعلی (برای تأیید)</label>
                  <input class="form-control" type="password" name="current_password" required autocomplete="current-password">
                </div>
                <div class="modal-footer"><button class="btn btn-primary rounded-pill px-4">ذخیره تغییرات</button></div>
              </form></div></div></div>
            <script>
            function openAccountModal(){var m=document.getElementById('accountModal');if(m.parentNode!==document.body)document.body.appendChild(m);bootstrap.Modal.getOrCreateInstance(m).show();}
            </script>
            {% endif %}
        </div>
    </div>
</nav>

<div class="container-fluid px-4 pb-5 flex-grow-1">
    {% with messages = get_flashed_messages(with_categories=true) %}
      {% if messages %}
        {% for c, m in messages %}
          <div class="alert alert-{{ c }} alert-dismissible fade show rounded-4 shadow-sm border-0" role="alert">
            <i class="bi bi-check-circle-fill me-1"></i> {{ m }}
            <button type="button" class="btn-close" data-bs-dismiss="alert"></button>
          </div>
        {% endfor %}
      {% endif %}
    {% endwith %}

    {% if not logged_in %}
    <div class="row justify-content-center align-items-center mt-5">
        <div class="col-md-5 col-lg-4">
            <div class="glass-card p-4 p-md-5 text-center">
                <h4 class="fw-bold mb-3">ورود به پنل مدیریت</h4>
                <form method="POST" action="/login">
                    <div class="form-floating mb-3">
                        <input type="text" name="username" class="form-control rounded-4 text-center fs-5" id="userInp" placeholder="نام کاربری" autocomplete="username" required autofocus>
                        <label for="userInp">نام کاربری</label>
                    </div>
                    <div class="form-floating mb-3">
                        <input type="password" name="password" class="form-control rounded-4 text-center fs-5" id="passInp" placeholder="رمز عبور" autocomplete="current-password" required>
                        <label for="passInp">رمز عبور</label>
                    </div>
                    <button type="submit" class="btn btn-primary btn-lg w-100 rounded-4 fw-bold shadow" style="background: var(--primary-grad); border: none;">
                        ورود به سامانه
                    </button>
                </form>
            </div>
        </div>
    </div>
    {% else %}
    <div class="smart-toolbar mb-4" role="search">
        <div class="fw-bold"><i class="bi bi-stars text-primary fs-5"></i> جست‌وجوی هوشمند</div>
        <input id="globalSearchInput" class="form-control" placeholder="جست‌وجوی تاریخ، شعبه، شماره سریال یا متن گزارش…" autocomplete="off" oninput="runGlobalSearch()" aria-label="جست‌وجوی سراسری">
        <button class="btn btn-primary rounded-pill px-3" onclick="runGlobalSearch()"><i class="bi bi-search"></i> جستجو</button>
        <button class="btn btn-outline-secondary rounded-pill px-3" onclick="window.print()"><i class="bi bi-printer"></i> چاپ داشبورد</button>
        <span class="status-chip"><i class="bi bi-shield-check"></i> نشست محافظت‌شده</span>
        <div id="globalSearchResults" class="global-search-results w-100"></div>
    </div>
    <div class="glass-card p-3 mb-4 d-flex justify-content-between align-items-center flex-wrap gap-3">
      <div><div class="fw-bold"><i class="bi bi-calendar3 text-primary me-1"></i> گزارش ماهانه</div>
      <div class="small text-muted">آمار و نمودارها مربوط به ماه انتخابی است؛ آرشیو پایین صفحه همیشه همه گزارش‌های ثبت‌شده را نگه می‌دارد.</div></div>
      <form method="GET" action="/" class="d-flex gap-2 align-items-center">
        <input type="month" name="month" class="form-control rounded-3" value="{{ selected_month }}" aria-label="انتخاب ماه میلادی">
        <button class="btn btn-primary rounded-pill px-3" type="submit">نمایش آمار</button>
        <a class="btn btn-outline-secondary rounded-pill" href="/">ماه جاری</a>
      </form>
    </div>
    <div class="glass-card p-3 mb-4">
      <div class="d-flex flex-wrap gap-2 align-items-center mb-2">
        <a class="btn btn-sm btn-outline-primary rounded-pill" href="/?month={{ prev_month }}">▶ {{ prev_month }}</a>
        <span class="badge bg-{{ 'success' if month_is_current else 'secondary' }} rounded-pill px-3 py-2">{{ selected_month }} — {{ 'ماه جاری' if month_is_current else 'آرشیو' }}</span>
        {% if next_month %}<a class="btn btn-sm btn-outline-primary rounded-pill" href="/?month={{ next_month }}">{{ next_month }} ◀</a>{% endif %}
        <span class="ms-auto d-flex gap-2 flex-wrap">
          <a class="btn btn-sm btn-success rounded-pill" href="/export_range?start={{ month_start_str }}&end={{ month_end_str }}&fmt=excel">اکسل ماه</a>
          <a class="btn btn-sm btn-danger rounded-pill" href="/export_range?start={{ month_start_str }}&end={{ month_end_str }}&fmt=pdf">PDF ماه</a>
          <a class="btn btn-sm btn-info text-white rounded-pill" href="/export_range?start={{ month_start_str }}&end={{ month_end_str }}&fmt=chart">نمودار ماه</a>
        </span>
      </div>
      <div class="small mb-2">
        {% if month_cmp.cur.days %}
          نسبت به {{ month_cmp.prev.key }}: کل صادره
          <b class="{{ 'text-success' if (month_cmp.issued_pct or 0) >= 0 else 'text-danger' }}">{% if month_cmp.issued_pct is none %}—{% else %}{{ '%+d' % month_cmp.issued_pct }}٪{% endif %}</b>
          ({{ month_cmp.cur.issued }} در برابر {{ month_cmp.prev.issued }}) · میانگین روزانه
          <b>{% if month_cmp.avg_pct is none %}—{% else %}{{ '%+d' % month_cmp.avg_pct }}٪{% endif %}</b>
        {% else %}برای این ماه هنوز گزارشی ثبت نشده است؛ آمار از صفر شروع می‌شود و ماه‌های قبل در آرشیو محفوظ‌اند.{% endif %}
      </div>
      {% if available_months %}<div class="d-flex flex-wrap gap-1">{% for mth in available_months %}<a class="badge rounded-pill text-decoration-none bg-{{ 'primary' if mth.key == selected_month else 'light text-dark border' }}" href="/?month={{ mth.key }}">{{ mth.key }} · {{ mth.total }}</a>{% endfor %}</div>{% endif %}
    </div>

    <div class="row g-3 mb-4">
        <div class="col-6 col-md-4 col-lg-2">
            <div class="stat-box bg-grad-primary h-100">
                <div class="small opacity-75">کل مانیفست صادره</div>
                <div class="fs-3 fw-bold mt-1">{{ stats.total_issued }}</div>
                <div class="small mt-1 opacity-75">فقره ثبت‌شده</div>
                <i class="bi bi-file-earmark-check stat-icon"></i>
            </div>
        </div>
        <div class="col-6 col-md-4 col-lg-2">
            <div class="stat-box bg-grad-success h-100">
                <div class="small opacity-75">کل صادرات</div>
                <div class="fs-3 fw-bold mt-1">{{ stats.total_export }}</div>
                <div class="small mt-1 opacity-75">دستگاه کامیون</div>
                <i class="bi bi-box-seam stat-icon"></i>
            </div>
        </div>
        <div class="col-6 col-md-4 col-lg-2">
            <div class="stat-box bg-grad-info h-100">
                <div class="small opacity-75">کل ترانزیت</div>
                <div class="fs-3 fw-bold mt-1">{{ stats.total_transit }}</div>
                <div class="small mt-1 opacity-75">دستگاه تردد</div>
                <i class="bi bi-truck stat-icon"></i>
            </div>
        </div>
        <div class="col-6 col-md-4 col-lg-2">
            <div class="stat-box bg-grad-warn h-100">
                <div class="small opacity-75">میانگین روزانه</div>
                <div class="fs-3 fw-bold mt-1">{{ stats.avg_daily }}</div>
                <div class="small mt-1 opacity-75">در {{ stats.total_days }} روز کاری</div>
                <i class="bi bi-graph-up-arrow stat-icon"></i>
            </div>
        </div>
        <div class="col-6 col-md-4 col-lg-2">
            <div class="stat-box bg-grad-gold h-100">
                <div class="small opacity-75">رکورد شلوغ‌ترین روز</div>
                <div class="fs-3 fw-bold mt-1">{{ stats.max_day_val }}</div>
                <div class="small mt-1 opacity-75">{{ stats.max_day_date }}</div>
                <i class="bi bi-trophy-fill stat-icon"></i>
            </div>
        </div>
        <div class="col-6 col-md-4 col-lg-2">
            <div class="stat-box bg-grad-danger h-100">
                <div class="small opacity-75">مجموع ابطالی</div>
                <div class="fs-3 fw-bold mt-1">{{ stats.total_cancelled }}</div>
                <div class="small mt-1 opacity-75">فقره مانیفست</div>
                <i class="bi bi-x-octagon-fill stat-icon"></i>
            </div>
        </div>
    </div>

    <ul class="nav nav-pills glass-card p-2 mb-4 gap-2 d-flex justify-content-center flex-wrap" id="mainTabs" role="tablist">
        <li class="nav-item">
            <button class="nav-link active" id="tab-dash-btn" data-bs-toggle="pill" data-bs-target="#tab-dash" type="button">
                <i class="bi bi-magic me-1"></i> ثبت هوشمند گزارش
            </button>
        </li>
        <li class="nav-item">
            <button class="nav-link" id="tab-acc-btn" data-bs-toggle="pill" data-bs-target="#tab-acc" type="button">
                <i class="bi bi-calculator-fill text-warning me-1"></i> 💰 جدول مالی و حسابداری ماهانه (اکسل)
            </button>
        </li>
        <li class="nav-item">
            <button class="nav-link" id="tab-archive-btn" data-bs-toggle="pill" data-bs-target="#tab-archive" type="button">
                <i class="bi bi-journal-album me-1"></i> آرشیو و رهگیری سریال ({{ reports|length }})
            </button>
        </li>
        <li class="nav-item">
            <button class="nav-link" id="tab-missing-btn" data-bs-toggle="pill" data-bs-target="#tab-missing" type="button">
                <i class="bi bi-calendar2-x me-1"></i> روزهای جامانده
                {% if missing_dates %}<span class="badge bg-danger rounded-pill ms-1">{{ missing_dates|length }}</span>{% endif %}
            </button>
        </li>
        <li class="nav-item">
            <button class="nav-link" id="tab-admin-btn" data-bs-toggle="pill" data-bs-target="#tab-admin" type="button">
                <i class="bi bi-sliders me-1"></i> مدیریت کاربران و دیتابیس
            </button>
        </li>
        <li class="nav-item">
            <button class="nav-link" id="tab-employees-btn" data-bs-toggle="pill" data-bs-target="#tab-employees" type="button">
                <i class="bi bi-person-badge me-1"></i> کارکنان شعب و تأیید آمار
                <span class="badge bg-danger rounded-pill ms-1" id="employeePendingBadge">{{ pending_reports|length }}</span>
            </button>
        </li>
    </ul>

    <div id="employeeLiveNotice" class="alert alert-info shadow-sm d-none align-items-center justify-content-between gap-3" role="status" aria-live="polite">
      <div><i class="bi bi-bell-fill me-2"></i><strong id="employeeLiveNoticeTitle">گزارش جدید دریافت شد</strong><div class="small mt-1" id="employeeLiveNoticeText">یک گزارش کارمندی منتظر بررسی است.</div></div>
      <button type="button" class="btn btn-sm btn-primary text-nowrap" onclick="openEmployeeReviewTab()">مشاهده و به‌روزرسانی</button>
    </div>

    <div class="tab-content">
        <!-- تب ۱: ثبت گزارش هوشمند -->
        <div class="tab-pane fade show active" id="tab-dash" role="tabpanel">
            <div class="row g-4">
                <div class="col-lg-5">
                    <div class="glass-card p-4">
                        <div class="d-flex justify-content-between align-items-center mb-3">
                            <h5 class="fw-bold text-primary m-0"><i class="bi bi-stars"></i> ثبت گزارش (تا ۴ بازه سریال)</h5>
                            <span class="badge bg-success text-success bg-opacity-10 px-3 py-2 rounded-pill" id="formModeBadge">آماده ثبت جدید</span>
                        </div>

                        <div class="magic-box mb-3">
                            <label class="small fw-bold text-primary mb-1 d-block">
                                <i class="bi bi-lightning-charge-fill text-warning"></i> پر کردن فوق‌سریع با ۴ عدد (تیرپارک، پارسیان، ترانزیت، صادرات):
                            </label>
                            <div class="input-group input-group-sm">
                                <input type="text" id="magic_inp" class="form-control fw-bold text-center" placeholder="مثلاً: 27 85 33 79">
                                <button type="button" class="btn btn-primary fw-bold px-3" onclick="applyMagicFill()">⚡ چیدمان خودکار</button>
                            </div>
                        </div>

                        <div class="p-2 px-3 rounded-4 bg-primary bg-opacity-10 border border-primary border-opacity-25 d-flex justify-content-between align-items-center mb-3 small fw-bold">
                            <span>🔢 جمع کل صادره: <b class="text-primary fs-5" id="liveTotalIssued">0</b> فقره</span>
                            <span>❌ کل ابطالی: <b class="text-danger fs-6" id="liveTotalCanc">0</b></span>
                        </div>

                        <form method="POST" action="/save_report" id="reportForm">
                            <div class="mb-3"><label class="form-label small">دلیل اصلاح (اگر گزارش این تاریخ از قبل وجود دارد)</label><input class="form-control" name="change_reason" id="change_reason" maxlength="500" placeholder="مثلاً اصلاح اشتباه سریال پس از تطبیق با دفتر"></div>
<input type="hidden" name="json_payload" id="json_payload">
                            <div class="row g-2 mb-3">
                                <div class="col-7">
                                    <label class="form-label small fw-bold">📅 تاریخ گزارش (شمسی/میلادی):</label>
                                    <input type="text" id="report_date" data-jdp class="form-control fw-bold text-center" value="{{ today_shamsi }}" onchange="onDateChange()" required>
                                </div>
                                <div class="col-5 d-flex align-items-end gap-1">
                                    <button type="button" class="btn btn-outline-primary btn-sm w-50 py-2 fw-bold" onclick="setQuickDate('{{ today_shamsi }}')">امروز</button>
                                    <button type="button" class="btn btn-outline-secondary btn-sm w-50 py-2 fw-bold" onclick="setQuickDate('{{ yesterday_shamsi }}')">دیروز</button>
                                </div>
                            </div>

                            <div id="smart_branches_container"></div>

                            <button type="button" class="btn btn-outline-primary btn-sm w-100 rounded-4 mb-3 fw-bold py-2" onclick="addBranch()">
                                <i class="bi bi-plus-circle-dotted"></i> افزودن شعبه جدید به این روز
                            </button>

                            <div class="row g-2 mb-3">
                                <div class="col-6">
                                    <label class="form-label small fw-bold"><i class="bi bi-truck text-info"></i> تعداد ترانزیت:</label>
                                    <input type="number" id="inp_transit" class="form-control fw-bold" value="0" oninput="updateLivePreviewText()" required min="0">
                                </div>
                                <div class="col-6">
                                    <label class="form-label small fw-bold"><i class="bi bi-box-seam text-success"></i> تعداد صادرات:</label>
                                    <input type="number" id="inp_export" class="form-control fw-bold" value="0" oninput="updateLivePreviewText()" required min="0">
                                </div>
                            </div>

                            <div class="form-check form-switch mb-3">
                                <input class="form-check-input" type="checkbox" id="send_tg" checked>
                                <label class="form-check-label small fw-bold" for="send_tg">ارسال خودکار عکس، PDF و بکاپ به تلگرام مدیر</label>
                            </div>

                            <button type="submit" class="btn btn-primary btn-lg w-100 rounded-4 fw-bold shadow" style="background: var(--primary-grad); border: none;">
                                <i class="bi bi-check2-circle me-1"></i> ذخیره نهایی گزارش در دیتابیس
                            </button>
                        </form>
                    </div>
                </div>

                <div class="col-lg-7">
                    <div class="glass-card p-4 mb-4">
                        <div class="d-flex justify-content-between align-items-center mb-2">
                            <h6 class="fw-bold text-primary m-0"><i class="bi bi-file-earmark-text-fill"></i> پیش‌نمایش زنده متن گزارش</h6>
                            <div class="btn-group btn-group-sm">
                                <button type="button" class="btn btn-dark rounded-start-pill px-3" onclick="copyLivePreview()">
                                    <i class="bi bi-clipboard-check"></i> کپی متن
                                </button>
                                <a href="#" id="liveWaLink" target="_blank" class="btn btn-wa rounded-end-pill px-3">
                                    <i class="bi bi-whatsapp"></i> ارسال به واتساپ
                                </a>
                            </div>
                        </div>
                        <textarea id="livePreviewBox" class="form-control border-0 p-3 small" rows="7" readonly></textarea>
                    </div>

                    <div class="glass-card p-4 mb-4">
                        <h6 class="fw-bold mb-3"><i class="bi bi-graph-up text-primary"></i> نمودار روند صدور مانیفست، صادرات و ترانزیت</h6>
                        <canvas id="lineChart" height="115"></canvas>
                    </div>

                    <div class="row g-4">
                        <div class="col-md-6">
                            <div class="glass-card p-4 h-100">
                                <h6 class="fw-bold mb-3"><i class="bi bi-pie-chart-fill text-success"></i> سهم تفکیکی هر شعبه</h6>
                                <canvas id="pieChart" height="175"></canvas>
                            </div>
                        </div>
                        <div class="col-md-6">
                            <div class="glass-card p-4 h-100">
                                <h6 class="fw-bold text-primary mb-3"><i class="bi bi-cloud-download-fill"></i> دانلود گزارش‌های بازه‌ای</h6>
                                <form method="GET" action="/export_range" class="row g-2">
                                    <div class="col-6">
                                        <label class="small fw-bold">از تاریخ:</label>
                                        <input type="text" name="start" data-jdp class="form-control form-control-sm text-center fw-bold" value="{{ first_of_month }}" required>
                                    </div>
                                    <div class="col-6">
                                        <label class="small fw-bold">تا تاریخ:</label>
                                        <input type="text" name="end" data-jdp class="form-control form-control-sm text-center fw-bold" value="{{ today_shamsi }}" required>
                                    </div>
                                    <div class="col-12 mt-3 d-grid gap-2">
                                        <button type="submit" name="fmt" value="excel" class="btn btn-success btn-sm fw-bold rounded-3">دانلود اکسل جامع ۳ برگی</button>
                                        <button type="submit" name="fmt" value="pdf" class="btn btn-danger btn-sm fw-bold rounded-3">دانلود PDF رسمی بازه</button>
                                        <div class="btn-group btn-group-sm">
                                            <button type="submit" name="fmt" value="jpg" class="btn btn-primary fw-bold">عکس تجمیعی</button>
                                            <button type="submit" name="fmt" value="chart" class="btn btn-info text-white fw-bold">عکس نمودار</button>
                                        </div>
                                    </div>
                                </form>
                            </div>
                        </div>
                    </div>
                </div>
            </div>
        </div>

        <!-- تب ۲: جدول مالی و حسابداری ماهانه مطابق فایل 10-2026.xlsm -->
        <div class="tab-pane fade" id="tab-acc" role="tabpanel">
            <div class="glass-card p-4">
                <div class="d-flex flex-column flex-xl-row justify-content-between align-items-xl-center gap-3 mb-4 pb-3 border-bottom">
                    <div>
                        <h5 class="fw-bold text-primary m-0">
                            <i class="bi bi-file-earmark-spreadsheet-fill text-success"></i> دفتر حسابداری و صورت‌حساب ماهانه مرز باشماق
                        </h5>
                        <p class="text-muted small m-0 mt-1">محاسبه خودکار مبلغ دینار، کسورات تحویلی و جدول دلاری ایران مقصد، سبکبار و متفرقه</p>
                    </div>
                    <div class="d-flex flex-wrap align-items-center gap-2">
                        <div class="input-group input-group-sm" style="width: auto;">
                            <span class="input-group-text fw-bold bg-dark text-white">ماه میلادی:</span>
                            <input type="month" id="acc_month_picker" class="form-control fw-bold text-center" value="{{ current_month }}" onchange="loadAccountingSheet(false)">
                        </div>
                        <div class="input-group input-group-sm" style="width: 195px;">
                            <span class="input-group-text fw-bold bg-secondary text-white">نرخ هر آمار:</span>
                            <input type="number" id="acc_dinar_rate" class="form-control fw-bold text-center" value="15000" oninput="recalcAccountingTable()">
                        </div>
                        <button type="button" class="btn btn-outline-primary btn-sm fw-bold rounded-3 px-3" onclick="loadAccountingSheet(true)">
                            <i class="bi bi-arrow-repeat"></i> فراخوانی خودکار آمار از گزارش‌ها
                        </button>
                        <button type="button" class="btn btn-success btn-sm fw-bold rounded-3 px-3 shadow-sm" onclick="saveAccountingSheet()">
                            <i class="bi bi-floppy-fill"></i> ذخیره جدول مالی
                        </button>
                        <button type="button" class="btn btn-danger btn-sm fw-bold rounded-3 px-3 shadow-sm" onclick="downloadAccountingPDF()"><i class="bi bi-file-pdf-fill"></i> PDF رسمی A4 افقی</button>
                        <button type="button" class="btn btn-warning btn-sm fw-bold rounded-3 px-3 shadow-sm" onclick="downloadAccountingExcel()">
                            <i class="bi bi-file-earmark-excel-fill"></i> دانلود اکسل با همین فرمت
                        </button>
                    </div>
                </div>

                <div class="row g-3 mb-3">
                    <div class="col-md-3"><div class="stat-box bg-grad-primary"><div class="small">کل درآمد دیناری</div><div id="live_total_dinar" class="fs-5 fw-bold">0</div></div></div>
                    <div class="col-md-3"><div class="stat-box bg-grad-danger"><div class="small">مجموع کسورات</div><div id="live_deductions" class="fs-5 fw-bold">0</div></div></div>
                    <div class="col-md-3"><div class="stat-box bg-grad-success"><div class="small">مانده خالص قابل مطالبه</div><div id="live_net" class="fs-5 fw-bold">0</div></div></div>
                    <div class="col-md-3"><div class="stat-box bg-grad-info"><div class="small">درآمد دلاری</div><div id="live_usd" class="fs-5 fw-bold">$0</div></div></div>
                </div>
                <div class="d-flex gap-2 mb-3"><button class="btn btn-outline-success btn-sm" onclick="shareAccountingWhatsApp()">🟢 اشتراک خلاصه در واتساپ</button></div>
                <div class="table-responsive">
                    <table class="table acc-table text-center align-middle mb-0">
                        <thead>
                            <tr>
                                <th style="min-width:95px;">تاریخ شمسی</th>
                                <th style="min-width:95px;">تاریخ میلادی</th>
                                <th style="min-width:65px;">آمار</th>
                                <th style="min-width:110px;">مبلغ دینار</th>
                                <th style="min-width:100px;">حواله عراق</th>
                                <th style="min-width:130px;">مبلغ تحویلی به آقای ترتیفی</th>
                                <th style="min-width:130px;">مبلغ تحویلی به ناصردینار</th>
                                <th style="min-width:90px;">مبلغ دلار</th>
                                <th class="acc-col-grey" style="min-width:90px;">نرخ تبدیل</th>
                                <th class="acc-col-pink-dark" style="min-width:95px;">ایران مقصد</th>
                                <th class="acc-col-pink-light" style="min-width:95px;">سبکبار</th>
                                <th class="acc-col-pink-light" style="min-width:95px;">متفرقه</th>
                                <th style="min-width:160px;">توضیحات</th>
                            </tr>
                        </thead>
                        <tbody id="acc_tbody"></tbody>
                        <tfoot>
                            <tr class="fw-bold">
                                <td colspan="2" class="bg-light">جمع ستون‌ها</td>
                                <td class="acc-cell-blue" id="acc_sum_amar">0</td>
                                <td class="acc-col-pink-dark fw-bold" id="acc_sum_dinar">0</td>
                                <td id="acc_sum_hawala">0</td>
                                <td id="acc_sum_torshizi">0</td>
                                <td id="acc_sum_naser">0</td>
                                <td id="acc_sum_usd">$0</td>
                                <td class="acc-col-grey"></td>
                                <td class="acc-col-pink-dark" id="acc_sum_im">0</td>
                                <td class="acc-col-pink-light" id="acc_sum_sb">0</td>
                                <td class="acc-col-pink-light" id="acc_sum_mt">0</td>
                                <td></td>
                            </tr>
                        </tfoot>
                    </table>
                </div>

                <div class="row g-4 mt-3 align-items-end">
                    <div class="col-md-3 text-center">
                        <div class="p-3 border-bottom border-2 border-dark fw-bold fs-6">سید آکو کهنه پوشی</div>
                        <div class="small text-muted mt-1">امضا و تایید نماینده مرز باشماق</div>
                    </div>
                    <div class="col-md-5">
                        <table class="table acc-table text-center mb-0">
                            <tbody>
                                <tr>
                                    <td class="fw-bold bg-light" style="width:45%;">جمع کل (دینار)</td>
                                    <td class="acc-col-pink-dark fw-bold fs-6" id="box_total_dinar">0</td>
                                </tr>
                                <tr>
                                    <td class="fw-bold bg-light">جمع کسورات</td>
                                    <td class="fw-bold" id="box_total_deduct">0</td>
                                </tr>
                                <tr>
                                    <td class="fw-bold bg-light">مبلغ نهایی (دینار)</td>
                                    <td class="fw-bold text-primary fs-5" id="box_final_dinar">0</td>
                                </tr>
                            </tbody>
                        </table>
                    </div>
                    <div class="col-md-4">
                        <table class="table acc-table text-center mb-0">
                            <tbody>
                                <tr>
                                    <td class="acc-cell-blue" style="width:50%;">جمع آمار</td>
                                    <td class="acc-cell-blue fs-6" id="box_amar_2">0</td>
                                </tr>
                                <tr>
                                    <td class="acc-col-pink-light fw-bold">سبکبار</td>
                                    <td class="acc-col-pink-light fw-bold" id="box_sb_usd">$ -</td>
                                </tr>
                                <tr>
                                    <td class="acc-col-pink-light fw-bold">ایران مقصد</td>
                                    <td class="acc-col-pink-light fw-bold" id="box_im_usd">$ -</td>
                                </tr>
                                <tr>
                                    <td class="acc-col-pink-light fw-bold">متفرقه</td>
                                    <td class="acc-col-pink-light fw-bold" id="box_mt_usd">$ -</td>
                                </tr>
                                <tr>
                                    <td class="acc-cell-green">جمع کل (دلار)</td>
                                    <td class="acc-cell-green fs-6" id="box_grand_usd">$ -</td>
                                </tr>
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
        </div>

        <!-- تب ۳: آرشیو و رهگیری سریال -->
        <div class="tab-pane fade" id="tab-archive" role="tabpanel">
            <div class="glass-card p-4 mb-4">
                <h6 class="fw-bold mb-3"><i class="bi bi-upc-scan text-primary"></i> رهگیری آنی شماره سریال مانیفست</h6>
                <div class="input-group">
                    <span class="input-group-text bg-primary text-white fw-bold">شماره سریال:</span>
                    <input type="number" id="serial_search_inp" class="form-control fw-bold" placeholder="مثلاً 25700...">
                    <button class="btn btn-primary fw-bold px-4" onclick="searchSerial()">جستجو</button>
                </div>
                <div id="serial_result" class="mt-3"></div><hr><h6 class="fw-bold">🕵️ بازرس هوشمند سریال‌ها</h6><p class="small text-muted">تشخیص پرش و تداخل بازه‌های سریال در گزارش‌های ثبت‌شده شعب</p><button class="btn btn-outline-danger btn-sm" onclick="auditSerials()">اجرای بررسی همه گزارش‌ها</button><div id="serial_audit_result" class="mt-3"></div>
            </div>

            <div class="glass-card p-4">
                <div class="d-flex justify-content-between align-items-center mb-3">
                    <h5 class="fw-bold m-0">آرشیو گزارش‌ها ({{ reports|length }} روز)</h5>
                    <input type="text" id="tableSearch" class="form-control form-control-sm w-auto" placeholder="🔍 جستجو..." onkeyup="filterTable()">
                </div>
                <div class="table-responsive" style="max-height: 600px;">
                    <table class="table table-hover align-middle text-center small" id="reportsTable">
                        <thead class="sticky-top">
                            <tr><th>تاریخ شمسی</th><th>میلادی</th><th>عملکرد شعب</th><th>ترانزیت / صادرات</th><th>عملیات</th></tr>
                        </thead>
                        <tbody>
                            {% for r in reports %}
                            <tr>
                                <td class="fw-bold text-primary fs-6">{{ r.shamsi_date }}</td>
                                <td class="text-muted">{{ r.date }}</td>
                                <td>
                                    <span class="badge bg-primary rounded-pill mb-1">{{ r.total_iss }} صادره</span>
                                    <div class="text-muted mt-1" style="font-size: 0.8rem;">
                                        {% for b in r.branches %}<b>{{ b.name }}</b>: {{ b.issued }} <code class="text-dark bg-light px-1 rounded">{{ b.range_text }}</code><br>{% endfor %}
                                    </div>
                                </td>
                                <td>
                                    <span class="badge bg-info text-dark">🚛 {{ r.transit }}</span><br>
                                    <span class="badge bg-success mt-1">📦 {{ r.export }}</span>
                                </td>
                                <td>
                                    <div class="btn-group btn-group-sm flex-wrap">
                                        <button type="button" class="btn btn-outline-warning text-dark fw-bold" onclick='loadIntoForm({{ r | tojson }})'>ویرایش</button>
                                        <a href="/download/jpg?date={{ r.date }}" class="btn btn-outline-primary">JPG</a>
                                        <a href="/download/pdf?date={{ r.date }}" class="btn btn-outline-danger">PDF</a>
                                        <a href="{{ r.wa_url }}" target="_blank" class="btn btn-wa"><i class="bi bi-whatsapp"></i></a>
                                        <form method="POST" action="/send_to_tg" class="d-inline"><input type="hidden" name="date" value="{{ r.date }}"><button type="submit" class="btn btn-info text-white" title="ارسال به تلگرام"><i class="bi bi-telegram"></i></button></form>
                                        <form method="POST" action="/delete_report" class="d-inline" onsubmit="return confirm('گزارش این تاریخ برای همیشه حذف شود؟')"><input type="hidden" name="date" value="{{ r.date }}"><button type="submit" class="btn btn-outline-secondary" title="حذف گزارش"><i class="bi bi-trash"></i></button></form>
                                    </div>
                                </td>
                            </tr>
                            {% endfor %}
                        </tbody>
                    </table>
                </div>
            </div>
        </div>

        <!-- تب ۴: روزهای جامانده -->
        <div class="tab-pane fade" id="tab-missing" role="tabpanel">
            <div class="glass-card p-4">
                <h5 class="fw-bold text-danger mb-3">روزهای بدون گزارش در ۳۰ روز گذشته</h5>
                {% if not missing_dates %}
                <div class="alert alert-success rounded-4 m-0">عالی! تمامی روزها ثبت شده‌اند.</div>
                {% else %}
                <div class="row g-2">
                    {% for m in missing_dates %}
                    <div class="col-md-4 col-lg-3">
                        <div class="p-3 border rounded-4 d-flex justify-content-between align-items-center bg-danger bg-opacity-10">
                            <div><div class="fw-bold text-danger">{{ m.shamsi }}</div><div class="small text-muted">{{ m.gregorian }}</div></div>
                            <button class="btn btn-sm btn-danger rounded-pill px-3 fw-bold" onclick="fillMissingDate('{{ m.shamsi }}')">➕ ثبت</button>
                        </div>
                    </div>
                    {% endfor %}
                </div>
                {% endif %}
            </div>
        </div>

        <!-- تب ۵: مدیریت کاربران و دیتابیس -->
        <div class="tab-pane fade" id="tab-admin" role="tabpanel">
            <div class="row g-4">
                <div class="col-lg-6">
                    <div class="glass-card p-4 mb-4">
                        <h5 class="fw-bold text-primary mb-3">مدیریت کاربران ربات تلگرام</h5>
                        <div class="table-responsive" style="max-height: 280px;">
                            <table class="table table-sm align-middle text-center small">
                                <thead><tr><th>شناسه</th><th>نام</th><th>وضعیت</th><th>عملیات</th></tr></thead>
                                <tbody>
                                    {% for uid, uname, fname, allowed, is_adm in tg_users %}
                                    <tr>
                                        <td><code>{{ uid }}</code></td>
                                        <td class="fw-bold">{{ fname or uname or 'کاربر' }}</td>
                                        <td>{% if is_adm %}👑 مدیر{% elif allowed %}✅ مجاز{% else %}🚫 مسدود{% endif %}</td>
                                        <td>
                                            {% if uid != admin_id %}
                                            <form method="POST" action="/toggle_user_web" class="d-inline"><input type="hidden" name="uid" value="{{ uid }}"><input type="hidden" name="allow" value="{{ 0 if allowed else 1 }}"><button type="submit" class="btn btn-sm btn-outline-primary py-0">تغییر وضعیت</button></form>
                                            {% endif %}
                                        </td>
                                    </tr>
                                    {% endfor %}
                                </tbody>
                            </table>
                        </div>
                    </div>
                    <div class="glass-card p-4">
                        <h5 class="fw-bold text-primary mb-3">ارسال اطلاعیه همگانی به تلگرام</h5>
                        <form method="POST" action="/broadcast_web">
                            <textarea name="message" class="form-control mb-2 small" rows="3" placeholder="متن پیام همگانی..." required></textarea>
                            <button type="submit" class="btn btn-primary w-100 rounded-3 fw-bold">ارسال به کاربران تلگرام</button>
                        </form>
                    </div>
                </div>
                <div class="col-lg-6">
                    <div class="glass-card p-4 mb-4">
                        <h5 class="fw-bold text-warning mb-3">بازیابی دیتابیس (Restore)</h5>
                        <form method="POST" action="/restore_db" enctype="multipart/form-data" class="d-flex gap-2">
                            <input type="file" name="db_file" accept=".db" class="form-control" required>
                            <button type="submit" class="btn btn-warning fw-bold px-4">آپلود</button>
                        </form>
                    </div>
                    <div class="glass-card p-4">
                        <h5 class="fw-bold text-secondary mb-3">آخرین فعالیت‌های سیستم</h5>
                        <div class="table-responsive" style="max-height: 280px;">
                            <table class="table table-sm small">
                                <thead><tr><th>زمان</th><th>کاربر</th><th>عملیات</th></tr></thead>
                                <tbody>
                                    {% for uid, act, ts in audit_logs %}
                                    <tr><td class="text-muted">{{ ts[:16] }}</td><td><code>{{ uid }}</code></td><td>{{ act }}</td></tr>
                                    {% endfor %}
                                </tbody>
                            </table>
                        </div>
                    </div>
                </div>
            </div>
        </div>
        <div class="tab-pane fade" id="tab-employees" role="tabpanel">
          <div class="glass-card p-3 mb-4">
            <div class="d-flex flex-wrap align-items-center justify-content-between gap-2 mb-2">
              <h5 class="fw-bold mb-0">وضعیت هوشمند گزارش امروز</h5>
              <span class="small text-muted">{{ branch_daily_status[0].date if branch_daily_status else today_shamsi }}</span>
            </div>
            <div class="row g-2">
              {% for bs in branch_daily_status %}
              <div class="col-md-6"><div class="border rounded-3 p-3 d-flex align-items-center justify-content-between gap-2">
                <div><div class="fw-bold">{{ bs.branch_name }}</div><div class="small text-muted">وضعیت گزارش روز جاری</div></div>
                <span class="badge text-bg-{{ bs.status_class }}">{{ bs.status }}</span>
              </div></div>
              {% endfor %}
            </div>
            <div class="small text-muted mt-2">گزارش‌های در انتظار تأیید هنوز جزو آمار نهایی محسوب نمی‌شوند. پس از تأیید، وضعیت شعبه به «ثبت نهایی شده» تغییر می‌کند.</div>
          </div>
          <div class="glass-card p-4 mb-4">
            <h5 class="fw-bold mb-3">🔐 نهایی‌سازی یک‌جای گزارش روزانه</h5>
            <p class="small text-muted">گزارش هر دو شعبه را یک‌جا بررسی و نهایی کن. ترانزیت و صادرات فقط توسط مدیر وارد می‌شود. روز بسته‌شده بدون بازگشایی مستند قابل تغییر نیست.</p>
            <form method="POST" action="/employees/finalize_day" class="row g-2 align-items-end">
              <div class="col-md-3"><label class="form-label">تاریخ شمسی</label><input class="form-control" name="report_date" value="{{ today_shamsi }}" required></div>
              <div class="col-md-2"><label class="form-label">ترانزیت (مدیر)</label><input class="form-control" name="transit" type="number" min="0" value="0" required></div>
              <div class="col-md-2"><label class="form-label">صادرات (مدیر)</label><input class="form-control" name="export" type="number" min="0" value="0" required></div>
              <div class="col-md-3"><label class="form-label">یادداشت نهایی‌سازی</label><input class="form-control" name="close_note" maxlength="500" placeholder="مثلاً بررسی و تطبیق شد" required></div>
              <div class="col-md-2"><button class="btn btn-success w-100" type="submit" onclick="return confirm('آیا از نهایی‌سازی و بستن این روز مطمئن هستید؟')">نهایی و بستن روز</button></div>
            </form>
            <hr><h6 class="fw-bold">روزهای بسته‌شده اخیر</h6>
            <div class="table-responsive"><table class="table table-sm align-middle"><thead><tr><th>تاریخ</th><th>وضعیت</th><th>ترانزیت</th><th>صادرات</th><th>توضیح/عملیات</th></tr></thead><tbody>
            {% for cw in closed_workdays %}<tr><td>{{ cw.shamsi_date }}</td><td>{{ 'بسته' if cw.status == 'closed' else 'بازگشایی‌شده' }}</td><td>{{ cw.transit }}</td><td>{{ cw.export }}</td><td><div class="small text-muted mb-1">{{ cw.close_note }}</div>{% if cw.status == 'closed' %}<form method="POST" action="/employees/reopen_day" class="d-flex gap-1"><input type="hidden" name="report_date" value="{{ cw.report_date }}"><input class="form-control form-control-sm" name="reason" placeholder="دلیل بازگشایی" required><button class="btn btn-sm btn-outline-warning" onclick="return confirm('بازگشایی روز ثبت و در لاگ ذخیره می‌شود. ادامه؟')">بازگشایی</button></form>{% endif %}</td></tr>{% else %}<tr><td colspan="5" class="text-muted text-center">هنوز روزی بسته نشده است.</td></tr>{% endfor %}
            </tbody></table></div>
          </div>
          <div class="row g-4">
            <div class="col-lg-4"><div class="glass-card p-4">
              <h5 class="fw-bold text-primary mb-3">ثبت شناسه کارمند تلگرام</h5>
              <p class="small text-muted">کارمند پس از ثبت شناسه عددی، از دکمه «ثبت آمار شعبه» در ربات استفاده می‌کند. ترانزیت و صادرات فقط دست مدیر می‌ماند.</p>
              <form method="POST" action="/employees/add" id="empAddForm">
                <label class="form-label">انتخاب سریع از کاربران ربات</label>
                <select class="form-select mb-3" id="empCandidate" onchange="fillEmpCandidate(this)">
                  <option value="">— انتخاب کنید (اختیاری) —</option>
                  {% for cd in employee_candidates %}<option value="{{ cd.user_id }}" data-name="{{ cd.first_name }}">{{ cd.first_name or 'بدون نام' }}{% if cd.username %} (@{{ cd.username }}){% endif %} — {{ cd.user_id }}{% if not cd.allowed %} ⏳{% endif %}</option>{% endfor %}
                </select>
                <label class="form-label">شناسه عددی تلگرام</label><input class="form-control mb-3" name="user_id" type="number" min="1" required>
                <label class="form-label">نام کارمند</label><input class="form-control mb-3" name="employee_name" maxlength="100" required>
                <label class="form-label">شعبه</label><select class="form-select mb-3" name="branch_name" required><option value="تیرپارک">تیرپارک</option><option value="پارسیان">پارسیان</option></select>
                <label class="form-label">سطح دسترسی</label><select class="form-select mb-3" name="access"><option value="submit">ثبت آمار شعبه</option><option value="view">فقط مشاهده (بدون ثبت آمار)</option></select>
                <label class="form-label">یادداشت (اختیاری)</label><input class="form-control mb-3" name="note" maxlength="200">
                <div class="form-check mb-3"><input class="form-check-input" type="checkbox" name="notify" value="1" id="empAddNotify" checked><label class="form-check-label small" for="empAddNotify">اطلاع‌رسانی به کارمند در تلگرام</label></div>
                <button class="btn btn-primary w-100 rounded-pill">ثبت / به‌روزرسانی کارمند</button>
              </form><div class="alert alert-info small mt-3 mb-0">کارمند فقط کافی است یک‌بار /start را در ربات بزند؛ سپس از فهرست بالا انتخاب می‌شود و نیازی به تایپ شناسه نیست (⏳ یعنی هنوز دسترسی ربات ندارد و با ثبت، فعال می‌شود).</div>
            </div>
            </div>
            <div class="col-lg-8"><div class="glass-card p-4"><h5 class="fw-bold text-warning mb-3">آمارهای منتظر تأیید ({{ pending_reports|length }})</h5><div class="table-responsive"><table class="table table-hover align-middle"><thead><tr><th>تاریخ</th><th>شعبه</th><th>کارمند</th><th>صادره</th><th>سریال</th><th>ابطالی</th><th>بازرسی سریال</th><th>عملیات</th></tr></thead><tbody>
            {% for p in pending_reports %}<tr><td>{{ p.shamsi_date }}</td><td>{{ p.branch_name }}</td><td>{{ p.employee_name or 'کارمند حذف‌شده' }}<br><small>{{ p.user_id }}</small></td><td>{{ p.issued }}</td><td>{{ p.serial_start }} تا {{ p.serial_end }}</td><td>{{ p.cancelled }}<br><small>{{ p.cancelled_serial }}</small></td><td class="small {% if '⚠️' in p.serial_warnings %}text-danger fw-bold{% else %}text-muted{% endif %}">{{ p.serial_warnings }}</td><td><form method="POST" action="/employees/review" class="d-flex gap-1"><input type="hidden" name="pending_id" value="{{ p.id }}"><button class="btn btn-sm btn-success" name="decision" value="approve">تأیید</button><button class="btn btn-sm btn-outline-danger" name="decision" value="reject">رد</button></form></td></tr>{% else %}<tr><td colspan="8" class="text-center text-muted p-4">موردی برای تأیید وجود ندارد.</td></tr>{% endfor %}
            </tbody></table></div><p class="small text-muted mb-0">تأیید هر ارسال، فقط اطلاعات همان شعبه را در گزارش تاریخ مربوطه ثبت می‌کند؛ ترانزیت و صادرات فقط از پنل مدیر ثبت می‌شود. پنل هر ۸ ثانیه گزارش جدید را بررسی می‌کند و اعلان زنده نشان می‌دهد.</p></div></div>

            <div class="col-12"><div class="glass-card p-4">
              <div class="d-flex flex-wrap gap-2 justify-content-between align-items-center mb-3">
                <h5 class="fw-bold mb-0">کارکنان ثبت‌شده <span class="badge bg-secondary rounded-pill">{{ employees|length }}</span></h5>
                <div class="d-flex gap-2"><input id="empSearch" class="form-control form-control-sm" placeholder="جستجو: نام، شناسه، شعبه…" oninput="filterEmployees()" style="min-width:220px">
                <select id="empFilter" class="form-select form-select-sm" onchange="filterEmployees()"><option value="">همه</option><option value="active">فعال</option><option value="inactive">غیرفعال</option><option value="view">فقط مشاهده</option></select></div>
              </div>
              <div class="table-responsive"><table class="table table-sm align-middle" id="empTable"><thead><tr><th>کارمند</th><th>تلگرام</th><th>شعبه</th><th>دسترسی</th><th>وضعیت</th><th>عملکرد</th><th>عملیات</th></tr></thead><tbody>
              {% for e in employees %}<tr data-search="{{ e.employee_name }} {{ e.user_id }} {{ e.branch_name }} {{ e.username }}" data-state="{{ 'active' if e.active else 'inactive' }}" data-view="{{ 0 if e.can_submit else 1 }}">
                <td class="fw-bold">{{ e.employee_name }}{% if e.note %}<br><small class="text-muted fw-normal">{{ e.note }}</small>{% endif %}</td>
                <td><code>{{ e.user_id }}</code>{% if e.username %}<br><small>@{{ e.username }}</small>{% endif %}{% if not e.started %}<br><small class="text-warning">ربات را استارت نکرده</small>{% endif %}</td>
                <td>{{ e.branch_name }}</td>
                <td>{% if e.can_submit %}<span class="badge bg-primary">ثبت آمار</span>{% else %}<span class="badge bg-info text-dark">فقط مشاهده</span>{% endif %}</td>
                <td>{% if e.active %}<span class="badge bg-success">فعال</span>{% else %}<span class="badge bg-secondary">غیرفعال</span>{% endif %}</td>
                <td class="small">{{ e.total }} ارسال: <span class="text-success">{{ e.approved }} تأیید</span> · <span class="text-danger">{{ e.rejected }} رد</span>{% if e.pending %} · <span class="text-warning">{{ e.pending }} منتظر</span>{% endif %}{% if e.last_date %}<br><small class="text-muted">آخرین: {{ e.last_date }}</small>{% endif %}</td>
                <td><div class="d-flex gap-1 flex-wrap">
                  <button type="button" class="btn btn-sm btn-outline-primary" data-uid="{{ e.user_id }}" data-name="{{ e.employee_name }}" data-branch="{{ e.branch_name }}" data-active="{{ e.active }}" data-submit="{{ e.can_submit }}" data-note="{{ e.note }}" onclick="openEmpEdit(this)">ویرایش</button>
                  <form method="POST" action="/employees/toggle"><input type="hidden" name="user_id" value="{{ e.user_id }}"><button class="btn btn-sm btn-outline-secondary">{{ 'غیرفعال‌سازی' if e.active else 'فعال‌سازی' }}</button></form>
                  <form method="POST" action="/employees/delete" onsubmit="return confirm('کارمند {{ e.employee_name }} حذف شود؟{% if e.pending %} (این کارمند {{ e.pending }} گزارش منتظر تأیید دارد){% endif %}')"><input type="hidden" name="user_id" value="{{ e.user_id }}"><input type="hidden" name="notify" value="1"><button class="btn btn-sm btn-outline-danger">حذف</button></form>
                </div></td></tr>{% else %}<tr><td colspan="7" class="text-center text-muted p-4">کارمندی ثبت نشده است.</td></tr>{% endfor %}
              </tbody></table></div></div></div>
            <div class="modal fade" id="empEditModal" tabindex="-1"><div class="modal-dialog modal-dialog-centered"><div class="modal-content">
              <form method="POST" action="/employees/update">
                <div class="modal-header"><h5 class="modal-title">ویرایش کارمند <code id="empEditId"></code></h5><button type="button" class="btn-close" data-bs-dismiss="modal"></button></div>
                <div class="modal-body">
                  <input type="hidden" name="user_id" id="empEditUid">
                  <label class="form-label">نام کارمند</label><input class="form-control mb-3" name="employee_name" id="empEditName" maxlength="100" required>
                  <label class="form-label">شعبه</label><select class="form-select mb-3" name="branch_name" id="empEditBranch"><option value="تیرپارک">تیرپارک</option><option value="پارسیان">پارسیان</option></select>
                  <label class="form-label">سطح دسترسی</label><select class="form-select mb-3" name="access" id="empEditAccess"><option value="submit">ثبت آمار شعبه</option><option value="view">فقط مشاهده</option></select>
                  <label class="form-label">وضعیت</label><select class="form-select mb-3" name="active" id="empEditActive"><option value="1">فعال</option><option value="0">غیرفعال</option></select>
                  <label class="form-label">یادداشت</label><input class="form-control mb-3" name="note" id="empEditNote" maxlength="200">
                  <div class="form-check"><input class="form-check-input" type="checkbox" name="notify" value="1" id="empEditNotify" checked><label class="form-check-label small" for="empEditNotify">اطلاع‌رسانی تغییرات به کارمند در تلگرام</label></div>
                </div>
                <div class="modal-footer"><button class="btn btn-primary rounded-pill px-4">ذخیره تغییرات</button></div>
              </form></div></div></div>
            <script>
            function fillEmpCandidate(sel){var o=sel.options[sel.selectedIndex];if(!o.value)return;var f=document.getElementById('empAddForm');f.user_id.value=o.value;if(!f.employee_name.value)f.employee_name.value=o.getAttribute('data-name')||'';}
            function openEmpEdit(b){var m=document.getElementById('empEditModal');if(m.parentNode!==document.body)document.body.appendChild(m);
              document.getElementById('empEditUid').value=b.dataset.uid;document.getElementById('empEditId').textContent=b.dataset.uid;
              document.getElementById('empEditName').value=b.dataset.name;document.getElementById('empEditBranch').value=b.dataset.branch;
              document.getElementById('empEditAccess').value=b.dataset.submit==='1'?'submit':'view';document.getElementById('empEditActive').value=b.dataset.active==='1'?'1':'0';
              document.getElementById('empEditNote').value=b.dataset.note||'';bootstrap.Modal.getOrCreateInstance(m).show();}
            function filterEmployees(){var q=(document.getElementById('empSearch').value||'').toLowerCase(),f=document.getElementById('empFilter').value;
              document.querySelectorAll('#empTable tbody tr[data-search]').forEach(function(r){var ok=r.dataset.search.toLowerCase().indexOf(q)>-1;
                if(f==='active')ok=ok&&r.dataset.state==='active';else if(f==='inactive')ok=ok&&r.dataset.state==='inactive';else if(f==='view')ok=ok&&r.dataset.view==='1';
                r.style.display=ok?'':'none';});}
            </script>
          </div>
        </div>
    </div>
    {% endif %}
</div>

<footer class="text-center py-4 mt-auto opacity-75 small border-top">
    <div class="container d-flex justify-content-between align-items-center">
        <span class="fw-bold">{{ footer }}</span>
        <span class="text-muted">مجهز به جدول حسابداری ماهانه دینار/دلار + ۴ بازه سریال</span>
    </div>
</footer>

<script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/js/bootstrap.bundle.min.js"></script>
<script>
function applyTheme(theme) {
    document.documentElement.setAttribute('data-bs-theme', theme);
    try { localStorage.setItem('bashmaq_theme', theme); } catch (e) {}
    const btn = document.querySelector('[aria-label="تغییر حالت روشن و تاریک"]');
    if (btn) btn.innerHTML = theme === 'dark' ? '<i class="bi bi-sun-fill"></i> حالت روشن' : '<i class="bi bi-moon-stars-fill"></i> حالت تاریک';
}
function toggleTheme() {
    const cur = document.documentElement.getAttribute('data-bs-theme') || 'light';
    applyTheme(cur === 'dark' ? 'light' : 'dark');
}
let savedTheme = 'light';
try { savedTheme = localStorage.getItem('bashmaq_theme') || 'light'; } catch (e) {}
applyTheme(savedTheme);
function updateLiveClock() {
    const el = document.getElementById('liveClock');
    if (!el) return;
    try { el.textContent = new Intl.DateTimeFormat('fa-IR', { dateStyle: 'short', timeStyle: 'short' }).format(new Date()); }
    catch (e) { el.textContent = 'سامانه آماده'; }
}
updateLiveClock(); setInterval(updateLiveClock, 30000);
if (window.jalaliDatepicker) jalaliDatepicker.startWatch({ separatorChar: '/', time: false });
</script>

{% if logged_in %}
<script>
const KNOWN_BRANCHES = {{ known_branches | safe }};
let lastSuggestions = {};
let currentGregorianDate = "{{ today_gregorian }}";
let currentShamsiDate = "{{ today_shamsi }}";

let formData = {
    branches: [
        { name: 'تیرپارک', issued: 0, cancelled: 0, cancelled_serial: '0', ranges: [{start: '', end: ''}] },
        { name: 'پارسیان', issued: 0, cancelled: 0, cancelled_serial: '0', ranges: [{start: '', end: ''}] }
    ]
};

let accRows = [];

async function loadAccountingSheet(syncFromReports) {
    const mp = document.getElementById('acc_month_picker').value || new Date().toISOString().slice(0,7);
    const resp = await fetch(`/api/accounting_sheet?month=${encodeURIComponent(mp)}&sync=${syncFromReports ? '1' : '0'}`);
    const data = await resp.json();
    if (!data.ok) return;
    document.getElementById('acc_dinar_rate').value = data.dinar_rate || 15000;
    accRows = data.rows || [];
    renderAccountingTable();
}

function renderAccountingTable() {
    let html = '';
    accRows.forEach((r, idx) => {
        html += `
            <tr>
                <td class="fw-bold">${escapeHtml(r.s_date)}</td>
                <td class="text-muted">${escapeHtml(r.g_date)}</td>
                <td><input type="number" class="acc-inp" value="${escapeHtml(r.amar)}" oninput="updateAccCell(${idx}, 'amar', this.value)"></td>
                <td class="fw-bold" id="acc_dinar_${idx}">0</td>
                <td><input type="number" class="acc-inp" value="${escapeHtml(r.iraq_hawala)}" oninput="updateAccCell(${idx}, 'iraq_hawala', this.value)"></td>
                <td><input type="number" class="acc-inp" value="${escapeHtml(r.torshizi)}" oninput="updateAccCell(${idx}, 'torshizi', this.value)"></td>
                <td><input type="number" class="acc-inp" value="${escapeHtml(r.naser)}" oninput="updateAccCell(${idx}, 'naser', this.value)"></td>
                <td class="fw-bold" id="acc_usd_${idx}">$0</td>
                <td class="acc-col-grey"><input type="number" class="acc-inp" value="${escapeHtml(r.ex_rate)}" oninput="updateAccCell(${idx}, 'ex_rate', this.value)"></td>
                <td class="acc-col-pink-dark"><input type="number" class="acc-inp" value="${escapeHtml(r.iran_maqsad)}" oninput="updateAccCell(${idx}, 'iran_maqsad', this.value)"></td>
                <td class="acc-col-pink-light"><input type="number" class="acc-inp" value="${escapeHtml(r.sabokbar)}" oninput="updateAccCell(${idx}, 'sabokbar', this.value)"></td>
                <td class="acc-col-pink-light"><input type="number" class="acc-inp" value="${escapeHtml(r.motafareqeh)}" oninput="updateAccCell(${idx}, 'motafareqeh', this.value)"></td>
                <td><input type="text" class="acc-inp" value="${escapeHtml(r.desc || '')}" oninput="updateAccCell(${idx}, 'desc', this.value)"></td>
            </tr>
        `;
    });
    document.getElementById('acc_tbody').innerHTML = html;
    recalcAccountingTable();
}

function updateAccCell(idx, field, val) {
    accRows[idx][field] = val;
    recalcAccountingTable();
}

function fmtNum(n) {
    return Math.round(n).toLocaleString('en-US');
}

function recalcAccountingTable() {
    const rate = parseFloat(document.getElementById('acc_dinar_rate').value) || 15000;
    let sumAmar = 0, sumDinar = 0, sumHawala = 0, sumTor = 0, sumNaser = 0, sumUsd = 0, sumIm = 0, sumSb = 0, sumMt = 0;

    accRows.forEach((r, idx) => {
        const amar = parseFloat(r.amar) || 0;
        const dinar = amar * rate;
        const hawala = parseFloat(r.iraq_hawala) || 0;
        const tor = parseFloat(r.torshizi) || 0;
        const naser = parseFloat(r.naser) || 0;
        const ex = parseFloat(r.ex_rate) || 0;
        const im = parseFloat(r.iran_maqsad) || 0;
        const sb = parseFloat(r.sabokbar) || 0;
        const mt = parseFloat(r.motafareqeh) || 0;

        const rowUsd = (ex > 0 ? (dinar / ex) : 0) + im + sb + mt;

        sumAmar += amar;
        sumDinar += dinar;
        sumHawala += hawala;
        sumTor += tor;
        sumNaser += naser;
        sumUsd += rowUsd;
        sumIm += im;
        sumSb += sb;
        sumMt += mt;

        const dEl = document.getElementById(`acc_dinar_${idx}`);
        if (dEl) dEl.innerText = fmtNum(dinar);
        const uEl = document.getElementById(`acc_usd_${idx}`);
        if (uEl) uEl.innerText = '$' + fmtNum(rowUsd);
    });

    document.getElementById('acc_sum_amar').innerText = fmtNum(sumAmar);
    document.getElementById('acc_sum_dinar').innerText = fmtNum(sumDinar);
    document.getElementById('acc_sum_hawala').innerText = fmtNum(sumHawala);
    document.getElementById('acc_sum_torshizi').innerText = fmtNum(sumTor);
    document.getElementById('acc_sum_naser').innerText = fmtNum(sumNaser);
    document.getElementById('acc_sum_usd').innerText = '$' + fmtNum(sumUsd);
    document.getElementById('acc_sum_im').innerText = fmtNum(sumIm);
    document.getElementById('acc_sum_sb').innerText = fmtNum(sumSb);
    document.getElementById('acc_sum_mt').innerText = fmtNum(sumMt);

    const totalDeduct = sumHawala + sumTor + sumNaser;
    document.getElementById('box_total_dinar').innerText = fmtNum(sumDinar);
    document.getElementById('box_total_deduct').innerText = fmtNum(totalDeduct);
    document.getElementById('box_final_dinar').innerText = fmtNum(sumDinar - totalDeduct);

    document.getElementById('box_amar_2').innerText = fmtNum(sumAmar);
    document.getElementById('box_sb_usd').innerText = sumSb > 0 ? ('$ ' + fmtNum(sumSb)) : '$ -';
    document.getElementById('box_im_usd').innerText = sumIm > 0 ? ('$ ' + fmtNum(sumIm)) : '$ -';
    document.getElementById('box_mt_usd').innerText = sumMt > 0 ? ('$ ' + fmtNum(sumMt)) : '$ -';
    // مجموع دلار باید تبدیل دینار به دلار (دینار / نرخ تبدیل) را هم شامل شود؛ sumUsd از rowUsdها جمع شده است.
    const grandUsd = sumUsd;
    document.getElementById('live_total_dinar').innerText = fmtNum(sumDinar);
    document.getElementById('live_deductions').innerText = fmtNum(totalDeduct);
    document.getElementById('live_net').innerText = fmtNum(sumDinar-totalDeduct);
    document.getElementById('live_usd').innerText = '$' + fmtNum(grandUsd);
    document.getElementById('box_grand_usd').innerText = grandUsd > 0 ? ('$ ' + fmtNum(grandUsd)) : '$ -';
}

function downloadAccountingPDF() { const mp=document.getElementById('acc_month_picker').value || new Date().toISOString().slice(0,7); window.location.href=`/export_accounting_pdf?month=${encodeURIComponent(mp)}`; }
function shareAccountingWhatsApp() {
 const mp=document.getElementById('acc_month_picker').value || new Date().toISOString().slice(0,7);
 const msg=`📊 خلاصه صورت‌حساب مالی ماه ${mp}\n💰 درآمد دیناری: ${document.getElementById('live_total_dinar').innerText}\n➖ کسورات: ${document.getElementById('live_deductions').innerText}\n✅ مانده خالص: ${document.getElementById('live_net').innerText} دینار\n💵 درآمد دلاری: ${document.getElementById('live_usd').innerText}`;
 window.open('https://wa.me/?text='+encodeURIComponent(msg),'_blank');
}
async function auditSerials() {
 const box=document.getElementById('serial_audit_result');
 box.textContent='در حال بررسی...';
 try {
   const resp=await fetch('/api/audit_serials', {credentials:'same-origin', cache:'no-store'});
   if (!resp.ok) throw new Error('درخواست بررسی ناموفق بود');
   const data=await resp.json();
   box.replaceChildren();
   const panel=document.createElement('div');
   panel.className='alert '+(data.issues && data.issues.length ? 'alert-warning' : 'alert-success');
   if (data.issues && data.issues.length) {
     const title=document.createElement('b'); title.textContent='موارد نیازمند بررسی:'; panel.appendChild(title);
     const list=document.createElement('ul');
     data.issues.forEach(item=>{const li=document.createElement('li'); li.textContent=String(item); list.appendChild(li);});
     panel.appendChild(list);
   } else {
     panel.textContent='تداخل یا پرش مشکوک بین بازه‌های سریال شناسایی نشد.';
   }
   box.appendChild(panel);
 } catch (err) { box.textContent='بررسی انجام نشد؛ صفحه را تازه‌سازی کرده و دوباره تلاش کنید.'; }
}

async function saveAccountingSheet() {
    const mp = document.getElementById('acc_month_picker').value || new Date().toISOString().slice(0,7);
    const rate = parseFloat(document.getElementById('acc_dinar_rate').value) || 15000;
    const resp = await fetch('/api/save_accounting_sheet', {
        method: 'POST',
        headers: {'Content-Type': 'application/json', 'X-CSRFToken': document.querySelector('meta[name=\"csrf-token\"]')?.content || ''},
        body: JSON.stringify({ month_key: mp, dinar_rate: rate, rows: accRows })
    });
    const res = await resp.json();
    if (res.ok) alert('✅ جدول مالی ماه ' + mp + ' با موفقیت در دیتابیس ذخیره شد!');
}

function downloadAccountingExcel() {
    const mp = document.getElementById('acc_month_picker').value || new Date().toISOString().slice(0,7);
    window.location.href = `/export_accounting_excel?month=${encodeURIComponent(mp)}`;
}

function renderForm() {
    let html = '';
    formData.branches.forEach((b, bIdx) => {
        let quickBtns = KNOWN_BRANCHES.map(bn =>
            `<button type="button" class="btn btn-outline-secondary btn-sm py-0 px-2" style="font-size:0.72rem;" onclick="selectBranchName(${bIdx}, '${bn}')">${bn}</button>`
        ).join('');

        let rangesHtml = '';
        b.ranges.forEach((r, rIdx) => {
            rangesHtml += `
                <div class="range-row-box row g-2 align-items-center">
                    <div class="col-5">
                        <label class="small text-primary fw-bold">شروع بازه ${rIdx+1}</label>
                        <input type="text" class="form-control form-control-sm fw-bold text-center" value="${escapeHtml(r.start)}" oninput="updateRange(${bIdx}, ${rIdx}, 'start', this.value)">
                    </div>
                    <div class="col-5">
                        <label class="small text-primary fw-bold">پایان بازه ${rIdx+1}</label>
                        <input type="text" class="form-control form-control-sm fw-bold text-center" value="${escapeHtml(r.end)}" oninput="updateRange(${bIdx}, ${rIdx}, 'end', this.value)">
                    </div>
                    <div class="col-2 text-center mt-4">
                        ${rIdx > 0 ? `<button type="button" class="btn btn-sm btn-outline-danger py-0 px-2" onclick="removeRange(${bIdx}, ${rIdx})"><i class="bi bi-trash"></i></button>` : ''}
                    </div>
                </div>
            `;
        });

        html += `
            <div class="branch-box">
                <div class="d-flex justify-content-between align-items-center mb-2">
                    <span class="fw-bold text-primary"><i class="bi bi-building"></i> شعبه ${bIdx+1}</span>
                    <div class="d-flex gap-1 align-items-center">
                        <div class="btn-group">${quickBtns}</div>
                        ${formData.branches.length > 1 ? `<button type="button" class="btn-close btn-sm ms-1" onclick="removeBranch(${bIdx})"></button>` : ''}
                    </div>
                </div>
                <div class="row g-2 mb-2">
                    <div class="col-6">
                        <label class="small text-muted fw-bold">نام شعبه</label>
                        <input type="text" class="form-control form-control-sm fw-bold" value="${escapeHtml(b.name)}" oninput="updateBranch(${bIdx}, 'name', this.value)" required>
                    </div>
                    <div class="col-6">
                        <label class="small text-muted fw-bold">کل صادره این شعبه</label>
                        <input type="number" class="form-control form-control-sm fw-bold text-success" value="${b.issued || ''}" oninput="updateBranch(${bIdx}, 'issued', this.value)" required min="0">
                    </div>
                </div>
                <div class="p-2 border rounded-4 mb-2">
                    <div class="d-flex justify-content-between align-items-center mb-2">
                        <span class="small fw-bold text-secondary">لیست بازه سریال‌ها (تا ۴ بازه)</span>
                        <button type="button" class="btn btn-sm btn-primary rounded-pill py-0 px-2" onclick="addRange(${bIdx})">+ افزودن بازه جدید</button>
                    </div>
                    ${rangesHtml}
                    <div class="small fw-bold mt-1 text-center" id="err_${bIdx}"></div>
                </div>
                <div class="row g-2">
                    <div class="col-5">
                        <label class="small text-muted">تعداد ابطالی</label>
                        <input type="number" class="form-control form-control-sm" value="${b.cancelled}" oninput="updateBranch(${bIdx}, 'cancelled', this.value)" min="0">
                    </div>
                    <div class="col-7">
                        <label class="small text-muted">سریال ابطالی</label>
                        <input type="text" class="form-control form-control-sm" value="${escapeHtml(b.cancelled_serial)}" oninput="updateBranch(${bIdx}, 'cancelled_serial', this.value)">
                    </div>
                </div>
            </div>
        `;
    });
    document.getElementById('smart_branches_container').innerHTML = html;
    validateTotals();
}

function selectBranchName(bIdx, name) {
    formData.branches[bIdx].name = name;
    if (lastSuggestions[name]) {
        formData.branches[bIdx].ranges[0].start = lastSuggestions[name].next_start;
        let iss = parseInt(formData.branches[bIdx].issued) || 0;
        if (iss > 0 && formData.branches[bIdx].ranges.length === 1) {
            formData.branches[bIdx].ranges[0].end = (parseInt(lastSuggestions[name].next_start) + iss - 1).toString();
        }
    }
    renderForm();
}

function updateBranch(bIdx, field, val) {
    formData.branches[bIdx][field] = val;
    if (field === 'issued' && formData.branches[bIdx].ranges.length === 1) {
        let iss = parseInt(val) || 0;
        let s = parseInt(formData.branches[bIdx].ranges[0].start) || 0;
        if (iss === 0) formData.branches[bIdx].ranges[0].end = '0';
        else if (iss > 0 && s > 0) formData.branches[bIdx].ranges[0].end = (s + iss - 1).toString();
        renderForm();
        return;
    }
    validateTotals();
}

function updateRange(bIdx, rIdx, field, val) {
    formData.branches[bIdx].ranges[rIdx][field] = val;
    if (field === 'start' && formData.branches[bIdx].ranges.length === 1) {
        let issued = parseInt(formData.branches[bIdx].issued) || 0;
        let s = parseInt(val) || 0;
        if (issued > 0 && s > 0) {
            formData.branches[bIdx].ranges[rIdx].end = (s + issued - 1).toString();
            renderForm();
            return;
        }
    }
    validateTotals();
}

function addRange(bIdx) {
    if (formData.branches[bIdx].ranges.length >= 4) { alert('حداکثر ۴ بازه مجاز است!'); return; }
    formData.branches[bIdx].ranges.push({start: '', end: ''});
    renderForm();
}
function removeRange(bIdx, rIdx) { formData.branches[bIdx].ranges.splice(rIdx, 1); renderForm(); }
function addBranch() { formData.branches.push({ name: '', issued: 0, cancelled: 0, cancelled_serial: '0', ranges: [{start: '', end: ''}] }); renderForm(); }
function removeBranch(bIdx) { formData.branches.splice(bIdx, 1); renderForm(); }

function applyMagicFill() {
    const raw = document.getElementById('magic_inp').value.trim();
    const nums = raw.replace(/-/g, ' ').split(/\\s+/).map(x => parseInt(x)).filter(x => !isNaN(x));
    if (nums.length !== 4) { alert('لطفاً ۴ عدد با فاصله وارد کنید'); return; }
    const [issTir, issPar, trCnt, exCnt] = nums;
    const sTir = (lastSuggestions['تیرپارک'] ? lastSuggestions['تیرپارک'].next_start : 1);
    const eTir = issTir > 0 ? (sTir + issTir - 1) : 0;
    const sPar = (lastSuggestions['پارسیان'] ? lastSuggestions['پارسیان'].next_start : 1);
    const ePar = issPar > 0 ? (sPar + issPar - 1) : 0;

    formData.branches = [
        { name: 'تیرپارک', issued: issTir, cancelled: 0, cancelled_serial: '0', ranges: [{start: issTir > 0 ? sTir : '0', end: issTir > 0 ? eTir : '0'}] },
        { name: 'پارسیان', issued: issPar, cancelled: 0, cancelled_serial: '0', ranges: [{start: issPar > 0 ? sPar : '0', end: issPar > 0 ? ePar : '0'}] }
    ];
    document.getElementById('inp_transit').value = trCnt;
    document.getElementById('inp_export').value = exCnt;
    renderForm();
}

function validateTotals() {
    let tIss = 0, tCanc = 0;
    formData.branches.forEach((b, idx) => {
        tIss += parseInt(b.issued) || 0;
        tCanc += parseInt(b.cancelled) || 0;
        let rSum = 0;
        b.ranges.forEach(r => {
            let s = parseInt(r.start) || 0;
            let e = parseInt(r.end) || 0;
            if (e >= s && s > 0) rSum += (e - s + 1);
        });
        let errEl = document.getElementById('err_' + idx);
        if (errEl) {
            let iss = parseInt(b.issued) || 0;
            if (iss > 0 && rSum !== iss) {
                errEl.innerHTML = `⚠️ جمع بازه‌ها (${rSum}) با کل صادره (${iss}) برابر نیست!`;
                errEl.className = 'small fw-bold mt-2 text-danger';
            } else if (iss > 0 && rSum === iss) {
                errEl.innerHTML = `✅ مجموع بازه‌ها منطبق است (${rSum} فقره)`;
                errEl.className = 'small fw-bold mt-2 text-success';
            } else errEl.innerHTML = '';
        }
    });
    document.getElementById('liveTotalIssued').innerText = tIss;
    document.getElementById('liveTotalCanc').innerText = tCanc;
    updateLivePreviewText();
}

function updateLivePreviewText() {
    const cnt = formData.branches.length;
    const titleMap = {1: 'یک شعبه', 2: 'دو شعبه', 3: 'سه شعبه', 4: 'چهار شعبه'};
    const headerB = titleMap[cnt] || (cnt + ' شعبه');
    let txt = `باسلام\\n\\n(${headerB} مرز باشماق)\\n\\nدر روز ${currentGregorianDate}\\nتاریخ شمسی: ${currentShamsiDate}\\n\\n`;
    let totalIss = 0;
    formData.branches.forEach((b, i) => {
        totalIss += parseInt(b.issued) || 0;
        let rStr = b.ranges.filter(r => r.start && r.end).map(r => `${r.start} الی ${r.end}`).join(' و ') || '0';
        txt += `شعبه ${i+1} (${b.name || '-'}) مرز باشماق\\n\\nتعداد صادره ${b.issued || 0} فقره مانیفست از سریال ${rStr} درمرز باشماق صادر گردید\\nتعداد ابطالی ${b.cancelled || 0} فقره\\nسریال ابطالی: ${b.cancelled_serial || '0'}\\n\\n`;
    });
    txt += `جمع کل صادرها : ${totalIss} فقره\\n\\n────────────────\\n\\n`;
    const tr = document.getElementById('inp_transit')?.value || 0;
    const ex = document.getElementById('inp_export')?.value || 0;
    txt += `شعبه مرز باشماق\\n\\n${currentGregorianDate}\\n${currentShamsiDate}\\n\\nتعداد صادره\\nترانزیت : ${tr}\\nصادرات : ${ex}`;
    const box = document.getElementById('livePreviewBox');
    if (box) box.value = txt;
    const waLink = document.getElementById('liveWaLink');
    if (waLink) waLink.href = 'https://wa.me/?text=' + encodeURIComponent(txt);
}

function copyLivePreview() {
    const box = document.getElementById('livePreviewBox');
    box.select();
    navigator.clipboard.writeText(box.value);
    alert('✅ متن کپی شد!');
}

document.getElementById('reportForm').onsubmit = function() {
    document.getElementById('json_payload').value = JSON.stringify({
        date_input: document.getElementById('report_date').value,
        transit: document.getElementById('inp_transit').value,
        export: document.getElementById('inp_export').value,
        send_tg: document.getElementById('send_tg').checked,
        branches: formData.branches
    });
};

function setQuickDate(dt) { document.getElementById('report_date').value = dt; onDateChange(); }

async function onDateChange() {
    const dt = document.getElementById('report_date').value.trim();
    if (!dt) return;
    try {
        const resp = await fetch('/api/date_info?date=' + encodeURIComponent(dt));
        const info = await resp.json();
        if (!info.ok) return;
        currentGregorianDate = info.gregorian;
        currentShamsiDate = info.shamsi;
        lastSuggestions = info.suggestions || {};

        const badge = document.getElementById('formModeBadge');
        if (info.existing) {
            badge.innerText = '✏️ ویرایش (' + info.shamsi + ')';
            badge.className = 'badge bg-warning text-dark px-3 py-2 rounded-pill';
            document.getElementById('inp_transit').value = info.existing.transit || 0;
            document.getElementById('inp_export').value = info.existing.export || 0;
            formData.branches = info.existing.branches.map(b => ({
                name: b.name, issued: b.issued, cancelled: b.cancelled, cancelled_serial: b.cancelled_serial || '0',
                ranges: b.ranges && b.ranges.length > 0 ? b.ranges : [{start: b.start, end: b.end}]
            }));
        } else {
            badge.innerText = '✨ ثبت جدید (' + info.shamsi + ')';
            badge.className = 'badge bg-success bg-opacity-10 text-success px-3 py-2 rounded-pill';
            formData.branches.forEach(b => {
                if (lastSuggestions[b.name]) b.ranges = [{start: lastSuggestions[b.name].next_start, end: ''}];
            });
        }
        renderForm();
    } catch (e) {}
}

function loadIntoForm(r) {
    document.getElementById('report_date').value = r.shamsi_date;
    currentGregorianDate = r.date;
    currentShamsiDate = r.shamsi_date;
    document.getElementById('inp_transit').value = r.transit || 0;
    document.getElementById('inp_export').value = r.export || 0;
    formData.branches = r.branches.map(b => ({
        name: b.name, issued: b.issued, cancelled: b.cancelled, cancelled_serial: b.cancelled_serial || '0',
        ranges: b.ranges && b.ranges.length > 0 ? b.ranges : [{start: b.start, end: b.end}]
    }));
    renderForm();
    new bootstrap.Tab(document.getElementById('tab-dash-btn')).show();
    window.scrollTo({ top: 0, behavior: 'smooth' });
}

function fillMissingDate(shamsiDate) {
    document.getElementById('report_date').value = shamsiDate;
    onDateChange();
    new bootstrap.Tab(document.getElementById('tab-dash-btn')).show();
}

function filterTable() {
    const q = document.getElementById('tableSearch').value.toLowerCase();
    document.querySelectorAll('#reportsTable tbody tr').forEach(r => {
        r.style.display = r.innerText.toLowerCase().includes(q) ? '' : 'none';
    });
}

async function searchSerial() {
    const s = document.getElementById('serial_search_inp').value;
    const resBox = document.getElementById('serial_result');
    if (!s) return;
    resBox.innerHTML = '⏳ جستجو...';
    const resp = await fetch('/api/search_serial?serial=' + encodeURIComponent(s));
    const data = await resp.json();
    if (data.results.length === 0) resBox.innerHTML = `<div class="alert alert-warning py-2 m-0">❌ یافت نشد.</div>`;
    else resBox.innerHTML = data.results.map(x => `<div class="alert alert-info py-2 mb-2">📅 <b>${escapeHtml(x.shamsi)}</b> | شعبه: <b>${escapeHtml(x.branch)}</b> | <code>${escapeHtml(x.range)}</code> | <b>${escapeHtml(x.status)}</b></div>`).join('');
}

async function runGlobalSearch() {
  const q = (document.getElementById('globalSearchInput')?.value || '').trim();
  const box = document.getElementById('globalSearchResults'); if (!box) return;
  if (q.length < 2) { box.innerHTML=''; return; }
  box.innerHTML='<div class="small text-muted p-2">در حال جستجو…</div>';
  try {
    const r=await fetch('/api/global_search?q='+encodeURIComponent(q)); const d=await r.json();
    box.innerHTML = !d.results.length ? '<div class="alert alert-warning mt-2 mb-0">نتیجه‌ای پیدا نشد.</div>' : '<div class="list-group mt-2">'+d.results.map(x=>`<a class="list-group-item list-group-item-action" href="#" data-search="${escapeHtml(q)}" onclick="document.getElementById('tableSearch').value=this.dataset.search;filterTable();new bootstrap.Tab(document.getElementById('tab-archive-btn')).show();return false"><b>${escapeHtml(x.shamsi)}</b> — ${escapeHtml(x.branch)} <span class="text-muted">${escapeHtml(x.detail)}</span></a>`).join('')+'</div>';
  } catch(e) { box.innerHTML='<div class="alert alert-danger mt-2">جستجو انجام نشد.</div>'; }
}
function escapeHtml(v){return String(v??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));}
let globalSearchTimer; const _runGlobalSearch=runGlobalSearch; runGlobalSearch=function(){clearTimeout(globalSearchTimer);globalSearchTimer=setTimeout(_runGlobalSearch,250)};

window.onload = function() {
    renderForm();
    onDateChange();
    loadAccountingSheet(false);
    new Chart(document.getElementById('lineChart').getContext('2d'), {
        type: 'line',
        data: {
            labels: {{ chart_labels | safe }},
            datasets: [
                { label: 'کل صادره', data: {{ chart_issued | safe }}, borderColor: '#4f46e5', backgroundColor: 'rgba(79,70,229,0.12)', tension: 0.35, fill: true, borderWidth: 3 },
                { label: 'صادرات', data: {{ chart_export | safe }}, borderColor: '#10b981', tension: 0.35, borderWidth: 2 },
                { label: 'ترانزیت', data: {{ chart_transit | safe }}, borderColor: '#f97316', tension: 0.35, borderWidth: 2 }
            ]
        },
        options: { responsive: true, plugins: { legend: { labels: { font: { family: 'Vazirmatn' } } } } }
    });
    new Chart(document.getElementById('pieChart').getContext('2d'), {
        type: 'doughnut',
        data: {
            labels: {{ pie_labels | safe }},
            datasets: [{ data: {{ pie_values | safe }}, backgroundColor: ['#4f46e5', '#10b981', '#f59e0b', '#8b5cf6', '#ec4899'] }]
        },
        options: { responsive: true, plugins: { legend: { position: 'bottom', labels: { font: { family: 'Vazirmatn' } } } } }
    });
};
</script>
<script>
// اعلان زنده پنل مدیر برای گزارش‌های جدید کارمندان
let lastEmployeePendingId = Number({{ latest_pending_id | default(0) }});
const initialEmployeePendingCount = Number({{ pending_reports|length }});
let employeeNoticeShownFor = 0;
async function pollEmployeeReports() {
    try {
        const response = await fetch('/api/employee_notifications', { credentials: 'same-origin', cache: 'no-store' });
        if (!response.ok) return;
        const data = await response.json();
        if (!data || !data.ok) return;
        const badge = document.getElementById('employeePendingBadge');
        if (badge) {
            badge.textContent = String(data.pending_count || 0);
            badge.classList.toggle('d-none', Number(data.pending_count || 0) === 0);
        }
        if (data.latest && Number(data.latest.id) > lastEmployeePendingId && Number(data.latest.id) !== employeeNoticeShownFor) {
            employeeNoticeShownFor = Number(data.latest.id);
            const notice = document.getElementById('employeeLiveNotice');
            const title = document.getElementById('employeeLiveNoticeTitle');
            const detail = document.getElementById('employeeLiveNoticeText');
            if (title) title.textContent = 'گزارش جدید کارمند دریافت شد';
            if (detail) detail.textContent = `${data.latest.employee} — شعبه ${data.latest.branch} — تاریخ ${data.latest.date}؛ برای بررسی وارد تب کارکنان شوید.`;
            if (notice) { notice.classList.remove('d-none'); notice.classList.add('d-flex'); }
            lastEmployeePendingId = Number(data.latest.id);
        }
    } catch (err) {
        // قطع موقت شبکه نباید عملکرد اصلی پنل را مختل کند.
    }
}
function openEmployeeReviewTab() {
    try { window.location.hash = 'tab-employees'; } catch (e) {}
    window.location.reload();
}
// اگر صفحه با لینک تب کارکنان باز شده باشد، همان تب را فعال کن.
document.addEventListener('DOMContentLoaded', function () {
    if (initialEmployeePendingCount > 0) {
        const notice = document.getElementById('employeeLiveNotice');
        const title = document.getElementById('employeeLiveNoticeTitle');
        const detail = document.getElementById('employeeLiveNoticeText');
        if (title) title.textContent = 'گزارش‌های کارمندان منتظر تأیید شما هستند';
        if (detail) detail.textContent = `${initialEmployeePendingCount} گزارش در صف بررسی است. برای مشاهده و تأیید، وارد تب کارکنان شوید.`;
        if (notice) { notice.classList.remove('d-none'); notice.classList.add('d-flex'); }
    }
    if (window.location.hash === '#tab-employees') {
        const tabButton = document.getElementById('tab-employees-btn');
        if (tabButton && window.bootstrap) bootstrap.Tab.getOrCreateInstance(tabButton).show();
    }
    pollEmployeeReports();
    window.setInterval(pollEmployeeReports, 8000);
});
</script>
{% endif %}
</body>
</html>
"""

# همه‌ی فرم‌های HTML دارای توکن CSRF هستند؛ از این طریق فرم‌های ثابت و فرم‌های تب‌ها پوشش داده می‌شوند.
WEB_TEMPLATE = re.sub(
    r"(<form\b[^>]*>)",
    r'\1<input type="hidden" name="csrf_token" value="{{ csrf_token() }}">',
    WEB_TEMPLATE,
    flags=re.IGNORECASE,
)


@flask_app.context_processor
def _inject_security_helpers():
    def csrf_token():
        token = session.get("_csrf_token")
        if not token:
            token = secrets.token_urlsafe(32)
            session["_csrf_token"] = token
        return token
    return {"csrf_token": csrf_token}


@flask_app.route("/api/accounting_sheet")
def api_get_accounting_sheet():
    if not session.get("logged_in"):
        return jsonify({"ok": False, "error": "unauthorized"}), 401
    m_str = request.args.get("month", default_month_key())
    sync_flag = request.args.get("sync", "0") == "1"
    y_i, m_i = parse_month_key(m_str, default_month_tuple())
    sheet = build_month_accounting_rows(y_i, m_i, sync_from_reports=sync_flag)
    return jsonify({"ok": True, **sheet})


@flask_app.route("/api/save_accounting_sheet", methods=["POST"])
def api_save_accounting_sheet():
    if not session.get("logged_in"):
        return jsonify({"ok": False, "error": "unauthorized"}), 401
    payload = request.get_json(force=True, silent=True) or {}
    m_key = str(payload.get("month_key") or default_month_key())
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", m_key):
        return jsonify({"ok": False, "error": "invalid month"}), 400
    save_financial_sheet(
        m_key,
        {
            "dinar_rate": safe_int(payload.get("dinar_rate", 15000)) or 15000,
            "rows": payload.get("rows") if isinstance(payload.get("rows"), list) else [],
        },
    )
    log_action(ADMIN_ID, f"Saved financial accounting sheet for {m_key}")
    # پس از ذخیره موفق، PDF صورت‌حساب به مدیر ربات ارسال می‌شود.
    try:
        yy, mm = map(int, m_key.split('-'))
        pdf_path = create_accounting_pdf_file(yy, mm)
        if GLOBAL_BOT_APP and GLOBAL_LOOP and _fin_pdf_allowed(m_key):
            async def _send_financial_pdf():
                with open(pdf_path, 'rb') as pf:
                    await GLOBAL_BOT_APP.bot.send_document(chat_id=ADMIN_ID, document=pf, filename=os.path.basename(pdf_path), caption=f'📊 صورت‌حساب مالی ماه {m_key}')
            asyncio.run_coroutine_threadsafe(_send_financial_pdf(), GLOBAL_LOOP)
    except Exception as exc:
        logger.exception('Financial PDF/Telegram send failed: %s', exc)
    return jsonify({"ok": True})


@flask_app.route("/export_accounting_pdf")
def web_export_accounting_pdf():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    y_i, m_i = parse_month_key(request.args.get("month", ""), default_month_tuple())
    path = create_accounting_pdf_file(y_i, m_i)
    return send_file(path, as_attachment=True, download_name=f"hesabdari_{y_i}_{m_i:02d}.pdf")


@flask_app.route("/export_accounting_excel")
def web_export_accounting_excel():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    m_str = request.args.get("month", default_month_key())
    y_i, m_i = parse_month_key(m_str, default_month_tuple())
    path = create_accounting_excel_file(y_i, m_i)
    return send_file(path, as_attachment=True, download_name=f"{m_i:02d}-{y_i}.xlsx")


@flask_app.route("/login", methods=["GET", "POST"])
def web_login():
    if request.method == "POST":
        ip = client_ip()
        now = time_module.time() if "time_module" in globals() else datetime.now().timestamp()
        # پاک‌سازی تلاش‌های منقضی‌شده و IPهای قدیمی برای جلوگیری از رشد بی‌نهایت حافظه
        for old_ip in list(_LOGIN_ATTEMPTS):
            _LOGIN_ATTEMPTS[old_ip] = [t for t in _LOGIN_ATTEMPTS[old_ip] if now - t < 600]
            if not _LOGIN_ATTEMPTS[old_ip]:
                _LOGIN_ATTEMPTS.pop(old_ip, None)
        attempts = _LOGIN_ATTEMPTS.get(ip, [])
        attempts = [t for t in attempts if now - t < 600]
        if len(attempts) >= 8:
            _LOGIN_ATTEMPTS[ip] = attempts
            flash("تعداد تلاش‌ها زیاد است؛ ۱۰ دقیقه دیگر دوباره امتحان کنید.", "danger")
        elif verify_web_login(request.form.get("username", ""), request.form.get("password", "")):
            _LOGIN_ATTEMPTS.pop(ip, None)
            creds = get_web_credentials()
            session.clear(); session["logged_in"] = True; session.permanent = True
            session["cred_ver"] = creds["ver"]; session["web_user"] = creds["username"]
            log_action(ADMIN_ID, f"Web login from {ip} as {creds['username']}")
            flash("خوش آمدید!", "success")
            return redirect(url_for("web_dashboard"))
        else:
            attempts.append(now); _LOGIN_ATTEMPTS[ip] = attempts
            log_action(0, f"Failed web login from {ip}")
            flash("❌ نام کاربری یا رمز عبور اشتباه است.", "danger")
    tod = iran_today()
    jd = jdatetime.date.fromgregorian(date=tod)
    return render_template_string(WEB_TEMPLATE, logged_in=False, today_shamsi=jd.strftime("%Y/%m/%d"), footer=FIXED_FOOTER)


@flask_app.route("/logout", methods=["POST"])
def web_logout():
    session.clear()
    return redirect(url_for("web_login"))


@flask_app.route("/api/date_info")
def api_date_info():
    if not session.get("logged_in"):
        return jsonify({"ok": False})
    try:
        d = parse_smart_date(request.args.get("date", ""))
        ds, sh_d = d.strftime("%d/%m/%Y"), shamsi(d)
    except Exception:
        return jsonify({"ok": False})
    ex = load_report(ds)
    sug = {}
    for bname in get_known_branches():
        le, ld = get_last_end_serial(bname, before_date_str=ds)
        if le:
            sug[bname] = {"last_end": le, "next_start": le + 1, "last_date": ld}
    return jsonify({"ok": True, "gregorian": ds, "shamsi": sh_d, "existing": ex, "suggestions": sug})




def get_closed_workday(report_date):
    conn = db()
    row = conn.execute("SELECT report_date,shamsi_date,status,transit,export,closed_by,closed_at,close_note FROM closed_workdays WHERE report_date=?", (report_date,)).fetchone()
    conn.close()
    if not row:
        return None
    keys = ["report_date","shamsi_date","status","transit","export","closed_by","closed_at","close_note"]
    return dict(zip(keys, row))


def log_report_change(report_date, action, reason, user_id):
    conn = db()
    conn.execute("INSERT INTO report_change_log(report_date,action,reason,user_id,created_at) VALUES(?,?,?,?,?)",
                 (report_date, action, reason.strip(), int(user_id), datetime.now(TEHRAN).isoformat(timespec="seconds")))
    conn.commit()
    conn.close()
    log_action(user_id, f"{action} report {report_date}; reason={reason.strip()}")


def serial_warnings_for_range(start, end, report_date, branch_name):
    """Return advisory warnings for overlapping ranges or suspicious serial gaps."""
    try:
        a, z = int(tr_num(start)), int(tr_num(end))
    except Exception:
        return []
    if a <= 0 or z < a:
        return ["⚠️ بازه سریال نامعتبر است."]
    intervals = []
    for rep in all_reports():
        if rep.get("date") == report_date:
            # Existing same-day data may be intentionally replaced by manager; still flag overlap.
            pass
        for b in rep.get("branches", []):
            for rg in b.get("ranges", [{"start": b.get("start", "0"), "end": b.get("end", "0")} ]):
                try:
                    x, y = int(tr_num(rg.get("start", "0"))), int(tr_num(rg.get("end", "0")))
                    if x > 0 and y >= x:
                        intervals.append((x, y, rep.get("shamsi_date", rep.get("date", "-")), b.get("name", "-")))
                except Exception:
                    continue
    warnings = []
    for x, y, dt, bn in intervals:
        if a <= y and z >= x:
            warnings.append(f"⚠️ هم‌پوشانی احتمالی با سریال {x} تا {y} در شعبه {bn} ({dt})")
            if len(warnings) >= 4:
                break
    if intervals:
        prior = [it for it in intervals if it[1] < a]
        if prior:
            last = max(prior, key=lambda it: it[1])
            if a > last[1] + 1:
                warnings.append(f"⚠️ پرش احتمالی سریال: آخرین پایان قبلی {last[1]} بوده و بازه جدید از {a} شروع می‌شود.")
    return warnings


def get_branch_employees():
    """فهرست کارمندان همراه با نام تلگرام، سطح دسترسی و آمار عملکرد."""
    conn = db()
    try:
        rows = conn.execute("""SELECT e.user_id, e.employee_name, e.branch_name, e.active, COALESCE(e.can_submit,1), COALESCE(e.note,''),
            e.created_at, COALESCE(u.username,''), COALESCE(u.first_name,''), CASE WHEN u.user_id IS NULL THEN 0 ELSE 1 END,
            (SELECT COUNT(*) FROM pending_branch_reports p WHERE p.user_id=e.user_id),
            (SELECT COUNT(*) FROM pending_branch_reports p WHERE p.user_id=e.user_id AND p.status='approved'),
            (SELECT COUNT(*) FROM pending_branch_reports p WHERE p.user_id=e.user_id AND p.status='rejected'),
            (SELECT COUNT(*) FROM pending_branch_reports p WHERE p.user_id=e.user_id AND p.status='pending'),
            (SELECT p.shamsi_date FROM pending_branch_reports p WHERE p.user_id=e.user_id ORDER BY p.id DESC LIMIT 1)
            FROM branch_employees e LEFT JOIN users u ON u.user_id=e.user_id
            ORDER BY e.branch_name, e.employee_name""").fetchall()
    finally:
        conn.close()
    keys = ["user_id", "employee_name", "branch_name", "active", "can_submit", "note", "created_at", "username", "first_name",
            "started", "total", "approved", "rejected", "pending", "last_date"]
    return [dict(zip(keys, r)) for r in rows]


EMP_BRANCHES = ("تیرپارک", "پارسیان")


def get_employee_candidates():
    """کاربرانی که ربات را استارت کرده‌اند و هنوز کارمند نیستند (برای انتخاب سریع به‌جای تایپ شناسه)."""
    conn = db()
    try:
        rows = conn.execute("""SELECT user_id, COALESCE(username,''), COALESCE(first_name,''), allowed FROM users
            WHERE is_admin=0 AND user_id NOT IN (SELECT user_id FROM branch_employees)
            ORDER BY created_at DESC LIMIT 50""").fetchall()
    finally:
        conn.close()
    return [{"user_id": r[0], "username": r[1], "first_name": r[2], "allowed": r[3]} for r in rows]


def update_employee(uid, name=None, branch=None, active=None, can_submit=None, note=None):
    cols = (("employee_name", name), ("branch_name", branch), ("active", active), ("can_submit", can_submit), ("note", note))
    sets = [f"{c}=?" for c, v in cols if v is not None]
    vals = [v for c, v in cols if v is not None]
    if not sets:
        return False
    conn = db()
    try:
        cur = conn.execute(f"UPDATE branch_employees SET {', '.join(sets)} WHERE user_id=?", (*vals, uid))
        conn.commit()
        return cur.rowcount > 0
    finally:
        conn.close()


def delete_employee(uid):
    """حذف کارمند؛ (حذف شد؟, تعداد گزارش‌های منتظر تأیید او) را برمی‌گرداند."""
    conn = db()
    try:
        pending = conn.execute("SELECT COUNT(*) FROM pending_branch_reports WHERE user_id=? AND status='pending'", (uid,)).fetchone()[0]
        cur = conn.execute("DELETE FROM branch_employees WHERE user_id=?", (uid,))
        conn.commit()
        return cur.rowcount > 0, pending
    finally:
        conn.close()


async def _safe_send(uid, text):
    try:
        await GLOBAL_BOT_APP.bot.send_message(chat_id=uid, text=text)
    except Exception:
        logger.info("notify failed for %s", uid)


def notify_user_async(uid, text):
    if GLOBAL_BOT_APP and GLOBAL_LOOP:
        try:
            asyncio.run_coroutine_threadsafe(_safe_send(uid, text), GLOBAL_LOOP)
        except Exception:
            logger.exception("notify_user_async failed")


def employee_change_message(old, active=None, can_submit=None, branch=None):
    parts = []
    if active is not None and bool(active) != bool(old.get("active")):
        parts.append("✅ دسترسی شما فعال شد." if active else "⏸ دسترسی شما توسط مدیر غیرفعال شد.")
    if can_submit is not None and bool(can_submit) != bool(old.get("can_submit", 1)):
        parts.append("📝 اجازه ثبت آمار برای شما فعال شد." if can_submit else "👁 دسترسی شما به «فقط مشاهده» تغییر کرد.")
    if branch and branch != old.get("branch_name"):
        parts.append(f"🏢 شعبه شما به {branch} تغییر کرد.")
    return "\n".join(parts)


def get_employee(uid, active_only=True):
    conn = db()
    q = "SELECT user_id, employee_name, branch_name, active, COALESCE(can_submit, 1), COALESCE(note, '') FROM branch_employees WHERE user_id=?"
    row = conn.execute(q, (uid,)).fetchone()
    conn.close()
    if not row or (active_only and not row[3]):
        return None
    return {"user_id": row[0], "employee_name": row[1], "branch_name": row[2], "active": row[3], "can_submit": row[4], "note": row[5]}


def get_pending_branch_reports():
    conn = db()
    rows = conn.execute("""SELECT p.id,p.user_id,p.report_date,p.shamsi_date,p.branch_name,p.issued,
        p.serial_start,p.serial_end,p.cancelled,p.cancelled_serial,p.created_at,u.employee_name
        FROM pending_branch_reports p LEFT JOIN branch_employees u ON u.user_id=p.user_id
        WHERE p.status='pending' ORDER BY p.created_at DESC""").fetchall()
    conn.close()
    keys = ["id","user_id","report_date","shamsi_date","branch_name","issued","serial_start","serial_end","cancelled","cancelled_serial","created_at","employee_name"]
    items = [dict(zip(keys, r)) for r in rows]
    for item in items:
        warnings = serial_warnings_for_range(item["serial_start"], item["serial_end"], item["report_date"], item["branch_name"])
        item["serial_warnings"] = " | ".join(warnings) if warnings else "بدون هشدار مشخص"
    return items


def latest_pending_report_id():
    """آخرین شناسه ثبت‌شده؛ برای تشخیص گزارش‌های جدید در پنل مدیر."""
    conn = db()
    row = conn.execute("SELECT COALESCE(MAX(id), 0) FROM pending_branch_reports").fetchone()
    conn.close()
    return int(row[0] or 0)


def get_branch_daily_status(report_date=None):
    """نمای خلاصه وضعیت دو شعبه برای یک روز، بدون دست‌کاری گزارش نهایی."""
    try:
        d = parse_smart_date(report_date) if report_date else iran_today()
    except Exception:
        d = iran_today()
    ds = d.strftime("%d/%m/%Y")
    shd = shamsi(d)
    saved = load_report(ds) or {}
    saved_names = {str(b.get("name", "")).strip() for b in saved.get("branches", [])}
    conn = db()
    rows = conn.execute("""SELECT branch_name, COUNT(*) FROM pending_branch_reports
        WHERE report_date=? AND status='pending' GROUP BY branch_name""", (ds,)).fetchall()
    conn.close()
    pending_counts = {r[0]: int(r[1]) for r in rows}
    result = []
    for branch in ("تیرپارک", "پارسیان"):
        if branch in saved_names:
            status, status_class = "ثبت نهایی شده", "success"
        elif pending_counts.get(branch, 0):
            status, status_class = f"در انتظار تأیید ({pending_counts[branch]})", "warning"
        else:
            status, status_class = "گزارش دریافت نشده", "secondary"
        result.append({"branch_name": branch, "status": status, "status_class": status_class,
                       "pending_count": pending_counts.get(branch, 0), "date": shd})
    return result


@flask_app.route("/api/employee_notifications")
def api_employee_notifications():
    if not session.get("logged_in"):
        return jsonify({"ok": False}), 401
    conn = db()
    row = conn.execute("""SELECT p.id, p.shamsi_date, p.branch_name, p.created_at,
        COALESCE(e.employee_name, 'کارمند'), p.status
        FROM pending_branch_reports p
        LEFT JOIN branch_employees e ON e.user_id=p.user_id
        ORDER BY p.id DESC LIMIT 1""").fetchone()
    count_row = conn.execute("SELECT COUNT(*) FROM pending_branch_reports WHERE status='pending'").fetchone()
    conn.close()
    latest = None
    if row:
        latest = {"id": int(row[0]), "date": row[1], "branch": row[2], "created_at": row[3], "employee": row[4], "status": row[5]}
    return jsonify({"ok": True, "latest_id": latest["id"] if latest else 0,
                    "pending_count": int(count_row[0] or 0), "latest": latest})


def review_employee_report(pid, approve=True, reviewer=ADMIN_ID):
    conn = db()
    row = conn.execute("""SELECT id,user_id,report_date,shamsi_date,branch_name,issued,serial_start,serial_end,cancelled,cancelled_serial
                          FROM pending_branch_reports WHERE id=? AND status='pending'""", (pid,)).fetchone()
    if not row:
        conn.close()
        return False, "این درخواست قبلاً بررسی شده یا وجود ندارد."
    _, uid, ds, shd, branch, issued, st, en, cancelled, cancelled_serial = row
    day_lock = get_closed_workday(ds)
    if approve and day_lock and day_lock.get("status") == "closed":
        conn.close()
        return False, f"روز {shd} سابقه نهایی‌سازی دارد؛ برای حفظ دلیل اصلاح، گزارش را از پنل در فرایند نهایی‌سازی روزانه بررسی کنید."
    if approve:
        existing = load_report(ds)
        if existing is None:
            existing = {"date": ds, "shamsi_date": shd, "branches": [], "transit": 0, "export": 0, "created_by": reviewer}
        branches = [b for b in existing.get("branches", []) if b.get("name") != branch]
        branches.append({"name": branch, "issued": int(issued), "ranges":[{"start":str(st),"end":str(en)}],
                         "start":str(st), "end":str(en), "range_text":f"{st} الی {en}",
                         "cancelled":int(cancelled), "cancelled_serial":str(cancelled_serial)})
        existing["branches"] = branches
        save_report(existing)
        action = "approved"
    else:
        action = "rejected"
    conn.execute("UPDATE pending_branch_reports SET status=?, reviewed_at=? WHERE id=?",
                 (action, datetime.now(TEHRAN).isoformat(timespec="seconds"), pid))
    conn.commit()
    conn.close()
    log_action(reviewer, f"{action} employee branch report #{pid} for {ds} ({branch})")
    return True, f"{'تأیید و ثبت' if approve else 'رد'} آمار شعبه {branch} برای {shd} انجام شد."


def get_recent_closed_workdays(limit=20):
    conn = db()
    rows = conn.execute("SELECT report_date,shamsi_date,status,transit,export,closed_by,closed_at,close_note FROM closed_workdays ORDER BY closed_at DESC LIMIT ?", (limit,)).fetchall()
    conn.close()
    keys = ["report_date","shamsi_date","status","transit","export","closed_by","closed_at","close_note"]
    return [dict(zip(keys, r)) for r in rows]


@flask_app.route("/employees/finalize_day", methods=["POST"])
def web_finalize_workday():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    try:
        d = parse_smart_date(request.form.get("report_date", ""))
        transit = safe_int(request.form.get("transit", "0"))
        export = safe_int(request.form.get("export", "0"))
        if transit is None or export is None:
            raise ValueError()
    except Exception:
        flash("تاریخ، ترانزیت یا صادرات معتبر نیست.", "danger")
        return redirect(url_for("web_dashboard") + "#tab-employees")
    ds, shd = d.strftime("%d/%m/%Y"), shamsi(d)
    reason = str(request.form.get("close_note", "")).strip()[:500]
    if not reason:
        flash("برای نهایی‌سازی روز، یادداشت یا توضیح کوتاه الزامی است.", "warning")
        return redirect(url_for("web_dashboard") + "#tab-employees")
    existing_lock = get_closed_workday(ds)
    if existing_lock and existing_lock.get("status") == "closed":
        flash("این روز قبلاً بسته شده است. برای تغییر ابتدا آن را با دلیل بازگشایی کنید.", "warning")
        return redirect(url_for("web_dashboard") + "#tab-employees")
    saved = load_report(ds) or {"date": ds, "shamsi_date": shd, "branches": [], "transit": 0, "export": 0, "created_by": ADMIN_ID}
    saved_map = {str(b.get("name", "")).strip(): b for b in saved.get("branches", [])}
    conn = db()
    rows = conn.execute("""SELECT id,branch_name,issued,serial_start,serial_end,cancelled,cancelled_serial,shamsi_date
        FROM pending_branch_reports WHERE report_date=? AND status='pending' ORDER BY id""", (ds,)).fetchall()
    pending_by_branch = {}
    for row in rows:
        pending_by_branch[row[1]] = row
    # Every required branch must either already exist in the saved report or have a pending submission.
    missing = [bn for bn in ("تیرپارک", "پارسیان") if bn not in saved_map and bn not in pending_by_branch]
    if missing:
        conn.close()
        flash("نهایی‌سازی ممکن نیست؛ گزارش این شعبه‌ها موجود نیست: " + "، ".join(missing), "warning")
        return redirect(url_for("web_dashboard") + "#tab-employees")
    for bn, row in pending_by_branch.items():
        rid, _, issued, st, en, cancelled, cancelled_serial, _ = row
        saved_map[bn] = {"name": bn, "issued": int(issued), "ranges":[{"start":str(st),"end":str(en)}], "start":str(st), "end":str(en), "range_text":f"{st} الی {en}", "cancelled":int(cancelled), "cancelled_serial":str(cancelled_serial)}
    conn.close()
    saved["branches"] = [saved_map[bn] for bn in ("تیرپارک", "پارسیان") if bn in saved_map] + [b for bn,b in saved_map.items() if bn not in ("تیرپارک", "پارسیان")]
    saved["transit"], saved["export"], saved["created_by"] = transit, export, ADMIN_ID
    save_report(saved)
    now = datetime.now(TEHRAN).isoformat(timespec="seconds")
    conn = db()
    conn.execute("""INSERT INTO closed_workdays(report_date,shamsi_date,status,transit,export,closed_by,closed_at,close_note)
        VALUES(?,?, 'closed', ?,?,?,?,?)
        ON CONFLICT(report_date) DO UPDATE SET shamsi_date=excluded.shamsi_date,status='closed',transit=excluded.transit,
        export=excluded.export,closed_by=excluded.closed_by,closed_at=excluded.closed_at,close_note=excluded.close_note""",
        (ds, shd, transit, export, ADMIN_ID, now, reason))
    staff_rows = conn.execute("SELECT DISTINCT user_id,branch_name FROM pending_branch_reports WHERE report_date=? AND status='pending'", (ds,)).fetchall()
    conn.execute("UPDATE pending_branch_reports SET status='approved', reviewed_at=? WHERE report_date=? AND status='pending'",
                 (now, ds))
    conn.commit()
    conn.close()
    log_report_change(ds, "نهایی‌سازی و بستن روز", reason, ADMIN_ID)
    if GLOBAL_BOT_APP and GLOBAL_LOOP:
        msg = f"🔐 روز کاری نهایی و بسته شد\n📅 {shd}\n🏢 تیرپارک و پارسیان بررسی شدند\n🚛 ترانزیت مدیر: {transit}\n📦 صادرات مدیر: {export}\n📝 توضیح: {reason}"
        try:
            asyncio.run_coroutine_threadsafe(GLOBAL_BOT_APP.bot.send_message(chat_id=ADMIN_ID, text=msg), GLOBAL_LOOP)
            for staff_uid, staff_branch in staff_rows:
                staff_msg = f"✅ گزارش شعبه {staff_branch} برای تاریخ {shd} در گزارش نهایی روز ثبت شد."
                asyncio.run_coroutine_threadsafe(GLOBAL_BOT_APP.bot.send_message(chat_id=staff_uid, text=staff_msg), GLOBAL_LOOP)
        except Exception:
            pass
    flash(f"روز {shd} با آمار هر دو شعبه نهایی و بسته شد.", "success")
    return redirect(url_for("web_dashboard") + "#tab-employees")


@flask_app.route("/employees/reopen_day", methods=["POST"])
def web_reopen_workday():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    ds = str(request.form.get("report_date", "")).strip()
    reason = str(request.form.get("reason", "")).strip()[:500]
    if not reason:
        flash("برای بازگشایی، ذکر دلیل الزامی است.", "warning")
        return redirect(url_for("web_dashboard") + "#tab-employees")
    try:
        ds = parse_smart_date(ds).strftime("%d/%m/%Y")
    except Exception:
        flash("تاریخ بازگشایی معتبر نیست.", "danger")
        return redirect(url_for("web_dashboard") + "#tab-employees")
    closed = get_closed_workday(ds)
    if not closed or closed.get("status") != "closed":
        flash("این روز بسته‌شده‌ای برای بازگشایی ندارد.", "warning")
        return redirect(url_for("web_dashboard") + "#tab-employees")
    conn = db()
    conn.execute("UPDATE closed_workdays SET status='open', close_note=close_note || ? WHERE report_date=?", (f" | بازگشایی: {reason}", ds))
    conn.commit()
    conn.close()
    log_report_change(ds, "بازگشایی روز", reason, ADMIN_ID)
    flash(f"روز {closed['shamsi_date']} با ثبت دلیل بازگشایی شد؛ پس از اصلاح، دوباره نهایی‌سازی کنید.", "success")
    return redirect(url_for("web_dashboard") + "#tab-employees")


@flask_app.route("/employees/add", methods=["POST"])
def web_employee_add():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    uid = safe_int(request.form.get("user_id", ""))
    name = request.form.get("employee_name", "").strip()[:100]
    branch = request.form.get("branch_name", "").strip()
    can_submit = 0 if request.form.get("access") == "view" else 1
    note = request.form.get("note", "").strip()[:200]
    if not uid or not name or branch not in EMP_BRANCHES:
        flash("اطلاعات کارمند معتبر نیست.", "danger")
    elif is_admin_id(uid):
        flash("این شناسه متعلق به مدیر است و نیازی به ثبت به‌عنوان کارمند ندارد.", "warning")
    else:
        conn = db()
        conn.execute("""INSERT INTO branch_employees(user_id,employee_name,branch_name,active,created_at,can_submit,note)
            VALUES(?,?,?,1,?,?,?) ON CONFLICT(user_id) DO UPDATE SET employee_name=excluded.employee_name,
            branch_name=excluded.branch_name, active=1, can_submit=excluded.can_submit, note=excluded.note""",
            (uid, name, branch, datetime.now(TEHRAN).isoformat(timespec="seconds"), can_submit, note))
        conn.execute("UPDATE access_requests SET status='approved' WHERE user_id=? AND status='pending'", (uid,))
        conn.commit()
        conn.close()
        log_action(ADMIN_ID, f"Registered branch employee {uid} for {branch} (submit={can_submit})")
        if request.form.get("notify") == "1":
            notify_user_async(uid, f"✅ شما به‌عنوان کارمند شعبه {branch} ثبت شدید"
                                   f"{'' if can_submit else ' (فقط مشاهده)'}. برای شروع /start را بزنید.")
        flash("کارمند ثبت شد. اکنون می‌تواند از ربات استفاده کند.", "success")
    return redirect(url_for("web_dashboard") + "#tab-employees")


@flask_app.route("/employees/update", methods=["POST"])
def web_employee_update():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    uid = safe_int(request.form.get("user_id", ""))
    old = get_employee(uid, active_only=False) if uid else None
    name = request.form.get("employee_name", "").strip()[:100]
    branch = request.form.get("branch_name", "").strip()
    if not old or not name or branch not in EMP_BRANCHES:
        flash("اطلاعات ویرایش معتبر نیست.", "danger")
        return redirect(url_for("web_dashboard") + "#tab-employees")
    active = 1 if request.form.get("active", "1") == "1" else 0
    can_submit = 0 if request.form.get("access") == "view" else 1
    note = request.form.get("note", "").strip()[:200]
    update_employee(uid, name=name, branch=branch, active=active, can_submit=can_submit, note=note)
    log_action(ADMIN_ID, f"Updated employee {uid}: branch={branch} active={active} submit={can_submit}")
    if request.form.get("notify") == "1":
        msg = employee_change_message(old, active=active, can_submit=can_submit, branch=branch)
        if msg:
            notify_user_async(uid, msg)
    flash("تغییرات کارمند ذخیره شد.", "success")
    return redirect(url_for("web_dashboard") + "#tab-employees")


@flask_app.route("/employees/toggle", methods=["POST"])
def web_employee_toggle():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    uid = safe_int(request.form.get("user_id", ""))
    old = get_employee(uid, active_only=False) if uid else None
    if old:
        new_active = 0 if old["active"] else 1
        update_employee(uid, active=new_active)
        log_action(ADMIN_ID, f"Toggled branch employee {uid} -> {new_active}")
        notify_user_async(uid, employee_change_message(old, active=new_active))
        flash(f"دسترسی {old['employee_name']} {'فعال' if new_active else 'غیرفعال'} شد.", "info")
    return redirect(url_for("web_dashboard") + "#tab-employees")


@flask_app.route("/employees/delete", methods=["POST"])
def web_employee_delete():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    uid = safe_int(request.form.get("user_id", ""))
    old = get_employee(uid, active_only=False) if uid else None
    if not old:
        flash("کارمند پیدا نشد.", "warning")
        return redirect(url_for("web_dashboard") + "#tab-employees")
    ok, pending = delete_employee(uid)
    if ok:
        log_action(ADMIN_ID, f"Deleted branch employee {uid} ({old['employee_name']})")
        if request.form.get("notify") == "1":
            notify_user_async(uid, "⛔ دسترسی کارمندی شما توسط مدیر حذف شد.")
        extra = f" {pending} گزارش منتظر تأیید او همچنان در صف بررسی است." if pending else ""
        flash(f"کارمند «{old['employee_name']}» حذف شد.{extra}", "warning")
    return redirect(url_for("web_dashboard") + "#tab-employees")


@flask_app.route("/employees/review", methods=["POST"])
def web_employee_review():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    pid = safe_int(request.form.get("pending_id", ""))
    approve = request.form.get("decision") == "approve"
    if pid:
        conn = db()
        employee_row = conn.execute("SELECT user_id,branch_name,shamsi_date FROM pending_branch_reports WHERE id=?", (pid,)).fetchone()
        conn.close()
        ok, msg = review_employee_report(pid, approve, ADMIN_ID)
        flash(msg, "success" if ok else "warning")
        if ok and GLOBAL_BOT_APP and GLOBAL_LOOP:
            try:
                asyncio.run_coroutine_threadsafe(GLOBAL_BOT_APP.bot.send_message(chat_id=ADMIN_ID, text=f"{'✅' if approve else '❌'} {msg}"), GLOBAL_LOOP)
                if employee_row:
                    uid, branch_name, shamsi_date = employee_row
                    staff_msg = (f"✅ گزارش شعبه {branch_name} برای تاریخ {shamsi_date} تأیید و ثبت شد." if approve else f"❌ گزارش شعبه {branch_name} برای تاریخ {shamsi_date} رد شد. لطفاً با مدیر هماهنگ کنید.")
                    asyncio.run_coroutine_threadsafe(GLOBAL_BOT_APP.bot.send_message(chat_id=uid, text=staff_msg), GLOBAL_LOOP)
            except Exception:
                pass
    return redirect(url_for("web_dashboard") + "#tab-employees")


async def employee_amar_status(update, context):
    """نمایش وضعیت پنج گزارش اخیر کارمند."""
    user = update.effective_user
    emp = get_employee(user.id, active_only=False)
    if not emp:
        await update.message.reply_text("شناسه شما در فهرست کارکنان ثبت نشده است. با مدیر سامانه تماس بگیرید.")
        return
    conn = db()
    rows = conn.execute("""SELECT id,shamsi_date,status,created_at,reviewed_at
        FROM pending_branch_reports WHERE user_id=? ORDER BY id DESC LIMIT 5""", (user.id,)).fetchall()
    conn.close()
    if not rows:
        await update.message.reply_text(
            f"👤 {emp['employee_name']} — شعبه {emp['branch_name']}\n"
            "هنوز گزارشی ثبت نکرده‌اید. از دکمه «📝 ثبت آمار شعبه» استفاده کنید.")
        return
    labels = {"pending": "⏳ در انتظار تأیید", "approved": "✅ تأیید و ثبت شد", "rejected": "❌ رد شد"}
    lines = [f"📋 وضعیت گزارش‌های اخیر — {emp['branch_name']}"]
    for rid, date_text, status, created, reviewed in rows:
        lines.append(f"#{rid} | {date_text} | {labels.get(status, status)}")
    lines.append("\nبرای ارسال گزارش جدید از دکمه «📝 ثبت آمار شعبه» استفاده کنید.")
    await update.message.reply_text("\n".join(lines))


async def employee_help(update, context):
    user = update.effective_user
    emp = get_employee(user.id)
    if emp:
        await update.message.reply_text(
            f"راهنمای کارمند: {emp['employee_name']} — شعبه {emp['branch_name']}\n\n"
            "از دکمه‌های پایین استفاده کنید: «ثبت آمار شعبه»، «وضعیت گزارش‌های من»، «وضعیت دو شعبه» و «لغو ثبت آمار».\n\n"
            "ثبت آمار مرحله‌به‌مرحله است: تاریخ میلادی (مثلاً 2026-10-23)، صادره، سریال شروع، سریال پایان، ابطالی و سریال ابطالی. هر بار فقط پاسخ همان مرحله را بفرستید.\n"
            "هر تاریخ برای هر شعبه فقط یک گزارش تأییدشده/در انتظار می‌پذیرد؛ اصلاح گزارش نهایی با مدیر است.")
    else:
        await update.message.reply_text("شما به‌عنوان کارمند فعال ثبت نشده‌اید. با مدیر سامانه تماس بگیرید.")


async def admin_branch_status(update, context):
    if not is_admin(update):
        emp = get_employee(update.effective_user.id)
        if emp:
            conn = db()
            rows = conn.execute("SELECT shamsi_date,status FROM pending_branch_reports WHERE report_date=? AND branch_name=? ORDER BY id DESC LIMIT 1", (iran_today().strftime("%d/%m/%Y"), emp["branch_name"])).fetchone()
            conn.close()
            state = ("در انتظار تأیید" if rows and rows[1] == "pending" else "تأیید و ثبت شده" if rows and rows[1] == "approved" else "رد شده" if rows and rows[1] == "rejected" else "گزارشی برای امروز ثبت نشده")
            await update.message.reply_text(f"📊 وضعیت امروز\nشعبه: {emp['branch_name']}\nوضعیت: {state}", reply_markup=employee_menu())
            return
        await update.message.reply_text("⛔ این گزینه فقط برای مدیر یا کارکنان ثبت‌شده است.")
        return
    try:
        target = parse_smart_date(context.args[0]) if context.args else iran_today()
    except Exception:
        await update.message.reply_text("تاریخ معتبر نیست. نمونه: /branchstatus 1405/07/30")
        return
    ds = target.strftime("%d/%m/%Y")
    shd = shamsi(target)
    saved = load_report(ds) or {}
    saved_map = {str(b.get("name", "")).strip(): b for b in saved.get("branches", [])}
    conn = db()
    rows = conn.execute("""SELECT branch_name, COUNT(*), MAX(id) FROM pending_branch_reports
        WHERE report_date=? AND status='pending' GROUP BY branch_name""", (ds,)).fetchall()
    conn.close()
    pending = {r[0]: (int(r[1]), int(r[2])) for r in rows}
    lines = [f"📊 وضعیت روزانه شعب — {shd}"]
    for branch in ("تیرپارک", "پارسیان"):
        if branch in saved_map:
            b = saved_map[branch]
            lines.append(f"✅ {branch}: ثبت نهایی | صادره: {b.get('issued', 0)} | ابطالی: {b.get('cancelled', 0)}")
        elif branch in pending:
            lines.append(f"⏳ {branch}: {pending[branch][0]} گزارش منتظر تأیید (آخرین درخواست #{pending[branch][1]})")
        else:
            lines.append(f"⚠️ {branch}: هنوز گزارش نهایی یا درخواست منتظر تأیید ندارد")
    lines.append(f"🚛 ترانزیت مدیر: {saved.get('transit', 0)} | 📦 صادرات مدیر: {saved.get('export', 0)}")
    lines.append("برای بررسی گزارش‌های کارکنان، تب «کارکنان شعب و تأیید آمار» در پنل را باز کنید.")
    await update.message.reply_text("\n".join(lines))


async def admin_workday_control(update, context):
    if not is_admin(update):
        await update.message.reply_text("⛔ این بخش فقط برای مدیر سامانه است.")
        return
    d = iran_today()
    ds = d.strftime("%d/%m/%Y")
    lock = get_closed_workday(ds)
    conn = db()
    rows = conn.execute("SELECT branch_name,COUNT(*) FROM pending_branch_reports WHERE report_date=? AND status='pending' GROUP BY branch_name", (ds,)).fetchall()
    conn.close()
    pending = {r[0]: int(r[1]) for r in rows}
    if lock and lock.get("status") == "closed":
        state = "🔒 امروز نهایی و بسته شده است."
    elif lock and lock.get("status") == "open":
        state = "🟡 امروز بازگشایی شده و پس از اصلاح باید دوباره نهایی شود."
    else:
        state = "🟢 امروز هنوز نهایی نشده است."
    msg = [f"🔐 کنترل روز کاری — {shamsi(d)}", state,
           f"تیرپارک: {pending.get('تیرپارک', 0)} گزارش منتظر تأیید",
           f"پارسیان: {pending.get('پارسیان', 0)} گزارش منتظر تأیید",
           "برای نهایی‌سازی یک‌جای هر دو شعبه، ترانزیت و صادرات و یادداشت نهایی، پنل سایت را در تب کارکنان باز کنید."]
    await update.message.reply_text("\n".join(msg), reply_markup=management_menu())


async def employee_amar_cancel(update, context):
    context.user_data.pop("employee_report_waiting", None)
    await update.message.reply_text("عملیات لغو شد.", reply_markup=get_menu(update))


async def employee_amar_start(update, context):
    user = update.effective_user
    emp = get_employee(user.id)
    if not emp:
        await update.message.reply_text("⛔ شناسه شما به‌عنوان کارمند فعال ثبت نشده است. با مدیر سامانه تماس بگیرید.")
        return
    if not emp.get("can_submit", 1):
        await update.message.reply_text("👁 دسترسی شما «فقط مشاهده» است و اجازه ثبت آمار ندارید. با مدیر هماهنگ کنید.")
        return
    context.user_data["employee_report_waiting"] = True
    context.user_data["employee_report_step"] = "date"
    context.user_data["employee_report_data"] = {}
    await update.message.reply_text(
        f"سلام {emp['employee_name']}؛ شعبه شما: {emp['branch_name']}\n\n"
        "📝 ثبت آمار را مرحله‌به‌مرحله انجام می‌دهیم.\n"
        "مرحله ۱ از ۶ — تاریخ میلادی را وارد کنید.\n"
        "قالب پیشنهادی: YYYY-MM-DD\nمثال: 2026-10-23\n\n"
        "برای لغو، دکمه «❌ لغو ثبت آمار» را بزنید.",
        reply_markup=employee_menu())


async def employee_amar_receive(update, context):
    user = update.effective_user
    if not context.user_data.get("employee_report_waiting"):
        return
    emp = get_employee(user.id)
    if not emp:
        context.user_data.pop("employee_report_waiting", None)
        context.user_data.pop("employee_report_step", None)
        context.user_data.pop("employee_report_data", None)
        await update.message.reply_text("دسترسی کارمندی شما فعال نیست.")
        return
    if not emp.get("can_submit", 1):
        context.user_data.pop("employee_report_waiting", None)
        context.user_data.pop("employee_report_step", None)
        context.user_data.pop("employee_report_data", None)
        await update.message.reply_text("👁 اجازه ثبت آمار برای شما غیرفعال شده است.")
        return

    step = context.user_data.get("employee_report_step", "date")
    data = context.user_data.setdefault("employee_report_data", {})
    raw = (update.message.text or "").strip()
    try:
        if step == "date":
            d = None
            for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y"):
                try:
                    d = datetime.strptime(raw, fmt).date()
                    break
                except ValueError:
                    pass
            if d is None or d.year < 2000 or d > (iran_today() + timedelta(days=1)):
                raise ValueError("تاریخ")
            data["date"] = d.strftime("%d/%m/%Y")
            data["shamsi_date"] = shamsi(d)
            context.user_data["employee_report_step"] = "issued"
            await update.message.reply_text("مرحله ۲ از ۶ — تعداد **صادره** را فقط به عدد وارد کنید:", parse_mode="Markdown")
            return

        if step == "issued":
            value = safe_int(tr_num(raw))
            if value is None or value < 0:
                raise ValueError("صادره")
            data["issued"] = value
            context.user_data["employee_report_step"] = "serial_start"
            await update.message.reply_text("مرحله ۳ از ۶ — **سریال شروع صادره** را وارد کنید:", parse_mode="Markdown")
            return

        if step == "serial_start":
            value = tr_num(raw).replace(" ", "")
            if not value.isdigit():
                raise ValueError("سریال شروع")
            data["serial_start"] = value
            context.user_data["employee_report_step"] = "serial_end"
            await update.message.reply_text("مرحله ۴ از ۶ — **سریال پایان صادره** را وارد کنید:", parse_mode="Markdown")
            return

        if step == "serial_end":
            value = tr_num(raw).replace(" ", "")
            if not value.isdigit() or int(value) < int(data["serial_start"]):
                raise ValueError("سریال پایان")
            if data["issued"] > 0 and int(value) - int(data["serial_start"]) + 1 != data["issued"]:
                await update.message.reply_text(
                    f"⚠️ بازه سریال ({int(value) - int(data['serial_start']) + 1} شماره) با تعداد صادره ({data['issued']}) برابر نیست.\n"
                    "لطفاً **سریال پایان صادره** را دوباره وارد کنید:", parse_mode="Markdown")
                return
            data["serial_end"] = value
            context.user_data["employee_report_step"] = "cancelled"
            await update.message.reply_text("مرحله ۵ از ۶ — تعداد **ابطالی** را وارد کنید (اگر ندارید، ۰):", parse_mode="Markdown")
            return

        if step == "cancelled":
            value = safe_int(tr_num(raw))
            if value is None or value < 0:
                raise ValueError("ابطالی")
            data["cancelled"] = value
            context.user_data["employee_report_step"] = "cancelled_serial"
            await update.message.reply_text("مرحله ۶ از ۶ — **سریال ابطالی** را وارد کنید؛ اگر ابطالی ندارید، ۰ بفرستید:", parse_mode="Markdown")
            return

        if step == "cancelled_serial":
            value = tr_num(raw).replace(" ", "")
            if not value.isdigit():
                raise ValueError("سریال ابطالی")
            if data.get("cancelled", 0) == 0:
                value = "0"
            elif value == "0":
                raise ValueError("سریال ابطالی")
            data["cancelled_serial"] = value

            ds, shd = data["date"], data["shamsi_date"]
            issued, cancelled = data["issued"], data["cancelled"]
            st, en, cs = data["serial_start"], data["serial_end"], data["cancelled_serial"]
            day_lock = get_closed_workday(ds)
            if day_lock and day_lock.get("status") == "closed":
                await update.message.reply_text("🔒 این روز توسط مدیر نهایی و بسته شده است؛ امکان ارسال گزارش جدید وجود ندارد. اگر اصلاح لازم است، مدیر باید روز را از پنل با ذکر دلیل بازگشایی کند.", reply_markup=employee_menu())
                context.user_data.pop("employee_report_waiting", None)
                context.user_data.pop("employee_report_step", None)
                context.user_data.pop("employee_report_data", None)
                return

            conn = db()
            duplicate = conn.execute("""SELECT id,status FROM pending_branch_reports
                WHERE report_date=? AND branch_name=? AND status IN ('pending','approved')
                ORDER BY id DESC LIMIT 1""", (ds, emp["branch_name"])).fetchone()
            if duplicate:
                conn.close()
                if duplicate[1] == "pending":
                    msg = (f"⏳ برای شعبه {emp['branch_name']} در تاریخ {shd} یک گزارش در انتظار تأیید وجود دارد (درخواست #{duplicate[0]}).\n"
                           "برای جلوگیری از ثبت تکراری، گزارش جدید ساخته نشد. اگر اشتباه است با مدیر تماس بگیرید.")
                else:
                    msg = (f"✅ گزارش شعبه {emp['branch_name']} برای تاریخ {shd} قبلاً تأیید شده است.\n"
                           "برای اصلاح آمار، لطفاً از مدیر بخواهید اصلاح را از پنل انجام دهد.")
                await update.message.reply_text(msg, reply_markup=employee_menu())
                context.user_data.pop("employee_report_waiting", None)
                context.user_data.pop("employee_report_step", None)
                context.user_data.pop("employee_report_data", None)
                return

            cur = conn.execute("""INSERT INTO pending_branch_reports(user_id,report_date,shamsi_date,branch_name,issued,serial_start,serial_end,cancelled,cancelled_serial,status,created_at)
                VALUES(?,?,?,?,?,?,?,?,?,'pending',?)""",
                (user.id, ds, shd, emp["branch_name"], issued, st, en, cancelled, cs, datetime.now(TEHRAN).isoformat(timespec="seconds")))
            pid = cur.lastrowid
            conn.commit()
            conn.close()
            context.user_data.pop("employee_report_waiting", None)
            context.user_data.pop("employee_report_step", None)
            context.user_data.pop("employee_report_data", None)
            await update.message.reply_text(
                f"✅ گزارش شما دریافت شد و در صف بررسی مدیر قرار گرفت.\n🏢 شعبه: {emp['branch_name']}\n📅 تاریخ میلادی: {datetime.strptime(ds, '%d/%m/%Y').strftime('%Y-%m-%d')}\n📅 تاریخ شمسی: {shd}\n🧾 صادره: {issued}\n🔢 سریال: {st} تا {en}\n❌ ابطالی: {cancelled}\n🧾 سریال ابطالی: {cs}\n\nتا قبل از تأیید مدیر، این آمار وارد گزارش نهایی نمی‌شود.",
                reply_markup=employee_menu())
            kb = InlineKeyboardMarkup([[
                InlineKeyboardButton("✅ تأیید و ثبت", callback_data=f"empapprove:{pid}"),
                InlineKeyboardButton("❌ رد", callback_data=f"empreject:{pid}")
            ]])
            serial_warnings = serial_warnings_for_range(st, en, ds, emp["branch_name"])
            admin_text = (
                f"🔔 گزارش جدید کارمند — نیازمند بررسی\n\n🆔 درخواست: #{pid}\n👤 کارمند: {emp['employee_name']}\n"
                f"🪪 شناسه تلگرام: {user.id}\n🏢 شعبه: {emp['branch_name']}\n📅 تاریخ شمسی: {shd}\n📅 تاریخ میلادی: {datetime.strptime(ds, '%d/%m/%Y').strftime('%Y-%m-%d')}\n"
                f"📦 صادره: {issued}\n🔢 سریال: {st} تا {en}\n❌ ابطالی: {cancelled}\n🧾 سریال ابطالی: {cs}\n\n"
                "ℹ️ ترانزیت و صادرات فقط توسط مدیر ثبت می‌شود.\nبرای بررسی از دکمه‌های زیر استفاده کنید یا پنل سایت را باز کنید.")
            if serial_warnings:
                admin_text += "\n\n🕵️ هشدارهای بررسی سریال:\n" + "\n".join(serial_warnings)
            try:
                await context.bot.send_message(chat_id=ADMIN_ID, text=admin_text, reply_markup=kb)
            except Exception as ex:
                logger.warning("Could not notify admin about employee report #%s: %s", pid, ex)
            log_action(user.id, f"Submitted pending branch report #{pid} ({emp['branch_name']})")
            return

        raise ValueError("مرحله نامعتبر")
    except Exception as ex:
        logger.warning("Employee report input validation failed at step %s: %s", step, ex)
        prompts = {
            "date": "تاریخ میلادی معتبر نیست. نمونه: 2026-10-23. دوباره تاریخ را وارد کنید:",
            "issued": "تعداد صادره باید عدد صحیح صفر یا بیشتر باشد. دوباره وارد کنید:",
            "serial_start": "سریال شروع باید عدد باشد. دوباره وارد کنید:",
            "serial_end": "سریال پایان باید عدد و بزرگ‌تر یا مساوی سریال شروع باشد. تعداد سریال‌ها نیز باید با صادره برابر باشد. دوباره وارد کنید:",
            "cancelled": "تعداد ابطالی باید عدد صحیح صفر یا بیشتر باشد. دوباره وارد کنید:",
            "cancelled_serial": "سریال ابطالی باید عدد باشد؛ اگر ابطالی ندارید، ۰ وارد کنید. دوباره وارد کنید:",
        }
        await update.message.reply_text("⚠️ " + prompts.get(step, "ورودی معتبر نیست؛ دوباره تلاش کنید."))


async def employee_review_callback(update, context):
    q = update.callback_query
    await q.answer()
    if not is_admin(update):
        await q.edit_message_text("⛔ فقط مدیر سامانه اجازه تأیید آمار را دارد.")
        return
    try:
        action, raw_id = q.data.split(":", 1)
        pid = int(raw_id)
        conn = db()
        employee_row = conn.execute("SELECT user_id,branch_name,shamsi_date FROM pending_branch_reports WHERE id=?", (pid,)).fetchone()
        conn.close()
        approved = action == "empapprove"
        ok, msg = review_employee_report(pid, approved, ADMIN_ID)
        await q.edit_message_text(("✅ " if ok else "⚠️ ") + msg)
        if ok and employee_row:
            uid, branch_name, shamsi_date = employee_row
            staff_msg = (f"✅ گزارش شعبه {branch_name} برای تاریخ {shamsi_date} تأیید و ثبت شد." if approved else f"❌ گزارش شعبه {branch_name} برای تاریخ {shamsi_date} رد شد. لطفاً با مدیر هماهنگ کنید.")
            try:
                await context.bot.send_message(chat_id=uid, text=staff_msg)
            except Exception:
                logger.info("Could not notify employee %s about report %s", uid, pid)
    except Exception:
        await q.edit_message_text("خطا در بررسی درخواست.")


@flask_app.route("/")
def web_dashboard():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    all_reps = all_reports()
    today = iran_today()
    selected_year, selected_mon = parse_month_key(request.args.get("month", ""), (today.year, today.month))
    selected_month = month_key(selected_year, selected_mon)
    prev_y, prev_m = shift_month(selected_year, selected_mon, -1)
    next_y, next_m = shift_month(selected_year, selected_mon, 1)
    next_month = month_key(next_y, next_m) if (next_y, next_m) <= (today.year, today.month) else ""
    m_start, m_end = month_bounds(selected_year, selected_mon)
    reps = []
    for report in all_reps:
        try:
            report_date = parse_smart_date(report.get("date", ""))
            if report_date.year == selected_year and report_date.month == selected_mon:
                reps.append(report)
        except Exception:
            continue
    kb = get_known_branches()
    td = len(reps)
    ti = sum(total_issued(r) for r in reps)

    max_val, max_dt = 0, "-"
    for r in reps:
        ri = total_issued(r)
        if ri > max_val:
            max_val = ri
            max_dt = r.get("shamsi_date", r.get("date", "-"))

    stats = {
        "total_days": td,
        "total_issued": ti,
        "avg_daily": round(ti / td, 1) if td else 0,
        "total_transit": sum(int(r.get("transit", 0)) for r in reps),
        "total_export": sum(int(r.get("export", 0)) for r in reps),
        "total_cancelled": sum(total_cancelled(r) for r in reps),
        "max_day_val": max_val,
        "max_day_date": max_dt,
    }
    sl = reps[-20:]
    cl = json.dumps([r.get("shamsi_date", r["date"]) for r in sl])
    ci = json.dumps([total_issued(r) for r in sl])
    ce = json.dumps([int(r.get("export", 0)) for r in sl])
    ct = json.dumps([int(r.get("transit", 0)) for r in sl])
    bt = {}
    for r in reps:
        for b in r.get("branches", []):
            bn = b.get("name", "سایر").strip()
            if bn not in {"aa", "ff", "B", "ت"}:
                bt[bn] = bt.get(bn, 0) + int(b.get("issued", 0))
    pl = json.dumps(list(bt.keys()), ensure_ascii=False)
    pv = json.dumps(list(bt.values()))
    # آمار و نمودارها فقط برای ماه انتخاب‌شده هستند؛ آرشیو باید تاریخچه کامل را نشان دهد.
    trps = []
    for r in reversed(all_reps):
        rc = dict(r)
        rc["total_iss"] = total_issued(r)
        rc["total_canc"] = total_cancelled(r)
        for b in rc.get("branches", []):
            b["range_text"] = get_branch_ranges_text(b)
        rc["wa_url"] = make_whatsapp_url(r)
        rc["full_text"] = f"{create_branch_report(r)}\n\n────────────────\n\n{create_trade_report(r)}"
        trps.append(rc)

    tod = iran_today()
    jd = jdatetime.date.fromgregorian(date=tod)
    return render_template_string(
        WEB_TEMPLATE,
        logged_in=True,
        stats=stats,
        reports=trps,
        known_branches=json.dumps(kb, ensure_ascii=False),
        today_shamsi=jd.strftime("%Y/%m/%d"),
        today_gregorian=tod.strftime("%d/%m/%Y"),
        current_month=tod.strftime("%Y-%m"),
        selected_month=selected_month,
        prev_month=month_key(prev_y, prev_m),
        next_month=next_month,
        month_is_current=(selected_month == today.strftime("%Y-%m")),
        month_cmp=month_compare(selected_year, selected_mon),
        available_months=months_with_data(18),
        month_start_str=m_start.strftime("%d/%m/%Y"),
        month_end_str=m_end.strftime("%d/%m/%Y"),
        yesterday_shamsi=jdatetime.date.fromgregorian(date=tod - timedelta(days=1)).strftime("%Y/%m/%d"),
        first_of_month=jdatetime.date(jd.year, jd.month, 1).strftime("%Y/%m/%d"),
        chart_labels=cl,
        chart_issued=ci,
        chart_export=ce,
        chart_transit=ct,
        pie_labels=pl,
        pie_values=pv,
        missing_dates=find_missing_dates(30),
        audit_logs=get_recent_logs(25),
        tg_users=users_rows(),
        employees=get_branch_employees(),
        employee_candidates=get_employee_candidates(),
        pending_reports=get_pending_branch_reports(),
        latest_pending_id=latest_pending_report_id(),
        branch_daily_status=get_branch_daily_status(),
        closed_workdays=get_recent_closed_workdays(),
        admin_id=ADMIN_ID,
        footer=FIXED_FOOTER,
    )


@flask_app.route("/save_report", methods=["POST"])
def web_save_report():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    ps = request.form.get("json_payload")
    if not ps:
        return redirect(url_for("web_dashboard"))
    try:
        p = json.loads(ps)
        d = parse_smart_date(p["date_input"])
    except Exception:
        flash("❌ تاریخ واردشده معتبر نیست.", "danger")
        return redirect(url_for("web_dashboard"))
    vb = []
    for b in p.get("branches", []):
        if not str(b.get("name", "")).strip():
            continue
        rngs = [{"start": tr_num(r.get("start", "0")), "end": tr_num(r.get("end", "0"))} for r in b.get("ranges", [])]
        vb.append({
            "name": str(b.get("name")).strip(),
            "issued": safe_int(b.get("issued", 0)) or 0,
            "ranges": rngs,
            "start": rngs[0]["start"] if rngs else "0",
            "end": rngs[-1]["end"] if rngs else "0",
            "cancelled": safe_int(b.get("cancelled", 0)) or 0,
            "cancelled_serial": tr_num(b.get("cancelled_serial", "0")),
        })
    dt = {
        "date": d.strftime("%d/%m/%Y"),
        "shamsi_date": shamsi(d),
        "branches": vb,
        "transit": safe_int(p.get("transit", 0)) or 0,
        "export": safe_int(p.get("export", 0)) or 0,
        "created_by": ADMIN_ID,
    }
    closed = get_closed_workday(dt["date"])
    if closed and closed.get("status") == "closed":
        flash("🔒 این روز نهایی و بسته شده است. ابتدا از بخش کنترل روز کاری، با ذکر دلیل آن را بازگشایی کنید.", "warning")
        return redirect(url_for("web_dashboard") + "#tab-employees")
    previous = load_report(dt["date"])
    reason = str(request.form.get("change_reason", "")).strip()[:500]
    if ((previous and previous != dt) or closed) and not reason:
        flash("برای تغییر گزارش موجود یا تاریخی که قبلاً نهایی شده، دلیل اصلاح را در قسمت دلیل تغییر وارد کنید.", "warning")
        return redirect(url_for("web_dashboard") + "#tab-employees")
    save_report(dt)
    if previous and previous != dt:
        log_report_change(dt["date"], "اصلاح گزارش از پنل", reason, ADMIN_ID)
    else:
        log_action(ADMIN_ID, f"Saved report via Web Panel for {dt['date']}")
    if p.get("send_tg") and GLOBAL_BOT_APP and GLOBAL_LOOP:
        asyncio.run_coroutine_threadsafe(notify_admin_from_web(dt), GLOBAL_LOOP)
    flash(f"✅ گزارش تاریخ {dt['shamsi_date']} با موفقیت ذخیره شد.", "success")
    return redirect(url_for("web_dashboard"))


@flask_app.route("/send_to_tg", methods=["POST"])
def web_send_to_tg():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    ds = request.form.get("date", "")
    d = load_report(ds)
    if d and GLOBAL_BOT_APP and GLOBAL_LOOP:
        asyncio.run_coroutine_threadsafe(notify_admin_from_web(d), GLOBAL_LOOP)
        flash(f"✈️ گزارش تاریخ {d.get('shamsi_date')} به تلگرام شما ارسال شد.", "info")
    return redirect(url_for("web_dashboard"))


@flask_app.route("/toggle_user_web", methods=["POST"])
def web_toggle_user():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    uid = safe_int(request.form.get("uid", ""))
    allow = request.form.get("allow", "1") == "1"
    if uid and uid != ADMIN_ID:
        set_allowed(uid, allow)
        flash(f"✅ وضعیت دسترسی کاربر {uid} تغییر کرد.", "info")
    return redirect(url_for("web_dashboard"))


@flask_app.route("/broadcast_web", methods=["POST"])
def web_broadcast():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    msg_text = request.form.get("message", "").strip()
    if msg_text and GLOBAL_BOT_APP and GLOBAL_LOOP:
        asyncio.run_coroutine_threadsafe(broadcast_from_web_async(msg_text), GLOBAL_LOOP)
        flash("📢 پیام همگانی به تلگرام ارسال شد.", "success")
    return redirect(url_for("web_dashboard"))


async def broadcast_from_web_async(text):
    conn = db()
    us = conn.execute("SELECT user_id FROM users WHERE allowed=1").fetchall()
    conn.close()
    for (ui,) in us:
        try:
            await GLOBAL_BOT_APP.bot.send_message(chat_id=ui, text=f"📢 اطلاعیه مدیریت سامانه مرز باشماق:\n\n{text}")
            await asyncio.sleep(0.05)
        except Exception:
            pass


@flask_app.route("/restore_db", methods=["POST"])
def web_restore_db():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    f = request.files.get("db_file")
    if not f or not (f.filename or "").lower().endswith(".db"):
        flash("❌ فایل .db انتخاب نشده است.", "danger")
        return redirect(url_for("web_dashboard"))
    tp = os.path.join(OUT_DIR, f"web_restore_{secrets.token_hex(4)}.db")
    f.save(tp)
    try:
        chk = sqlite3.connect(tp)
        try:
            integrity = chk.execute("PRAGMA integrity_check").fetchone()
            tables = {r[0] for r in chk.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            report_cols = {r[1] for r in chk.execute("PRAGMA table_info(reports)")} if "reports" in tables else set()
        finally:
            chk.close()
        if not integrity or integrity[0] != "ok":
            raise ValueError("بررسی سلامت فایل ناموفق بود")
        if "reports" not in tables or not {"date", "data", "created_at"}.issubset(report_cols):
            raise ValueError("ساختار جدول گزارش‌ها با این سامانه سازگار نیست")
        # پیش از بازیابی، بکاپ سازگار SQLite تهیه می‌شود؛ کپی ساده‌ی فایل ممکن است WAL را جا بیندازد.
        if os.path.exists(DB_FILE):
            sd = os.path.join(BASE_DIR, "backups")
            os.makedirs(sd, exist_ok=True)
            backup_path = os.path.join(sd, f"before_restore_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db")
            src_conn = sqlite3.connect(DB_FILE, timeout=20)
            dst_conn = sqlite3.connect(backup_path, timeout=20)
            try:
                src_conn.backup(dst_conn)
            finally:
                dst_conn.close()
                src_conn.close()
            verify = sqlite3.connect(backup_path, timeout=20)
            try:
                backup_check = verify.execute("PRAGMA integrity_check").fetchone()
                if not backup_check or backup_check[0] != "ok":
                    raise ValueError("تهیه نسخه پشتیبان از دیتابیس فعلی تأیید نشد")
            finally:
                verify.close()
        _replace_db_file(tp)
        init_db()
        log_action(ADMIN_ID, "Restored database via Web Panel")
        flash("✅ دیتابیس بازیابی شد؛ نسخه قبلی در پوشه backups نگهداری شد.", "success")
    except Exception as ex:
        safe_remove(tp)
        flash(f"❌ بازیابی انجام نشد: {str(ex)[:200]}", "danger")
    return redirect(url_for("web_dashboard"))


async def notify_admin_from_web(d):
    try:
        bot = GLOBAL_BOT_APP.bot
        await bot.send_message(chat_id=ADMIN_ID, text=f"🌐 گزارش ارسالی از پنل وب:\n\n{create_branch_report(d)}")
        jpg = create_jpg_report(d)
        with open(jpg, "rb") as f:
            await bot.send_photo(chat_id=ADMIN_ID, photo=f, caption="🖼️ تصویر رسمی گزارش", reply_markup=whatsapp_share_kb(d))
        os.remove(jpg)
        pdf = create_pdf_report(d)
        with open(pdf, "rb") as f:
            await bot.send_document(chat_id=ADMIN_ID, document=f, caption="📄 فایل PDF رسمی گزارش")
        os.remove(pdf)
        await send_auto_backup_to_admin(bot, reason=f"بکاپ خودکار وب ({d['date']})")
    except Exception:
        logger.exception("notify_admin_from_web failed")


@flask_app.route("/download/<fmt>")
def web_download_single(fmt):
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    ds = request.args.get("date", "")
    d = load_report(ds)
    if not d:
        return redirect(url_for("web_dashboard"))
    if fmt == "jpg":
        return send_file(create_jpg_report(d), as_attachment=True, download_name=f"bashmaq_{ds.replace('/','-')}.jpg")
    elif fmt == "pdf":
        return send_file(create_pdf_report(d), as_attachment=True, download_name=f"bashmaq_{ds.replace('/','-')}.pdf")
    return redirect(url_for("web_dashboard"))


@flask_app.route("/export_range")
def web_export_range():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    try:
        s = parse_smart_date(request.args.get("start", ""))
        e = parse_smart_date(request.args.get("end", ""))
    except Exception:
        return redirect(url_for("web_dashboard"))
    fmt = request.args.get("fmt", "excel")
    if fmt == "pdf":
        return send_file(create_range_pdf(s, e), as_attachment=True, download_name="bashmaq_range.pdf")
    elif fmt == "jpg":
        return send_file(create_range_jpg(s, e), as_attachment=True, download_name="bashmaq_range.jpg")
    elif fmt == "chart":
        cp = create_chart_report(s, e)
        if cp:
            return send_file(cp, as_attachment=True, download_name="bashmaq_chart.png")
        return redirect(url_for("web_dashboard"))
    return send_file(create_excel_report(s, e), as_attachment=True, download_name="bashmaq_range.xlsx")


@flask_app.route("/delete_report", methods=["POST"])
def web_delete_report():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    ds = request.form.get("date", "")
    _cw = get_closed_workday(ds)
    if _cw and _cw.get("status") == "closed":
        flash("🔒 روز بسته‌شده قابل حذف نیست؛ ابتدا با ذکر دلیل از بخش کنترل روز کاری بازگشایی کنید.", "warning")
    elif delete_report(ds):
        flash("🗑 گزارش حذف شد.", "warning")
        log_action(ADMIN_ID, f"Deleted report {ds} via Web Panel")
    return redirect(url_for("web_dashboard"))


@flask_app.route("/backup_db")
def web_backup_db():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    backup_dir = os.path.join(BASE_DIR, "backups")
    os.makedirs(backup_dir, exist_ok=True)
    backup_path = os.path.join(backup_dir, f"bashmaq_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db")
    src = sqlite3.connect(DB_FILE, timeout=20); dst = sqlite3.connect(backup_path)
    try: src.backup(dst)
    finally: dst.close(); src.close()
    check = sqlite3.connect(backup_path)
    try:
        integrity = check.execute("PRAGMA integrity_check").fetchone()
        if not integrity or integrity[0] != "ok":
            check.close()
            os.remove(backup_path)
            flash("❌ بررسی سلامت نسخه پشتیبان ناموفق بود؛ فایل ارائه نشد.", "danger")
            return redirect(url_for("web_dashboard"))
    finally:
        try: check.close()
        except Exception: pass
    # keep the newest 30 backups to avoid filling the disk
    old_backups = sorted([os.path.join(backup_dir, n) for n in os.listdir(backup_dir) if n.endswith('.db')], key=os.path.getmtime, reverse=True)
    for stale in old_backups[30:]:
        try: os.remove(stale)
        except OSError: pass
    log_action(ADMIN_ID, "Created database backup")
    return send_file(backup_path, as_attachment=True, download_name=os.path.basename(backup_path))


@flask_app.route("/api/global_search")
def api_global_search():
    if not session.get("logged_in"): return jsonify({"results": []}), 401
    q = request.args.get("q", "").strip().casefold()
    if len(q) < 2: return jsonify({"results": []})
    results=[]
    for rep in reversed(all_reports()):
        sh=rep.get("shamsi_date", "-"); gd=rep.get("date", "-")
        for b in rep.get("branches", []):
            rng=get_branch_ranges_text(b)
            hay=f"{sh} {gd} {b.get('name','')} {rng} {b.get('issued','')} {b.get('cancelled','')} {rep.get('transit','')} {rep.get('export','')}".casefold()
            if q in hay:
                results.append({"shamsi":sh,"date":gd,"branch":str(b.get("name","-")),"detail":f"سریال {rng} | صادره {b.get('issued',0)}"})
                if len(results)>=50: break
        if len(results)>=50: break
    return jsonify({"results":results})


@flask_app.route("/api/smart_insights")
def api_smart_insights():
    if not session.get("logged_in"): return jsonify({"insights": []}), 401
    reps=all_reports(); insights=[]
    total=sum(total_issued(r) for r in reps)
    insights.append({"title":"نمای کلی عملکرد", "text":f"{len(reps)} روز گزارش ثبت شده و مجموع صادره‌ها {total:,} فقره است."})
    if reps:
        best=max(reps,key=total_issued)
        insights.append({"title":"بیشترین عملکرد روزانه", "text":f"{best.get('shamsi_date',best.get('date','-'))} با {total_issued(best):,} فقره صادره، بیشترین عملکرد را داشته است."})
        today=iran_today(); month_start=date(today.year,today.month,1)
        recent=[]
        for r in reps:
            try:
                dt=datetime.strptime(r.get('date',''),'%d/%m/%Y').date()
                if dt>=month_start: recent.append(r)
            except Exception: pass
        insights.append({"title":"عملکرد ماه جاری", "text":f"در این ماه {len(recent)} گزارش و {sum(total_issued(r) for r in recent):,} فقره صادره ثبت شده است."})
    gaps=find_missing_dates(30)
    insights.append({"title":"کنترل روزهای جامانده", "text":f"{len(gaps)} روز بدون گزارش در بازه بررسی ۳۰ روزه شناسایی شد." if gaps else "در بازه ۳۰ روزه روز جامانده‌ای شناسایی نشد."})
    # quick overlapping serial check
    intervals=[]
    for r in reps:
        for b in r.get('branches',[]):
            for rg in b.get('ranges',[{'start':b.get('start','0'),'end':b.get('end','0')} ]):
                try:
                    a,z=int(tr_num(rg.get('start','0'))),int(tr_num(rg.get('end','0')))
                    if a>0 and z>=a: intervals.append((a,z,r.get('shamsi_date','-'),b.get('name','-')))
                except Exception: pass
    intervals.sort(); overlaps=sum(1 for i in range(1,len(intervals)) if intervals[i][0]<=intervals[i-1][1])
    insights.append({"title":"سلامت سریال‌ها", "text":f"در بررسی سریع {overlaps} مورد هم‌پوشانی احتمالی بازه یافت شد؛ برای جزئیات از بازرس سریال استفاده کنید." if overlaps else "در بررسی سریع بازه‌های سریال، هم‌پوشانی مجاوری پیدا نشد."})
    insights.extend(extra_smart_insights())
    return jsonify({"insights":insights})


@flask_app.route("/api/audit_serials")
def api_audit_serials():
    if not session.get("logged_in"): return jsonify({"issues": []}), 401
    items=[]
    for rep in all_reports():
        for b in rep.get('branches', []):
            for rg in b.get('ranges', [{'start':b.get('start','0'),'end':b.get('end','0')} ]):
                try:
                    st,en=int(tr_num(rg.get('start','0'))),int(tr_num(rg.get('end','0')))
                    if st>0 and en>=st: items.append((st,en,rep.get('date','?'),b.get('name','?')))
                    elif st>0 or en>0: pass
                except Exception: pass
    items.sort()
    issues=[]; prev=None
    for cur in items:
        if prev and cur[0] <= prev[1]:
            issues.append(f'تداخل سریال {cur[0]} الی {min(cur[1],prev[1])}: شعبه {prev[3]} ({prev[2]}) با شعبه {cur[3]} ({cur[2]})')
        elif prev and cur[0] > prev[1]+1:
            issues.append(f'پرش احتمالی از {prev[1]} تا {cur[0]} بین {prev[3]} ({prev[2]}) و {cur[3]} ({cur[2]})')
        if not prev or cur[1] > prev[1]: prev=cur
    return jsonify({"issues": issues[:300]})


@flask_app.route("/api/search_serial")
def api_search_serial():
    if not session.get("logged_in"):
        return jsonify({"results": []})
    tgt = safe_int(request.args.get("serial", ""))
    res = []
    if not tgt:
        return jsonify({"results": []})
    for r in all_reports():
        for b in r.get("branches", []):
            rngs = b.get("ranges", [{"start": b.get("start", "0"), "end": b.get("end", "0")}])
            isc = str(tgt) in [x.strip() for x in tr_num(b.get("cancelled_serial", "")).replace("،", ",").split(",")]
            fnd = False
            for rn in rngs:
                try:
                    st, en = int(tr_num(rn["start"])), int(tr_num(rn["end"]))
                    if st <= tgt <= en and st > 0:
                        fnd = True
                        break
                except Exception:
                    continue
            if fnd or isc:
                res.append({"date": r["date"], "shamsi": r.get("shamsi_date", "-"), "branch": b.get("name", "-"), "range": get_branch_ranges_text(b), "status": "❌ ابطالی" if isc else "✅ معتبر"})
    return jsonify({"results": res})


def run_flask_server():
    try:
        from waitress import serve
        logger.info("Web server: waitress on port %s", PORT)
        serve(flask_app, host="0.0.0.0", port=PORT, threads=8)
    except ImportError:
        logger.warning("waitress نصب نیست؛ سرور توسعه Flask اجرا می‌شود (پیشنهاد: pip install waitress)")
        flask_app.run(host="0.0.0.0", port=PORT, use_reloader=False, threaded=True)


# ============================================================
# بخش دوم: ربات تلگرام
# ============================================================
async def send_auto_backup_to_admin(bot, reason="بکاپ"):
    if not os.path.exists(DB_FILE):
        try:
            await bot.send_message(chat_id=ADMIN_ID, text="🚨 هشدار پشتیبان‌گیری: فایل دیتابیس پیدا نشد.")
        except Exception:
            pass
        return False
    backup_dir = os.path.join(BASE_DIR, "backups")
    os.makedirs(backup_dir, exist_ok=True)
    path = os.path.join(backup_dir, f"bashmaq_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db")
    verified = False
    try:
        src = sqlite3.connect(DB_FILE, timeout=20)
        dst = sqlite3.connect(path)
        try:
            src.backup(dst)
        finally:
            dst.close()
            src.close()
        check = sqlite3.connect(path)
        try:
            result = check.execute("PRAGMA integrity_check").fetchone()
            if not result or result[0] != "ok":
                raise RuntimeError(f"integrity_check failed: {result}")
            table_count = check.execute("SELECT COUNT(*) FROM sqlite_master WHERE type='table'").fetchone()[0]
            verified = True
        finally:
            check.close()
        # Retain the latest 30 verified backups locally.
        old_backups = sorted([os.path.join(backup_dir, n) for n in os.listdir(backup_dir) if n.endswith('.db')], key=os.path.getmtime, reverse=True)
        for stale in old_backups[30:]:
            try: os.remove(stale)
            except OSError: pass
        with open(path, "rb") as f:
            await bot.send_document(chat_id=ADMIN_ID, document=f, filename=os.path.basename(path), caption=f"💾 {reason} — نسخه بررسی‌شده\n📅 {shamsi(iran_today())}\n✅ سلامت دیتابیس: OK | جدول‌ها: {table_count}")
        log_action(ADMIN_ID, f"Verified backup created: {os.path.basename(path)} ({reason})")
        return True
    except Exception as ex:
        logger.exception("Backup verification/send failed")
        try:
            await bot.send_message(chat_id=ADMIN_ID, text=f"🚨 هشدار مهم: پشتیبان‌گیری یا بررسی سلامت دیتابیس ناموفق بود.\nعلت: {str(ex)[:500]}")
        except Exception:
            pass
        if os.path.exists(path) and not verified:
            try: os.remove(path)
            except OSError: pass
        return False


async def nightly_backup_job(ctx: ContextTypes.DEFAULT_TYPE):
    await send_auto_backup_to_admin(ctx.bot, "بکاپ شبانه")


async def nightly_reminder_job(ctx: ContextTypes.DEFAULT_TYPE):
    d = iran_today()
    ds = d.strftime("%d/%m/%Y")
    if not load_report(ds):
        try:
            await ctx.bot.send_message(chat_id=ADMIN_ID, text=f"⏰ گزارش امروز (`{shamsi(d)}`) ثبت نشده است.", parse_mode="Markdown")
        except Exception:
            pass


async def start(u, c):
    c.user_data.clear()
    save_user(u.effective_user)
    if not can_use(u):
        return ConversationHandler.END
    await u.message.reply_text("سلام 🌹\nبه سامانه هوشمند مرز باشماق (نسخه رویال ۲۰۲۶) خوش آمدید.", reply_markup=get_menu(u))
    return ConversationHandler.END


async def send_monthly_financial_files(u, c):
    if not can_use(u): return
    now = iran_today(); month_key = f"{now.year:04d}-{now.month:02d}"
    try:
        pdf = create_accounting_pdf_file(now.year, now.month)
        xlsx = create_accounting_excel_file(now.year, now.month)
        with open(pdf, 'rb') as f: await u.message.reply_document(f, filename=os.path.basename(pdf), caption=f'📄 PDF صورت‌حساب مالی {month_key}')
        with open(xlsx, 'rb') as f: await u.message.reply_document(f, filename=os.path.basename(xlsx), caption=f'📊 اکسل صورت‌حساب مالی {month_key}')
    except Exception as e:
        logger.exception('Monthly financial files failed: %s', e)
        await u.message.reply_text('❌ تهیه فایل‌های مالی ناموفق بود. ابتدا جدول ماه جاری را در سایت ذخیره کنید.')


async def my_id(u, c):
    await u.message.reply_text(f"`{u.effective_user.id}`", parse_mode="Markdown")


async def cancel(u, c):
    c.user_data.clear()
    if u.callback_query:
        await u.callback_query.answer()
        await u.callback_query.message.reply_text("❌ لغو شد.", reply_markup=get_menu(u))
    else:
        await u.message.reply_text("❌ لغو شد.", reply_markup=get_menu(u))
    return ConversationHandler.END


async def back(u, c):
    c.user_data.clear()
    await u.message.reply_text("🏠 منوی اصلی", reply_markup=get_menu(u))


async def error_handler(u, c):
    err = c.error
    logger.error("Telegram error:", exc_info=err)
    key = type(err).__name__
    now = time_module.time()
    if now - _ERROR_LAST.get(key, 0) > 300:
        _ERROR_LAST[key] = now
        tb = "".join(traceback.format_exception(type(err), err, err.__traceback__))[-700:]
        try:
            await c.bot.send_message(chat_id=ADMIN_ID, text=f"🚨 خطای ربات: {key}\n{tb}")
        except Exception:
            pass
    try:
        if isinstance(u, Update) and u.effective_message:
            await u.effective_message.reply_text("⚠️ خطایی رخ داد؛ دوباره تلاش کنید یا /start را بزنید.")
    except Exception:
        pass


def branch_Count_kb():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("2 شعبه", callback_data="bcount:2"), InlineKeyboardButton("3 شعبه", callback_data="bcount:3")],
        [InlineKeyboardButton("1 شعبه", callback_data="bcount:1"), InlineKeyboardButton("4 شعبه", callback_data="bcount:4")],
    ])


def branch_name_kb():
    rws, rw = [], []
    for nm in get_known_branches():
        rw.append(InlineKeyboardButton(f"🏢 {nm}", callback_data=f"bname:{nm}"))
        if len(rw) == 2:
            rws.append(rw)
            rw = []
    if rw:
        rws.append(rw)
    return InlineKeyboardMarkup(rws) if rws else None


async def new_report(u, c):
    save_user(u.effective_user)
    if not is_admin(u):
        return ConversationHandler.END
    c.user_data.clear()
    c.user_data["branches"] = []
    c.user_data["mode"] = "new"
    kb = InlineKeyboardMarkup([[InlineKeyboardButton("📅 امروز", callback_data="date_today"), InlineKeyboardButton("⏪ دیروز", callback_data="date_yesterday")]])
    await u.message.reply_text("📅 تاریخ گزارش را وارد کنید:", reply_markup=kb)
    return DATE


async def edit_start(u, c):
    if not is_admin(u):
        return ConversationHandler.END
    await u.message.reply_text("📅 تاریخ گزارش جهت ویرایش:", reply_markup=ReplyKeyboardRemove())
    return EDIT_DATE


async def get_edit_date(u, c):
    try:
        ds = norm_date(u.message.text)
    except ValueError:
        await u.message.reply_text("❌ تاریخ اشتباه است.")
        return EDIT_DATE
    old = load_report(ds)
    if not old:
        await u.message.reply_text("❌ گزارشی یافت نشد.", reply_markup=get_menu(u))
        return ConversationHandler.END
    c.user_data.clear()
    c.user_data.update(old)
    c.user_data["date"] = ds
    c.user_data["shamsi_date"] = shamsi(parse_smart_date(ds))
    kb = InlineKeyboardMarkup([
        [InlineKeyboardButton("🚛 فقط ترانزیت و صادرات", callback_data="ed_trade")],
        [InlineKeyboardButton("🏢 ویرایش کامل", callback_data="ed_full")],
        [InlineKeyboardButton("❌ انصراف", callback_data="ed_cancel")],
    ])
    await u.message.reply_text(f"📋 ویرایش گزارش {ds}:", reply_markup=kb)
    return EDIT_CHOICE


async def edit_choice_callback(u, c):
    q = u.callback_query
    await q.answer()
    if q.data == "ed_cancel":
        c.user_data.clear()
        await q.edit_message_text("❌ لغو شد.")
        await q.message.reply_text("🏠", reply_markup=admin_menu())
        return ConversationHandler.END
    if q.data == "ed_trade":
        await q.edit_message_text("🚛 ویرایش ترانزیت:")
        await q.message.reply_text("تعداد ترانزیت جدید:")
        return TRANSIT
    c.user_data["branches"] = []
    await q.edit_message_text("🏢 ویرایش کامل:")
    await q.message.reply_text("تعداد شعب:", reply_markup=branch_Count_kb())
    return COUNT


async def date_quick_callback(u, c):
    q = u.callback_query
    await q.answer()
    d = iran_today() if q.data == "date_today" else iran_today() - timedelta(days=1)
    return await process_selected_date(q.message, c, d)


async def get_date(u, c):
    try:
        d = parse_smart_date(u.message.text)
    except ValueError:
        await u.message.reply_text("❌ تاریخ اشتباه است.")
        return DATE
    return await process_selected_date(u.message, c, d)


async def process_selected_date(m, c, d):
    ds = d.strftime("%d/%m/%Y")
    c.user_data["date"] = ds
    c.user_data["shamsi_date"] = shamsi(d)
    old = load_report(ds)
    if old:
        kb = InlineKeyboardMarkup([[InlineKeyboardButton("✏️ ویرایش", callback_data="dup_replace")], [InlineKeyboardButton("❌ لغو", callback_data="dup_cancel")]])
        await m.reply_text(f"⚠️ برای {ds} قبلاً گزارش ثبت شده:", reply_markup=kb)
        return DATE
    await m.reply_text(f"✅ تاریخ: {ds}\nتعداد شعبه‌ها:", reply_markup=branch_Count_kb())
    return COUNT


async def duplicate_callback(u, c):
    q = u.callback_query
    await q.answer()
    if q.data == "dup_cancel":
        c.user_data.clear()
        await q.edit_message_text("❌ لغو شد.")
        await q.message.reply_text("🏠", reply_markup=admin_menu())
        return ConversationHandler.END
    old = load_report(c.user_data["date"])
    if old:
        c.user_data.update(old)
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton("🚛 ترانزیت/صادرات", callback_data="ed_trade")],
            [InlineKeyboardButton("🏢 ویرایش کامل", callback_data="ed_full")],
            [InlineKeyboardButton("❌ انصراف", callback_data="ed_cancel")],
        ])
        await q.edit_message_text("کدام بخش ویرایش شود؟", reply_markup=kb)
        return EDIT_CHOICE
    await q.edit_message_text("تعداد شعبه‌ها:")
    await q.message.reply_text("تعداد:", reply_markup=branch_Count_kb())
    return COUNT


async def count_callback(u, c):
    q = u.callback_query
    await q.answer()
    c.user_data["count"] = int(q.data.split(":")[1])
    c.user_data["current_branch"] = 1
    await q.edit_message_text(f"تعداد: {c.user_data['count']}")
    await q.message.reply_text("نام شعبه 1:", reply_markup=branch_name_kb())
    return NAME


async def get_count(u, c):
    n = safe_int(u.message.text)
    if not n or not (1 <= n <= 10):
        await u.message.reply_text("❌ عدد ۱ تا ۱۰")
        return COUNT
    c.user_data["count"] = n
    c.user_data["current_branch"] = 1
    await u.message.reply_text("نام شعبه 1:", reply_markup=branch_name_kb())
    return NAME


async def name_callback(u, c):
    q = u.callback_query
    await q.answer()
    name = q.data.split(":", 1)[1]
    c.user_data["branch_name"] = name
    await q.edit_message_text(f"شعبه: {name}")
    await q.message.reply_text(f"کل تعداد صادره «{name}»:")
    return ISSUED


async def get_name(u, c):
    c.user_data["branch_name"] = u.message.text.strip()
    await u.message.reply_text(f"کل تعداد صادره «{c.user_data['branch_name']}»:")
    return ISSUED


async def get_issued(u, c):
    n = safe_int(u.message.text)
    if n is None:
        await u.message.reply_text("❌ فقط عدد")
        return ISSUED
    c.user_data["issued"] = n
    c.user_data["current_ranges"] = []
    c.user_data["collected_issued"] = 0
    if n == 0:
        c.user_data["current_ranges"] = [{"start": "0", "end": "0"}]
        kb = InlineKeyboardMarkup([[InlineKeyboardButton("0️⃣ بدون ابطالی", callback_data="canc_zero")]])
        await u.message.reply_text("صادره 0. ابطالی؟", reply_markup=kb)
        return CANCELLED
    le, _ = get_last_end_serial(c.user_data.get("branch_name", ""), before_date_str=c.user_data.get("date"))
    if le:
        kb = InlineKeyboardMarkup([[InlineKeyboardButton(f"⚡ ثبت سریع {le+1} الی {le+n}", callback_data=f"autoserial:{le+1}:{le+n}")]])
        await u.message.reply_text(f"سریال شروع **بازه اول**؟ (پیشنهاد از {le+1})", reply_markup=kb, parse_mode="Markdown")
        return START
    await u.message.reply_text("سریال شروع **بازه اول**؟", parse_mode="Markdown")
    return START


async def auto_serial_callback(u, c):
    q = u.callback_query
    await q.answer()
    _, st, en = q.data.split(":")
    c.user_data["current_ranges"] = [{"start": st, "end": en}]
    c.user_data["range_text"] = f"{st} الی {en}"
    await q.edit_message_text(f"✅ بازه {st} الی {en} ثبت شد.")
    kb = InlineKeyboardMarkup([[InlineKeyboardButton("0️⃣ بدون ابطالی", callback_data="canc_zero")]])
    await q.message.reply_text("تعداد ابطالی؟", reply_markup=kb)
    return CANCELLED


async def get_start(u, c):
    s = tr_num(u.message.text)
    if not s:
        return START
    c.user_data["temp_start"] = s
    rem = c.user_data["issued"] - c.user_data.get("collected_issued", 0)
    if s.isdigit() and rem > 0:
        se = int(s) + rem - 1
        kb = InlineKeyboardMarkup([[InlineKeyboardButton(f"✅ پایان: {se}", callback_data=f"endserial:{se}")]])
        await u.message.reply_text("سریال پایان بازه را وارد یا انتخاب کنید:", reply_markup=kb)
        return END
    await u.message.reply_text("سریال پایان این بازه:")
    return END


async def end_serial_callback(u, c):
    q = u.callback_query
    await q.answer()
    return await process_end_serial(q.message, c, q.data.split(":")[1], q)


async def get_end(u, c):
    return await process_end_serial(u.message, c, tr_num(u.message.text))


async def process_end_serial(msg, c, eval_str, q=None):
    sval = c.user_data["temp_start"]
    iss = c.user_data["issued"]
    col = c.user_data.get("collected_issued", 0)
    rngs = c.user_data.get("current_ranges", [])
    try:
        chk = int(eval_str) - int(sval) + 1
    except Exception:
        await msg.reply_text("❌ نامعتبر.")
        return END
    if chk <= 0:
        await msg.reply_text("❌ پایان بزرگتر از شروع باشد.")
        return END

    if col + chk < iss:
        if len(rngs) >= 3:
            await msg.reply_text("❌ حداکثر ۴ بازه! ثبت را تکرار کنید:\n\nسریال شروع:", parse_mode="Markdown")
            c.user_data["current_ranges"] = []
            c.user_data["collected_issued"] = 0
            return START
        c.user_data["current_ranges"].append({"start": sval, "end": eval_str})
        c.user_data["collected_issued"] = col + chk
        if q:
            await q.edit_message_text(f"✅ بازه {sval} تا {eval_str} ثبت شد.")
        await msg.reply_text(f"⚠️ تا الان {col + chk} فقره ثبت شد. {iss - (col + chk)} فقره باقی مانده.\n\n**سریال شروع بازه بعدی** را وارد کنید:", parse_mode="Markdown")
        return START
    elif col + chk == iss:
        if len(rngs) >= 4:
            await msg.reply_text("❌ حداکثر ۴ بازه! ثبت را تکرار کنید:\n\nسریال شروع:", parse_mode="Markdown")
            c.user_data["current_ranges"] = []
            c.user_data["collected_issued"] = 0
            return START
        c.user_data["current_ranges"].append({"start": sval, "end": eval_str})
        rs = " و ".join([f"{r['start']} الی {r['end']}" for r in c.user_data["current_ranges"]])
        c.user_data["range_text"] = rs
        if q:
            await q.edit_message_text(f"✅ بازه‌ها تکمیل شد: {rs}")
        await msg.reply_text("✅ مجموع بازه‌ها کامل شد.\nتعداد ابطالی؟", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("0️⃣ بدون ابطالی", callback_data="canc_zero")]]))
        return CANCELLED
    else:
        await msg.reply_text(f"❌ تعداد ({chk}) بیشتر از باقیمانده ({iss - col}) است! پایان را درست وارد کنید:")
        return END


async def cancelled_zero_callback(u, c):
    q = u.callback_query
    await q.answer()
    c.user_data["cancelled"] = 0
    c.user_data["cancelled_serial"] = "0"
    await q.edit_message_text("✅ ابطالی: 0")
    return await save_branch_bot(q.message, c)


async def get_cancelled(u, c):
    n = safe_int(u.message.text)
    if n is None:
        return CANCELLED
    c.user_data["cancelled"] = n
    if n > 0:
        await u.message.reply_text("سریال ابطالی:")
        return CANCELLED_SERIAL
    c.user_data["cancelled_serial"] = "0"
    return await save_branch_bot(u.message, c)


async def get_cancelled_serial(u, c):
    c.user_data["cancelled_serial"] = tr_num(u.message.text)
    return await save_branch_bot(u.message, c)


async def save_branch_bot(m, c):
    rngs = c.user_data.get("current_ranges", [])
    b = {
        "name": c.user_data["branch_name"],
        "issued": c.user_data["issued"],
        "ranges": rngs,
        "start": rngs[0]["start"] if rngs else "0",
        "end": rngs[-1]["end"] if rngs else "0",
        "range_text": c.user_data.get("range_text", "0"),
        "cancelled": c.user_data["cancelled"],
        "cancelled_serial": c.user_data["cancelled_serial"],
    }
    c.user_data["branches"].append(b)
    cur, tot = c.user_data["current_branch"], c.user_data["count"]
    if cur < tot:
        c.user_data["current_branch"] += 1
        await m.reply_text(f"نام شعبه {cur + 1}:", reply_markup=branch_name_kb())
        return NAME
    await m.reply_text("✅ اطلاعات شعب تکمیل شد.\n🚛 تعداد ترانزیت:")
    return TRANSIT


async def get_transit(u, c):
    v = safe_int(u.message.text)
    if v is None:
        await u.message.reply_text("❌ فقط عدد وارد کنید (مثلاً 0 یا 12):")
        return TRANSIT
    c.user_data["transit"] = v
    await u.message.reply_text("📦 تعداد صادرات:")
    return EXPORT


async def get_export(u, c):
    v = safe_int(u.message.text)
    if v is None:
        await u.message.reply_text("❌ فقط عدد وارد کنید (مثلاً 0 یا 12):")
        return EXPORT
    c.user_data["export"] = v
    c.user_data["created_by"] = u.effective_user.id
    prv = f"پیش‌نمایش:\n\n{create_branch_report(c.user_data)}\n\n{create_trade_report(c.user_data)}\n\nتأیید؟"
    kb = InlineKeyboardMarkup([[InlineKeyboardButton("✅ ذخیره", callback_data="conf_save")], [InlineKeyboardButton("❌ لغو", callback_data="conf_cancel")]])
    await u.message.reply_text(prv, reply_markup=kb)
    return CONFIRM_SAVE


async def fast_report_start(u, c):
    if not is_admin(u):
        return ConversationHandler.END
    ds = iran_today().strftime("%d/%m/%Y")
    lt, _ = get_last_end_serial("تیرپارک", ds)
    lp, _ = get_last_end_serial("پارسیان", ds)
    await u.message.reply_text(
        f"⚡ ثبت فوق‌سریع:\nشروع تیرپارک: {(lt+1) if lt else 1}\nشروع پارسیان: {(lp+1) if lp else 1}\n\n۴ عدد با فاصله بفرستید (مثال: `27 85 33 79`)",
        reply_markup=ReplyKeyboardRemove(),
        parse_mode="Markdown",
    )
    return FAST_REPORT_STATE


async def fast_report_receive(u, c):
    ls = [x.strip() for x in tr_num(u.message.text).splitlines() if x.strip()]
    if not ls:
        return FAST_REPORT_STATE
    try:
        if len(ls) == 1:
            d = iran_today()
            nums = [int(x) for x in ls[0].replace("-", " ").split() if x.isdigit()]
        else:
            d = parse_smart_date(ls[0])
            nums = [int(x) for x in ls[1].replace("-", " ").split() if x.isdigit()]
        if len(nums) != 4:
            raise ValueError()
    except Exception:
        await u.message.reply_text("❌ فرمت اشتباه است.")
        return FAST_REPORT_STATE
    it, ip, tr, ex = nums
    ds, sh_d = d.strftime("%d/%m/%Y"), shamsi(d)
    lt, _ = get_last_end_serial("تیرپارک", ds)
    lp, _ = get_last_end_serial("پارسیان", ds)
    stt, ent = (lt + 1 if lt and it > 0 else 1), (lt + it if lt and it > 0 else (it if it > 0 else 0))
    stp, enp = (lp + 1 if lp and ip > 0 else 1), (lp + ip if lp and ip > 0 else (ip if ip > 0 else 0))
    bl = []
    if it > 0:
        bl.append({"name": "تیرپارک", "issued": it, "ranges": [{"start": str(stt), "end": str(ent)}], "start": str(stt), "end": str(ent), "range_text": f"{stt} الی {ent}", "cancelled": 0, "cancelled_serial": "0"})
    if ip > 0:
        bl.append({"name": "پارسیان", "issued": ip, "ranges": [{"start": str(stp), "end": str(enp)}], "start": str(stp), "end": str(enp), "range_text": f"{stp} الی {enp}", "cancelled": 0, "cancelled_serial": "0"})
    c.user_data.clear()
    c.user_data.update({"date": ds, "shamsi_date": sh_d, "branches": bl, "transit": tr, "export": ex, "created_by": u.effective_user.id})
    kb = InlineKeyboardMarkup([[InlineKeyboardButton("✅ ذخیره", callback_data="conf_save")], [InlineKeyboardButton("❌ لغو", callback_data="conf_cancel")]])
    await u.message.reply_text(f"پیش‌نمایش:\n\n{create_branch_report(c.user_data)}\n\nتأیید؟", reply_markup=kb)
    return CONFIRM_SAVE


async def confirm_save_callback(u, c):
    q = u.callback_query
    await q.answer()
    if q.data == "conf_cancel":
        c.user_data.clear()
        await q.edit_message_text("❌ لغو شد.")
        await q.message.reply_text("🏠", reply_markup=admin_menu())
        return ConversationHandler.END
    if not c.user_data.get("branches") or "date" not in c.user_data:
        await q.edit_message_text("⌛ این پیش‌نمایش منقضی شده است؛ ثبت را از ابتدا شروع کنید.")
        return ConversationHandler.END
    closed = get_closed_workday(c.user_data["date"])
    if closed and closed.get("status") == "closed":
        await q.edit_message_text("🔒 این تاریخ سابقه نهایی‌سازی دارد. برای هر اصلاح، از پنل سایت و بخش کنترل روز کاری استفاده کن تا دلیل تغییر ثبت شود.")
        c.user_data.clear()
        await q.message.reply_text("🏠 منوی اصلی", reply_markup=admin_menu())
        return ConversationHandler.END
    save_report(c.user_data)
    log_action(u.effective_user.id, f"Saved report for {c.user_data['date']}")
    await q.edit_message_text("✅ ذخیره شد. ارسال فایل‌ها...")
    try:
        await q.message.reply_text(create_branch_report(c.user_data))
        await q.message.reply_text(create_trade_report(c.user_data), reply_markup=whatsapp_share_kb(c.user_data))
        jpg = create_jpg_report(c.user_data)
        with open(jpg, "rb") as f:
            await q.message.reply_photo(f, caption="🖼️ تصویر رسمی گزارش")
        os.remove(jpg)
        pdf = create_pdf_report(c.user_data)
        with open(pdf, "rb") as fp:
            await q.message.reply_document(fp, caption="📄 فایل PDF رسمی")
        os.remove(pdf)
    except Exception as e:
        logger.error(f"Send err: {e}")
    try:
        warns = analyze_report_anomalies(c.user_data)
        if warns:
            await q.message.reply_text("🧠 بررسی هوشمند گزارش:\n\n" + "\n".join(warns))
    except Exception:
        logger.exception("anomaly analysis failed")
    await send_auto_backup_to_admin(c.bot, f"بکاپ خودکار ({c.user_data['date']})")
    c.user_data.clear()
    await q.message.reply_text("✅ پایان.", reply_markup=admin_menu())
    return ConversationHandler.END


async def search_report(u, c):
    if not can_use(u):
        return ConversationHandler.END
    await u.message.reply_text("📅 تاریخ گزارش:", reply_markup=ReplyKeyboardRemove())
    return SEARCH_DATE


async def get_search_date(u, c):
    try:
        ds = norm_date(u.message.text)
    except Exception:
        await u.message.reply_text("❌ اشتباه.")
        return SEARCH_DATE
    data = load_report(ds)
    if not data:
        await u.message.reply_text("❌ یافت نشد.", reply_markup=get_menu(u))
        return ConversationHandler.END
    await u.message.reply_text(create_branch_report(data))
    await u.message.reply_text(create_trade_report(data), reply_markup=whatsapp_share_kb(data))
    try:
        jpg = create_jpg_report(data)
        with open(jpg, "rb") as f:
            await u.message.reply_photo(f)
        os.remove(jpg)
    except Exception:
        pass
    await u.message.reply_text("✅ انجام شد.", reply_markup=get_menu(u))
    return ConversationHandler.END


async def serial_search_start(u, c):
    if not can_use(u):
        return ConversationHandler.END
    await u.message.reply_text("🔢 شماره سریال:", reply_markup=ReplyKeyboardRemove())
    return SERIAL_SEARCH_STATE


async def get_serial_search(u, c):
    t = safe_int(u.message.text)
    res = []
    if not t:
        return SERIAL_SEARCH_STATE
    for r in all_reports():
        for b in r.get("branches", []):
            rngs = b.get("ranges", [{"start": b.get("start", "0"), "end": b.get("end", "0")}])
            isc = str(t) in [x.strip() for x in tr_num(b.get("cancelled_serial", "")).replace("،", ",").split(",")]
            fnd = False
            for rn in rngs:
                try:
                    st, en = int(tr_num(rn["start"])), int(tr_num(rn["end"]))
                    if st <= t <= en and st > 0:
                        fnd = True
                        break
                except Exception:
                    continue
            if fnd or isc:
                res.append(f"📅 {r.get('shamsi_date')} | {b.get('name')} | {'❌ ابطال' if isc else '✅ معتبر'}")
    if not res:
        await u.message.reply_text("❌ یافت نشد.", reply_markup=get_menu(u))
    else:
        await u.message.reply_text("🔎 نتایج:\n\n" + "\n".join(res), reply_markup=get_menu(u))
    return ConversationHandler.END


async def range_start(u, c):
    if not can_use(u):
        return ConversationHandler.END
    kb = InlineKeyboardMarkup([[InlineKeyboardButton("📆 هفته", callback_data="qr_week"), InlineKeyboardButton("📅 ماه", callback_data="qr_shamsi_month")]])
    await u.message.reply_text("📊 انتخاب کنید یا شروع را بنویسید:", reply_markup=kb)
    return RANGE_START


async def quick_range_callback(u, c):
    q = u.callback_query
    await q.answer()
    td = iran_today()
    if q.data == "qr_week":
        s, e = td - timedelta(days=(td.weekday() + 2) % 7), td
    elif q.data == "qr_shamsi_month":
        jd = jdatetime.date.fromgregorian(date=td)
        s, e = jdatetime.date(jd.year, jd.month, 1).togregorian(), td
    else:
        return ConversationHandler.END
    await q.edit_message_text("⏳ در حال استخراج...")
    await send_range_files_tg(q.message, s, e, get_menu(u))
    return ConversationHandler.END


async def get_range_start(u, c):
    try:
        c.user_data["range_start"] = parse_smart_date(u.message.text)
    except Exception:
        return RANGE_START
    await u.message.reply_text("📅 تاریخ پایان:")
    return RANGE_END


async def get_range_end(u, c):
    try:
        e = parse_smart_date(u.message.text)
    except Exception:
        return RANGE_END
    s = c.user_data.get("range_start")
    await send_range_files_tg(u.message, s, e, get_menu(u))
    c.user_data.clear()
    return ConversationHandler.END


async def send_range_files_tg(m, s, e, mu):
    await m.reply_text(range_text(s, e), reply_markup=whatsapp_range_kb(s, e))
    await m.reply_text("⬆️ آمار محاسبه شد.", reply_markup=mu)
    rws, *_ = range_summary(s, e)
    if not rws:
        return
    files = []
    try:
        jpg = create_range_jpg(s, e)
        files.append(jpg)
        with open(jpg, "rb") as f:
            await m.reply_photo(f)
        pdf = create_range_pdf(s, e)
        files.append(pdf)
        with open(pdf, "rb") as fp:
            await m.reply_document(fp)
    except Exception:
        logger.exception("range files send failed")
    finally:
        for p in files:
            safe_remove(p)


async def month_start(u, c):
    if not can_use(u):
        return ConversationHandler.END
    months = [x["key"] for x in months_with_data(6)]
    rows = [["ماه جاری", "ماه قبل"]] + [months[i:i + 3] for i in range(0, len(months), 3)] + [["❌ لغو"]]
    await u.message.reply_text("📅 ماه را انتخاب کنید یا میلادی تایپ کنید (مثال: 2026/10).\nماه‌های قدیمی هم همیشه ذخیره‌اند.",
                               reply_markup=ReplyKeyboardMarkup(rows, resize_keyboard=True))
    return MONTH


async def get_month(u, c):
    raw = tr_num(u.message.text.strip())
    today = iran_today()
    if raw in ("ماه جاری", "جاری"):
        ym = (today.year, today.month)
    elif raw in ("ماه قبل", "قبل"):
        ym = shift_month(today.year, today.month, -1)
    else:
        ym = parse_month_key(raw)
        if ym is None:  # سازگاری با ورودی قدیمی شمسی مثل 07/1405
            try:
                a, b = map(int, raw.replace("-", "/").split("/"))
                jy, jm = (b, a) if b > 1300 else (a, b)
                s, e = jalali_month_bounds(jdatetime.date(jy, jm, 1))
                await send_range_files_tg(u.message, s, e, get_menu(u))
                return ConversationHandler.END
            except Exception:
                await u.message.reply_text("⚠️ قالب معتبر نیست. مثل 2026/10 بنویسید یا یکی از دکمه‌ها را بزنید.")
                return MONTH
    await send_month_files_tg(u.message, ym[0], ym[1], get_menu(u))
    return ConversationHandler.END


async def today(u, c):
    d = iran_today()
    ds = d.strftime("%d/%m/%Y")
    dt = load_report(ds)
    if not dt:
        await u.message.reply_text("❌ ثبت نشده.", reply_markup=get_menu(u))
        return
    try:
        await u.message.reply_text(create_branch_report(dt))
        await u.message.reply_text(create_trade_report(dt), reply_markup=whatsapp_share_kb(dt))
        jpg = create_jpg_report(dt)
        with open(jpg, "rb") as f:
            await u.message.reply_photo(f)
        os.remove(jpg)
        await u.message.reply_text("✅", reply_markup=get_menu(u))
    except Exception:
        logger.exception("today report failed")
        await u.message.reply_text("⚠️ ارسال گزارش امروز با خطا مواجه شد.")


async def management(u, c):
    if is_admin(u):
        await u.message.reply_text("⚙️ پنل مدیریت پیشرفته", reply_markup=management_menu())


async def show_missing_days(u, c):
    if not is_admin(u):
        return
    m = find_missing_dates(30)
    if not m:
        await u.message.reply_text("✅ همه روزها گزارش دارند.", reply_markup=management_menu())
        return
    await u.message.reply_text("⚠️ **روزهای بدون گزارش:**\n" + "\n".join(f"• `{x['shamsi']}`" for x in m), reply_markup=management_menu(), parse_mode="Markdown")


async def system_dashboard(u, c):
    if not is_admin(u):
        return
    t = iran_today()
    rs = all_reports()
    txt = month_summary_text(t.year, t.month)
    txt += f"\n\n📁 کل گزارش‌های آرشیو: {len(rs)} | ✅ مجموع همه ماه‌ها: {sum(total_issued(r) for r in rs)}"
    ms = months_with_data(6)
    if ms:
        txt += "\n🗂 ماه‌های دارای آمار: " + " | ".join(f"{x['key']} ({x['total']})" for x in ms)
    txt += f"\n🌐 سایت روی پورت {PORT}"
    await u.message.reply_text(txt, reply_markup=management_menu())


async def show_audit_logs(u, c):
    if not is_admin(u):
        return
    await u.message.reply_text("📜 **۲۰ فعالیت اخیر:**\n" + "\n".join(f"• `{t[:16]}` | `{i}`: {a}" for i, a, t in get_recent_logs(20)), reply_markup=management_menu(), parse_mode="Markdown")


async def delete_start(u, c):
    if not is_admin(u):
        return ConversationHandler.END
    await u.message.reply_text("📅 تاریخ گزارش جهت حذف:", reply_markup=ReplyKeyboardRemove())
    return DELETE_DATE


async def get_delete_date(u, c):
    try:
        ds = norm_date(u.message.text)
    except Exception:
        await u.message.reply_text("❌ تاریخ اشتباه است.")
        return DELETE_DATE
    if not load_report(ds):
        await u.message.reply_text("❌ یافت نشد.", reply_markup=management_menu())
        return ConversationHandler.END
    kb = InlineKeyboardMarkup([[InlineKeyboardButton("🗑 بله", callback_data=f"del_yes:{ds}")], [InlineKeyboardButton("❌ لغو", callback_data="del_no")]])
    await u.message.reply_text(f"حذف گزارش {ds}؟", reply_markup=kb)
    return DELETE_DATE


async def delete_callback(u, c):
    q = u.callback_query
    await q.answer()
    if q.data != "del_no":
        ds = q.data.split(":", 1)[1]
        delete_report(ds)
        await q.edit_message_text(f"🗑 {ds} حذف شد.")
    else:
        await q.edit_message_text("❌ لغو شد.")
    await q.message.reply_text("⚙️", reply_markup=management_menu())
    return ConversationHandler.END


async def list_reports(u, c):
    if not is_admin(u):
        return
    dts = list_dates()
    await u.message.reply_text(f"📋 **گزارش‌ها ({len(dts)} مورد):**\n\n" + "\n".join(f"• {x}" for x in dts[:80]) if dts else "گزارشی نیست.", reply_markup=management_menu(), parse_mode="Markdown")


async def user_management(u, c):
    if not is_admin(u):
        return
    btns = [[InlineKeyboardButton(f"{f or un or ui} | {'👑 مدیر' if adm else ('✅ فعال' if al else '🚫 مسدود')}", callback_data=f"user:{ui}")] for ui, un, f, al, adm in users_rows()[:50]]
    await u.message.reply_text("👥 کاربران:", reply_markup=InlineKeyboardMarkup(btns))


async def user_callback(u, c):
    q = u.callback_query
    await q.answer()
    ui = int(q.data.split(":")[1])
    kb = InlineKeyboardMarkup([
        [InlineKeyboardButton("🚫 غیرفعال", callback_data=f"ua:{ui}:0"), InlineKeyboardButton("✅ فعال", callback_data=f"ua:{ui}:1")],
        [InlineKeyboardButton("👑 مدیر", callback_data=f"ua:{ui}:admin"), InlineKeyboardButton("👤 حذف مدیر", callback_data=f"ua:{ui}:unadmin")],
    ])
    await q.edit_message_text(f"کاربر `{ui}`:", reply_markup=kb, parse_mode="Markdown")


async def user_action(u, c):
    q = u.callback_query
    await q.answer()
    _, uis, act = q.data.split(":")
    ui = int(uis)
    if act in ("0", "1") and ui != ADMIN_ID:
        set_allowed(ui, act == "1")
    elif act == "admin":
        add_admin(ui)
    elif act == "unadmin" and ui != ADMIN_ID:
        remove_admin(ui)
    await q.edit_message_text("✅ اعمال شد.")


async def backup(u, c):
    if is_admin(u):
        await send_auto_backup_to_admin(c.bot, "بکاپ دستی")


async def send_report_start(u, c):
    if not is_admin(u):
        return ConversationHandler.END
    await u.message.reply_text("📅 تاریخ گزارش:", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("📅 امروز", callback_data="sdate_today"), InlineKeyboardButton("⏪ دیروز", callback_data="sdate_yesterday")]]))
    return SEND_DATE


async def send_date_quick_callback(u, c):
    q = u.callback_query
    await q.answer()
    d = iran_today() if q.data == "sdate_today" else iran_today() - timedelta(days=1)
    ds = d.strftime("%d/%m/%Y")
    if not load_report(ds):
        await q.edit_message_text("❌ ثبت نشده.")
        return ConversationHandler.END
    c.user_data["send_date"] = ds
    return await ask_send_platform(q.message, ds)


async def get_send_date(u, c):
    try:
        ds = norm_date(u.message.text)
    except Exception:
        return SEND_DATE
    if not load_report(ds):
        await u.message.reply_text("❌ پیدا نشد.")
        return SEND_DATE
    c.user_data["send_date"] = ds
    return await ask_send_platform(u.message, ds)


async def ask_send_platform(m, ds):
    kb = InlineKeyboardMarkup([
        [InlineKeyboardButton("🟢 واتساپ", callback_data="plat_wa")],
        [InlineKeyboardButton("✈️ تلگرام", callback_data="plat_tg")],
        [InlineKeyboardButton("❌ انصراف", callback_data="plat_cancel")],
    ])
    await m.reply_text(f"✅ گزارش `{ds}` آماده ارسال:", reply_markup=kb, parse_mode="Markdown")
    return SEND_PLATFORM


async def send_platform_callback(u, c):
    q = u.callback_query
    await q.answer()
    ds = c.user_data.get("send_date")
    dt = load_report(ds)
    if q.data == "plat_cancel" or not dt:
        await q.edit_message_text("❌ لغو شد.")
        return ConversationHandler.END
    if q.data == "plat_wa":
        await q.edit_message_text("🟢 شماره را بدهید یا کلیک کنید:", reply_markup=whatsapp_share_kb(dt))
        return SEND_WA_PHONE
    btns = [[InlineKeyboardButton(f"👤 {f or un or ui}", callback_data=f"sendu:{ui}")] for ui, un, f, al, _ in users_rows()[:20] if al]
    await q.edit_message_text("✈️ گیرنده:", reply_markup=InlineKeyboardMarkup(btns) if btns else None)
    return SEND_USER


async def get_send_wa_phone(u, c):
    p = normalize_phone(u.message.text)
    if not p:
        return SEND_WA_PHONE
    dt = load_report(c.user_data.get("send_date"))
    await u.message.reply_text(f"✅ لینک `+{p}`:", reply_markup=whatsapp_share_kb(dt, phone=p), parse_mode="Markdown")
    return ConversationHandler.END


async def send_user_callback(u, c):
    q = u.callback_query
    await q.answer()
    return await execute_send_to_user(q.message, c, int(q.data.split(":")[1]))


async def get_send_user(u, c):
    ui = safe_int(u.message.text)
    return await execute_send_to_user(u.message, c, ui) if ui else SEND_USER


async def execute_send_to_user(m, c, ui):
    dt = load_report(c.user_data.get("send_date"))
    if dt:
        try:
            await c.bot.send_message(chat_id=ui, text=create_branch_report(dt))
            await c.bot.send_message(chat_id=ui, text=create_trade_report(dt), reply_markup=whatsapp_share_kb(dt))
            await m.reply_text("✅ ارسال شد.", reply_markup=admin_menu())
        except Exception:
            await m.reply_text("❌ خطا در ارسال.")
    return ConversationHandler.END


async def branch_stats_start(u, c):
    if not is_admin(u):
        return ConversationHandler.END
    await u.message.reply_text("🏷 نام شعبه:", reply_markup=branch_name_kb())
    return BRANCH_SEARCH_STATE


async def branch_stats_callback(u, c):
    q = u.callback_query
    await q.answer()
    tb = q.data.split(":", 1)[1]
    await q.edit_message_text(f"🏢 شعبه: {tb}")
    return await compute_and_send_branch_stats(q.message, c, tb)


async def get_branch_stats(u, c):
    return await compute_and_send_branch_stats(u.message, c, u.message.text.strip())


async def compute_and_send_branch_stats(m, c, tb):
    rs, ti, tc, cnt = all_reports(), 0, 0, 0
    for r in rs:
        for b in r.get("branches", []):
            if tb in b.get("name", ""):
                ti += int(b.get("issued", 0))
                tc += int(b.get("cancelled", 0))
                cnt += 1
    le, ld = get_last_end_serial(tb)
    await m.reply_text(f"🔍 **آمار شعبه: {tb}**\n\n• فعالیت: {cnt} روز\n• صادره: {ti} | ابطالی: {tc}\n• آخرین پایان: `{le or '-'}` (`{ld or '-'}`)", reply_markup=management_menu(), parse_mode="Markdown")
    c.user_data.clear()
    return ConversationHandler.END


async def broadcast_start(u, c):
    if not is_admin(u):
        return ConversationHandler.END
    await u.message.reply_text("📢 متن پیام:", reply_markup=ReplyKeyboardRemove())
    return BROADCAST_STATE


async def broadcast_send(u, c):
    t = u.message.text
    conn = db()
    us = conn.execute("SELECT user_id FROM users WHERE allowed=1").fetchall()
    conn.close()
    s = 0
    for (ui,) in us:
        try:
            await c.bot.send_message(chat_id=ui, text=f"📢 اطلاعیه:\n\n{t}")
            s += 1
            await asyncio.sleep(0.05)
        except Exception:
            pass
    await u.message.reply_text(f"✅ ارسال به {s} کاربر.", reply_markup=management_menu())
    return ConversationHandler.END


async def restore_db_start(u, c):
    if not is_admin(u):
        return ConversationHandler.END
    await u.message.reply_text("📥 فایل `.db`:", reply_markup=ReplyKeyboardRemove())
    return RESTORE_DB_STATE


async def restore_db_receive(u, c):
    doc = u.message.document
    if not doc or not doc.file_name.endswith(".db"):
        await u.message.reply_text("❌ نامعتبر")
        return RESTORE_DB_STATE
    tp = os.path.join(BASE_DIR, "temp_restore.db")
    fo = await doc.get_file()
    await fo.download_to_drive(tp)
    try:
        test = sqlite3.connect(tp)
        integrity = test.execute("PRAGMA integrity_check").fetchone()
        tables = test.execute("SELECT COUNT(*) FROM sqlite_master WHERE type='table'").fetchone()[0]
        test.close()
        if not integrity or integrity[0] != "ok" or tables < 1:
            raise ValueError("نسخه پشتیبان سالم نیست")
        # Keep a safety copy of the live database before replacing it.
        if os.path.exists(DB_FILE):
            safety = os.path.join(BASE_DIR, "backups")
            os.makedirs(safety, exist_ok=True)
            shutil.copy2(DB_FILE, os.path.join(safety, f"before_restore_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"))
        _replace_db_file(tp)
        init_db()
        log_action(u.effective_user.id, "Restored database from verified Telegram backup")
        await u.message.reply_text("✅ دیتابیس بازیابی شد و بررسی سلامت فایل موفق بود. نسخه قبلی نیز در پوشه backups نگهداری شد.", reply_markup=management_menu())
    except Exception as ex:
        if os.path.exists(tp):
            try: os.remove(tp)
            except OSError: pass
        await u.message.reply_text(f"❌ بازیابی انجام نشد؛ فایل جایگزین نشد. علت: {str(ex)[:250]}", reply_markup=management_menu())
    return ConversationHandler.END


# ============================================================
# ماژول‌های افزوده‌شده: امنیت، تحلیل هوشمند، پیش‌بینی، سلامت سیستم
# ============================================================
START_TIME = time_module.time()
_REPORTS_LOCK = threading.Lock()
_REPORTS_CACHE = {"raw": None}
_FIN_PDF_LAST = {}
PERSIAN_WEEKDAYS = ["دوشنبه", "سه‌شنبه", "چهارشنبه", "پنجشنبه", "جمعه", "شنبه", "یکشنبه"]  # indexed by date.weekday()
ADMIN_ONLY_CB = ("user:", "ua:", "del_", "conf_", "acc:", "ed_", "dup_", "date_", "bcount:", "bname:",
                 "autoserial:", "endserial:", "canc_zero", "sdate_", "plat_", "sendu:", "emp:", "empx:")
_ERROR_LAST = {}


def invalidate_reports_cache():
    with _REPORTS_LOCK:
        _REPORTS_CACHE["raw"] = None


def default_month_key():
    return iran_today().strftime("%Y-%m")


def default_month_tuple():
    t = iran_today()
    return t.year, t.month


def _fin_pdf_allowed(key, min_gap=600):
    now = time_module.time()
    if now - _FIN_PDF_LAST.get(key, 0) < min_gap:
        return False
    _FIN_PDF_LAST[key] = now
    return True


def client_ip():
    if TRUST_PROXY:
        xff = request.headers.get("X-Forwarded-For", "")
        if xff:
            return xff.split(",")[0].strip()
    return request.remote_addr or "unknown"


def safe_remove(path):
    try:
        if path and os.path.exists(path):
            os.remove(path)
    except OSError:
        pass


def cleanup_generated_files(max_age=3600):
    """حذف فایل‌های موقت گزارش (jpg/pdf/xlsx/png) که قدیمی شده‌اند."""
    now = time_module.time()
    removed = 0
    legacy = ("gozaresh_", "jam_bazeh_", "range_report_", "bashmaq_report_", "chart_", "hesabdari_", "web_restore_")
    for folder in (OUT_DIR, BASE_DIR):
        try:
            names = os.listdir(folder)
        except OSError:
            continue
        for n in names:
            if folder == BASE_DIR and not n.startswith(legacy):
                continue
            if not n.lower().endswith((".jpg", ".pdf", ".xlsx", ".png", ".db")):
                continue
            p = os.path.join(folder, n)
            try:
                if os.path.isfile(p) and now - os.path.getmtime(p) > max_age:
                    os.remove(p)
                    removed += 1
            except OSError:
                pass
    return removed


def _replace_db_file(tp):
    """پیش از تعویض دیتابیس، WAL فعلی را checkpoint می‌کند تا تراکنش تأییدشده از دست نرود."""
    if os.path.exists(DB_FILE):
        checkpoint = sqlite3.connect(DB_FILE, timeout=20)
        try:
            checkpoint.execute("PRAGMA busy_timeout=20000")
            checkpoint.execute("PRAGMA wal_checkpoint(TRUNCATE)")
        finally:
            checkpoint.close()
    for suf in ("-wal", "-shm"):
        sidecar = DB_FILE + suf
        if os.path.exists(sidecar):
            os.remove(sidecar)
    os.replace(tp, DB_FILE)


# ------------------------------------------------------------
# تحلیل هوشمند: ناهنجاری، پیش‌بینی، سلامت سریال‌ها
# ------------------------------------------------------------
def _report_date(r):
    try:
        return parse_smart_date(r.get("date", ""))
    except Exception:
        return None


def daily_series(days=90, until=None, exclude=None):
    until = until or iran_today()
    start = until - timedelta(days=days)
    out = []
    for r in all_reports():
        d = _report_date(r)
        if d is None or d < start or d > until:
            continue
        if exclude and r.get("date") == exclude:
            continue
        out.append((d, total_issued(r)))
    return out


def _branch_intervals(exclude_date=None):
    items = []
    for rep in all_reports():
        if exclude_date and rep.get("date") == exclude_date:
            continue
        for b in rep.get("branches", []):
            for rg in b.get("ranges", [{"start": b.get("start", "0"), "end": b.get("end", "0")}]):
                try:
                    x, y = int(tr_num(rg.get("start", "0"))), int(tr_num(rg.get("end", "0")))
                except Exception:
                    continue
                if x > 0 and y >= x:
                    items.append((x, y, rep.get("shamsi_date", rep.get("date", "-")), b.get("name", "-")))
    return items


def analyze_report_anomalies(data):
    """بررسی هوشمند یک گزارش: جهش/افت غیرعادی، ابطالی زیاد، ناهماهنگی سریال و هم‌پوشانی."""
    warns = []
    try:
        d = parse_smart_date(data["date"])
    except Exception:
        return warns
    total, canc = total_issued(data), total_cancelled(data)
    hist = [v for _, v in daily_series(60, until=d - timedelta(days=1), exclude=data.get("date"))]
    if len(hist) >= 7:
        mean, sd = statistics.mean(hist), statistics.pstdev(hist)
        if mean > 0 and total == 0:
            warns.append(f"⚠️ مجموع صادره صفر است ولی میانگین اخیر {mean:.0f} فقره است.")
        elif sd > 0 and abs(total - mean) / sd >= 2.5:
            pct = (total - mean) / mean * 100 if mean else 0
            warns.append(f"📈 مجموع صادره ({total}) به‌طور غیرعادی {'بالاتر' if total > mean else 'پایین‌تر'} از میانگین ({mean:.0f}) است ({pct:+.0f}٪).")
        elif sd == 0 and mean > 0 and abs(total - mean) / mean > 0.5:
            warns.append(f"📈 مجموع صادره ({total}) بیش از ۵۰٪ با میانگین اخیر ({mean:.0f}) فاصله دارد.")
    if total + canc > 0 and canc >= 3 and canc / (total + canc) >= 0.10:
        warns.append(f"❌ نسبت ابطالی بالاست: {canc} ابطالی از {total + canc} ({canc / (total + canc) * 100:.0f}٪).")
    others = _branch_intervals(exclude_date=data.get("date"))
    for b in data.get("branches", []):
        name = b.get("name", "-")
        span, valid = 0, True
        for rg in b.get("ranges", [{"start": b.get("start", "0"), "end": b.get("end", "0")}]):
            try:
                x, y = int(tr_num(rg.get("start", "0"))), int(tr_num(rg.get("end", "0")))
            except Exception:
                valid = False
                continue
            if x <= 0 and y <= 0:
                continue
            if y < x:
                warns.append(f"⚠️ شعبه {name}: بازه سریال نامعتبر ({x} تا {y}).")
                valid = False
                continue
            span += y - x + 1
            for ox, oy, odate, oname in others:
                if x <= oy and y >= ox:
                    warns.append(f"🔁 شعبه {name}: هم‌پوشانی سریال {x}-{y} با {ox}-{oy} ({oname}، {odate}).")
                    break
        issued, cancelled = int(b.get("issued", 0)), int(b.get("cancelled", 0))
        if valid and span and span not in (issued, issued + cancelled):
            warns.append(f"🧮 شعبه {name}: طول بازه سریال ({span}) با تعداد صادره ({issued}) هم‌خوان نیست.")
    return warns[:8]


def forecast_next_days(n=7):
    today = iran_today()
    series = daily_series(90, until=today)
    if len(series) < 5 or n <= 0:
        return []
    vals = [v for _, v in series]
    overall = statistics.mean(vals)
    ew = vals[0]
    for v in vals[1:]:
        ew = 0.3 * v + 0.7 * ew
    by_wd = {}
    for d, v in series:
        by_wd.setdefault(d.weekday(), []).append(v)
    out = []
    for i in range(1, n + 1):
        d = today + timedelta(days=i)
        wd = by_wd.get(d.weekday(), [])
        wm = statistics.mean(wd) if len(wd) >= 3 else overall
        out.append((d, round(0.6 * wm + 0.4 * ew)))
    return out


def jalali_month_bounds(jd):
    start = jdatetime.date(jd.year, jd.month, 1)
    nxt = jdatetime.date(jd.year + 1, 1, 1) if jd.month == 12 else jdatetime.date(jd.year, jd.month + 1, 1)
    return start.togregorian(), nxt.togregorian() - timedelta(days=1)


def month_projection():
    today = iran_today()
    s_, e_ = month_bounds(today.year, today.month)
    so_far = sum(v for d, v in daily_series(40, until=today) if s_ <= d <= today)
    left = max((e_ - today).days, 0)
    fc = forecast_next_days(left)
    proj = so_far + sum(v for _, v in fc)
    prev_total = month_totals(*shift_month(today.year, today.month, -1))["issued"]
    return {"month": f"{today.year}/{today.month:02d}", "so_far": so_far, "projected": proj, "days_left": left, "prev_total": prev_total}


def serial_audit_issues(limit=300):
    items = sorted(_branch_intervals())
    issues, prev = [], None
    for cur in items:
        if prev and cur[0] <= prev[1]:
            issues.append(f"تداخل سریال {cur[0]} الی {min(cur[1], prev[1])}: شعبه {prev[3]} ({prev[2]}) با {cur[3]} ({cur[2]})")
        elif prev and cur[0] > prev[1] + 1:
            issues.append(f"پرش احتمالی از {prev[1]} تا {cur[0]} بین {prev[3]} ({prev[2]}) و {cur[3]} ({cur[2]})")
        if not prev or cur[1] > prev[1]:
            prev = cur
    return issues[:limit]


def recent_anomalies(count=7):
    out = []
    for r in all_reports()[-count:]:
        w = analyze_report_anomalies(r)
        if w:
            out.append((r.get("shamsi_date", r.get("date", "-")), w))
    return out


def best_weekday_text():
    by = {}
    for d, v in daily_series(90):
        by.setdefault(d.weekday(), []).append(v)
    by = {k: statistics.mean(v) for k, v in by.items() if len(v) >= 3}
    if len(by) < 3:
        return None
    hi, lo = max(by, key=by.get), min(by, key=by.get)
    return f"پرکارترین روز هفته {PERSIAN_WEEKDAYS[hi]} (میانگین {by[hi]:.0f}) و کم‌کارترین {PERSIAN_WEEKDAYS[lo]} (میانگین {by[lo]:.0f}) است."


def branch_share_text(days=30):
    until = iran_today()
    start = until - timedelta(days=days)
    totals = {}
    for r in all_reports():
        d = _report_date(r)
        if d is None or not (start <= d <= until):
            continue
        for b in r.get("branches", []):
            totals[b.get("name", "-")] = totals.get(b.get("name", "-"), 0) + int(b.get("issued", 0))
    s = sum(totals.values())
    if not s:
        return None
    return " | ".join(f"{n}: {v} ({v / s * 100:.0f}٪)" for n, v in sorted(totals.items(), key=lambda x: -x[1]))


def extra_smart_insights():
    ins = []
    try:
        fc = forecast_next_days(7)
        if fc:
            ins.append({"title": "پیش‌بینی ۷ روز آینده", "text": "؛ ".join(f"{shamsi(d)} ({PERSIAN_WEEKDAYS[d.weekday()]}): {v}" for d, v in fc)})
        mp = month_projection()
        if mp["so_far"]:
            t = f"تا امروز {mp['so_far']:,} فقره؛ برآورد پایان ماه {mp['projected']:,} فقره ({mp['days_left']} روز مانده)."
            if mp["prev_total"]:
                t += f" ماه قبل: {mp['prev_total']:,} ({(mp['projected'] - mp['prev_total']) / mp['prev_total'] * 100:+.0f}٪)."
            ins.append({"title": f"برآورد ماه {mp['month']}", "text": t})
        bw = best_weekday_text()
        if bw:
            ins.append({"title": "الگوی هفتگی", "text": bw})
        bs = branch_share_text()
        if bs:
            ins.append({"title": "سهم شعب (۳۰ روز اخیر)", "text": bs})
        an = recent_anomalies(7)
        ins.append({"title": "ناهنجاری‌های اخیر", "text": "؛ ".join(f"{d}: {w[0]}" for d, w in an) if an else "در ۷ گزارش اخیر مورد مشکوکی دیده نشد."})
    except Exception:
        logger.exception("extra insights failed")
    return ins


def build_smart_analysis_text():
    lines = ["🧠 تحلیل هوشمند مرز باشماق", ""]
    for x in extra_smart_insights():
        lines.append(f"▪️ {x['title']}:\n{x['text']}\n")
    iss = serial_audit_issues(5)
    lines.append("🔢 سلامت سریال‌ها:\n" + ("\n".join(iss) if iss else "تداخل یا پرش مشکوکی یافت نشد."))
    return "\n".join(lines)[:4000]


def system_health_info():
    info = {"uptime_h": round((time_module.time() - START_TIME) / 3600, 1), "db_ok": False}
    try:
        conn = db()
        try:
            info["db_ok"] = conn.execute("PRAGMA quick_check").fetchone()[0] == "ok"
            info["reports"] = conn.execute("SELECT COUNT(*) FROM reports").fetchone()[0]
            info["pending"] = conn.execute("SELECT COUNT(*) FROM pending_branch_reports WHERE status='pending'").fetchone()[0]
            info["users"] = conn.execute("SELECT COUNT(*) FROM users").fetchone()[0]
        finally:
            conn.close()
        info["db_mb"] = round(os.path.getsize(DB_FILE) / 1048576, 2)
    except Exception:
        logger.exception("health check failed")
    bdir = os.path.join(BASE_DIR, "backups")
    try:
        files = [os.path.join(bdir, n) for n in os.listdir(bdir) if n.endswith(".db")]
        info["last_backup_h"] = round((time_module.time() - max(os.path.getmtime(f) for f in files)) / 3600, 1) if files else None
    except OSError:
        info["last_backup_h"] = None
    try:
        info["disk_free_gb"] = round(shutil.disk_usage(BASE_DIR).free / 1073741824, 1)
    except OSError:
        info["disk_free_gb"] = None
    info["bot"] = bool(GLOBAL_BOT_APP)
    info["font"] = bool(FONT_PATH)
    info["missing_30d"] = len(find_missing_dates(30))
    return info


def build_health_text():
    h = system_health_info()
    yn = lambda v: "✅" if v else "❌"
    lb = h.get("last_backup_h")
    return (
        "🩺 سلامت سیستم\n\n"
        f"{yn(h['db_ok'])} دیتابیس | حجم: {h.get('db_mb', '?')} MB\n"
        f"📁 گزارش‌ها: {h.get('reports', '?')} | 👥 کاربران: {h.get('users', '?')}\n"
        f"⏳ گزارش‌های منتظر تأیید: {h.get('pending', '?')}\n"
        f"{yn(h['bot'])} ربات | {yn(h['font'])} فونت فارسی\n"
        f"💾 آخرین بکاپ: {'نامشخص' if lb is None else f'{lb} ساعت پیش'}\n"
        f"💽 فضای آزاد: {h.get('disk_free_gb', '?')} GB\n"
        f"📅 روزهای بدون گزارش (۳۰ روز): {h['missing_30d']}\n"
        f"⏱ مدت اجرا: {h['uptime_h']} ساعت"
    )


# ------------------------------------------------------------
# امنیت وب: بررسی منشأ درخواست (CSRF)، هدرهای امنیتی، سلامت‌سنجی
# ------------------------------------------------------------
_STATE_CHANGING_GET = set()
_ALLOWED_HOSTS = {h.strip().lower() for h in os.getenv("ALLOWED_HOSTS", "").split(",") if h.strip()}
_CSRF_CHECK = os.getenv("CSRF_ORIGIN_CHECK", "1").strip() != "0"


def _same_site_request():
    origin = request.headers.get("Origin") or request.headers.get("Referer") or ""
    if origin and origin != "null":
        host = (urllib.parse.urlparse(origin).hostname or "").lower()
        allowed = {(urllib.parse.urlparse("//" + request.host).hostname or "").lower()} | _ALLOWED_HOSTS
        if TRUST_PROXY and request.headers.get("X-Forwarded-Host"):
            allowed.add((urllib.parse.urlparse("//" + request.headers["X-Forwarded-Host"].split(",")[0].strip()).hostname or "").lower())
        return host in allowed
    sfs = request.headers.get("Sec-Fetch-Site", "")
    return sfs in ("", "same-origin", "none")


@flask_app.before_request
def _web_security_guard():
    if not _CSRF_CHECK:
        return None
    if request.method in ("POST", "PUT", "PATCH", "DELETE") or request.path in _STATE_CHANGING_GET:
        if not _same_site_request():
            logger.warning("Blocked cross-site request to %s", request.path)
            return jsonify({"ok": False, "error": "cross-site request blocked"}), 403
        if request.method in ("POST", "PUT", "PATCH", "DELETE"):
            expected = session.get("_csrf_token", "")
            supplied = request.headers.get("X-CSRFToken", "") or request.form.get("csrf_token", "")
            if request.is_json and not supplied:
                payload = request.get_json(silent=True) or {}
                supplied = str(payload.get("csrf_token", ""))
            if not expected or not supplied or not hmac.compare_digest(str(expected), str(supplied)):
                logger.warning("Blocked request with missing/invalid CSRF token: %s", request.path)
                return jsonify({"ok": False, "error": "درخواست معتبر نیست؛ صفحه را تازه‌سازی کنید."}), 403
    return None


@flask_app.errorhandler(413)
def _upload_too_large(_error):
    if request.path.startswith("/api/"):
        return jsonify({"ok": False, "error": "حجم درخواست بیش از حد مجاز است."}), 413
    flash("حجم فایل بیش از حد مجاز است. فایل کوچک‌تری انتخاب کنید.", "danger")
    return redirect(url_for("web_dashboard" if session.get("logged_in") else "web_login"))


@flask_app.after_request
def _web_security_headers(resp):
    resp.headers.setdefault("X-Content-Type-Options", "nosniff")
    resp.headers.setdefault("X-Frame-Options", "SAMEORIGIN")
    resp.headers.setdefault("Referrer-Policy", "same-origin")
    if request.path.startswith("/api/"):
        resp.headers["Cache-Control"] = "no-store"
    return resp


@flask_app.route("/healthz")
def web_healthz():
    ok = False
    try:
        conn = db()
        try:
            conn.execute("SELECT 1").fetchone()
            ok = True
        finally:
            conn.close()
    except Exception:
        pass
    return jsonify({"ok": ok, "bot": bool(GLOBAL_BOT_APP)}), (200 if ok else 503)


@flask_app.route("/api/forecast")
def api_forecast():
    if not session.get("logged_in"):
        return jsonify({"ok": False}), 401
    fc = forecast_next_days(7)
    return jsonify({"ok": True,
                    "forecast": [{"date": d.strftime("%d/%m/%Y"), "shamsi": shamsi(d), "weekday": PERSIAN_WEEKDAYS[d.weekday()], "predicted": v} for d, v in fc],
                    "month": month_projection()})


@flask_app.route("/api/anomalies")
def api_anomalies():
    if not session.get("logged_in"):
        return jsonify({"ok": False}), 401
    return jsonify({"ok": True,
                    "recent": [{"date": d, "warnings": w} for d, w in recent_anomalies(14)],
                    "serial_issues": serial_audit_issues(50)})


# ------------------------------------------------------------
# ربات: دروازه دسترسی، درخواست دسترسی، تحلیل و گزارش‌های خودکار
# ------------------------------------------------------------
def _access_state(uid):
    conn = db()
    try:
        u = conn.execute("SELECT allowed FROM users WHERE user_id=?", (uid,)).fetchone()
        r = conn.execute("SELECT status FROM access_requests WHERE user_id=?", (uid,)).fetchone()
    finally:
        conn.close()
    return u, (r[0] if r else None)


def _set_access_request(uid, status):
    conn = db()
    conn.execute("INSERT INTO access_requests(user_id,status,requested_at) VALUES(?,?,?) "
                 "ON CONFLICT(user_id) DO UPDATE SET status=excluded.status",
                 (uid, status, datetime.now(TEHRAN).isoformat(timespec="seconds")))
    conn.commit()
    conn.close()


async def access_gate(u, c):
    """همه‌ی به‌روزرسانی‌ها از اینجا عبور می‌کنند؛ کاربر غیرمجاز و callbackهای مدیریتی برای غیرمدیر مسدود می‌شوند."""
    user = u.effective_user
    if user is None:
        return
    q = u.callback_query
    if is_allowed(user.id):
        if q and q.data and q.data.startswith(ADMIN_ONLY_CB) and not is_admin_id(user.id):
            await q.answer("⛔ این بخش فقط برای مدیر است.", show_alert=True)
            raise ApplicationHandlerStop
        return
    if u.message and (u.message.text or "").split("@")[0].strip() == "/id":
        return
    row, req = _access_state(user.id)
    save_user(user)
    if q:
        await q.answer("⛔ دسترسی ندارید.", show_alert=True)
    elif u.effective_message:
        if row is None:
            _set_access_request(user.id, "pending")
            await u.effective_message.reply_text("🔒 درخواست دسترسی شما برای مدیر ارسال شد. پس از تأیید، پیام می‌گیرید.\nشناسه شما: " + str(user.id))
            name = (user.first_name or "-") + (f" (@{user.username})" if user.username else "")
            kb = InlineKeyboardMarkup([[InlineKeyboardButton("✅ تأیید", callback_data=f"acc:ok:{user.id}"),
                                        InlineKeyboardButton("🚫 رد", callback_data=f"acc:no:{user.id}")]])
            try:
                await c.bot.send_message(chat_id=ADMIN_ID, text=f"🔔 درخواست دسترسی جدید\nنام: {name}\nشناسه: {user.id}", reply_markup=kb)
            except Exception:
                logger.exception("access request notify failed")
        elif req == "pending":
            await u.effective_message.reply_text("⏳ درخواست شما در انتظار تأیید مدیر است.")
        else:
            await u.effective_message.reply_text("🚫 دسترسی شما غیرفعال است. با مدیر تماس بگیرید.")
    raise ApplicationHandlerStop


async def access_callback(u, c):
    q = u.callback_query
    await q.answer()
    if not is_admin_id(u.effective_user.id):
        return
    try:
        _, act, uid_s = q.data.split(":")
        uid = int(uid_s)
    except Exception:
        return
    ok = act == "ok"
    set_allowed(uid, ok)
    _set_access_request(uid, "approved" if ok else "denied")
    log_action(u.effective_user.id, f"Access {'approved' if ok else 'denied'} for {uid}")
    await q.edit_message_text(f"{'✅ تأیید شد' if ok else '🚫 رد شد'}: {uid}")
    try:
        if ok:
            await c.bot.send_message(chat_id=uid, text="✅ دسترسی شما تأیید شد. /start را بزنید.", reply_markup=user_menu())
        else:
            await c.bot.send_message(chat_id=uid, text="🚫 درخواست دسترسی شما رد شد.")
    except Exception:
        pass


async def smart_analysis(u, c):
    if not is_admin(u):
        return
    await u.message.reply_text(build_smart_analysis_text(), reply_markup=management_menu())


async def health_cmd(u, c):
    if not is_admin(u):
        return
    await u.message.reply_text(build_health_text(), reply_markup=management_menu())


async def help_cmd(u, c):
    if is_admin(u):
        txt = ("📖 راهنمای مدیر\n/start منو | /id شناسه | /help راهنما\n"
               "/analysis تحلیل و پیش‌بینی | /health سلامت سیستم | /webpass نام کاربری پنل | /webreset رمز جدید پنل\n"
               "/amar ثبت آمار شعبه | /branchstatus وضعیت دو شعبه\n"
               "دکمه «⚡ ثبت فوق‌سریع»: تعداد تیرپارک، پارسیان، ترانزیت، صادرات را با فاصله بفرستید.\n"
               "هر ثبت گزارش به‌صورت خودکار از نظر ناهنجاری و هم‌پوشانی سریال بررسی می‌شود.")
    else:
        txt = "📖 از منوی پایین استفاده کنید. شناسه شما با /id نمایش داده می‌شود."
    await u.message.reply_text(txt)


async def weekly_summary_job(ctx: ContextTypes.DEFAULT_TYPE):
    if iran_today().weekday() != 5:  # شنبه
        return
    end = iran_today() - timedelta(days=1)
    start = end - timedelta(days=6)
    rows, iss, tr, ex, canc, bt = range_summary(start, end)
    pr = range_summary(start - timedelta(days=7), end - timedelta(days=7))
    chg = f" ({(iss - pr[1]) / pr[1] * 100:+.0f}٪ نسبت به هفته قبل)" if pr[1] else ""
    txt = (f"🗓 خلاصه هفتگی {shamsi(start)} تا {shamsi(end)}\n\nصادره: {iss}{chg}\nابطالی: {canc}\n"
           f"ترانزیت: {tr} | صادرات: {ex}\nروزهای ثبت‌شده: {len(rows)} از ۷\n")
    if bt:
        txt += "\n" + "\n".join(f"▫️ {n}: {v}" for n, v in bt.items())
    fc = forecast_next_days(7)
    if fc:
        txt += f"\n\n🔮 پیش‌بینی هفته پیش‌رو: {sum(v for _, v in fc)} فقره"
    try:
        await ctx.bot.send_message(chat_id=ADMIN_ID, text=txt)
    except Exception:
        logger.exception("weekly summary failed")


async def monthly_summary_job(ctx: ContextTypes.DEFAULT_TYPE):
    """اول هر ماه میلادی: خلاصه و اکسل ماه گذشته به مدیر ارسال می‌شود و آمار ماه جدید از صفر شروع می‌شود."""
    today = iran_today()
    if today.day != 1:
        return
    y, m = shift_month(today.year, today.month, -1)
    s_, e_ = month_bounds(y, m)
    try:
        await ctx.bot.send_message(chat_id=ADMIN_ID, text="🗓 پایان ماه میلادی\n\n" + month_summary_text(y, m) +
                                   f"\n\n🆕 آمار ماه {month_key(today.year, today.month)} از صفر شروع شد؛ ماه‌های قبل در آرشیو (پنل و «📈 آمار ماهانه») محفوظ‌اند.")
        if month_totals(y, m)["days"]:
            path = create_excel_report(s_, e_)
            with open(path, "rb") as f:
                await ctx.bot.send_document(chat_id=ADMIN_ID, document=f, filename=os.path.basename(path), caption=f"📊 اکسل ماه {month_key(y, m)}")
    except Exception:
        logger.exception("monthly summary failed")


async def morning_gap_job(ctx: ContextTypes.DEFAULT_TYPE):
    y = iran_today() - timedelta(days=1)
    if not load_report(y.strftime("%d/%m/%Y")):
        try:
            await ctx.bot.send_message(chat_id=ADMIN_ID, text=f"🌅 گزارش دیروز ({shamsi(y)}) ثبت نشده است. با «⚡ ثبت فوق‌سریع» ثبت کنید.")
        except Exception:
            pass


async def cleanup_job(ctx: ContextTypes.DEFAULT_TYPE):
    cleanup_generated_files()


def _employee_detail(uid):
    e = get_employee(uid, active_only=False)
    if not e:
        return None
    txt = (f"👷 {e['employee_name']}\nشناسه: {e['user_id']}\nشعبه: {e['branch_name']}\n"
           f"وضعیت: {'✅ فعال' if e['active'] else '⏸ غیرفعال'}\n"
           f"دسترسی: {'📝 ثبت آمار' if e['can_submit'] else '👁 فقط مشاهده'}" + (f"\nیادداشت: {e['note']}" if e.get('note') else ""))
    kb = InlineKeyboardMarkup([
        [InlineKeyboardButton("⏸ غیرفعال‌سازی" if e["active"] else "✅ فعال‌سازی", callback_data=f"empx:t:{uid}"),
         InlineKeyboardButton("👁 فقط مشاهده" if e["can_submit"] else "📝 اجازه ثبت", callback_data=f"empx:s:{uid}")],
        [InlineKeyboardButton("🗑 حذف", callback_data=f"empx:d:{uid}"), InlineKeyboardButton("🔙 فهرست", callback_data="emp:list")],
    ])
    return txt, kb


def _employees_list_view():
    emps = get_branch_employees()
    if not emps:
        return "هنوز کارمندی ثبت نشده است. ثبت از پنل وب انجام می‌شود.", None
    rows = [[InlineKeyboardButton(f"{'✅' if e['active'] else '⏸'} {e['employee_name']} | {e['branch_name']}{'' if e['can_submit'] else ' 👁'}", callback_data=f"emp:{e['user_id']}")] for e in emps[:50]]
    return f"👷 کارمندان ({len(emps)}):", InlineKeyboardMarkup(rows)


async def employees_admin_menu(u, c):
    if not is_admin(u):
        return
    txt, kb = _employees_list_view()
    await u.message.reply_text(txt, reply_markup=kb)


async def employees_callback(u, c):
    q = u.callback_query
    await q.answer()
    if not is_admin_id(u.effective_user.id):
        return
    arg = q.data.split(":", 1)[1]
    if arg == "list":
        txt, kb = _employees_list_view()
    else:
        d = _employee_detail(safe_int(arg) or 0)
        txt, kb = d if d else ("کارمند پیدا نشد.", None)
    await q.edit_message_text(txt, reply_markup=kb)


async def employees_action_callback(u, c):
    q = u.callback_query
    await q.answer()
    if not is_admin_id(u.effective_user.id):
        return
    try:
        _, act, uid_s = q.data.split(":")
        uid = int(uid_s)
    except Exception:
        return
    old = get_employee(uid, active_only=False)
    if not old:
        await q.edit_message_text("کارمند پیدا نشد.")
        return
    note = ""
    if act == "t":
        new = 0 if old["active"] else 1
        update_employee(uid, active=new)
        note = employee_change_message(old, active=new)
    elif act == "s":
        new = 0 if old["can_submit"] else 1
        update_employee(uid, can_submit=new)
        note = employee_change_message(old, can_submit=new)
    elif act == "d":
        kb = InlineKeyboardMarkup([[InlineKeyboardButton("🗑 بله، حذف شود", callback_data=f"empx:dy:{uid}"), InlineKeyboardButton("❌ انصراف", callback_data=f"emp:{uid}")]])
        await q.edit_message_text(f"حذف «{old['employee_name']}»؟", reply_markup=kb)
        return
    elif act == "dy":
        ok, pending = delete_employee(uid)
        log_action(u.effective_user.id, f"Deleted branch employee {uid} via bot")
        try:
            await c.bot.send_message(chat_id=uid, text="⛔ دسترسی کارمندی شما توسط مدیر حذف شد.")
        except Exception:
            pass
        txt, kb = _employees_list_view()
        await q.edit_message_text(f"🗑 حذف شد." + (f" ({pending} گزارش منتظر او در صف است)" if pending else "") + "\n\n" + txt, reply_markup=kb)
        return
    if act in ("t", "s"):
        log_action(u.effective_user.id, f"Employee {uid} changed via bot ({act})")
        if note:
            try:
                await c.bot.send_message(chat_id=uid, text=note)
            except Exception:
                pass
    d = _employee_detail(uid)
    await q.edit_message_text(d[0], reply_markup=d[1])


# ============================================================
# آمار ماهانه (میلادی): هر ماه جدا محاسبه و همیشه قابل مشاهده است
# ============================================================
def month_bounds(year, month):
    start = date(year, month, 1)
    end = date(year + (month == 12), (month % 12) + 1, 1) - timedelta(days=1)
    return start, end


def month_key(year, month):
    return f"{year:04d}-{month:02d}"


def shift_month(year, month, delta):
    idx = year * 12 + (month - 1) + delta
    return idx // 12, idx % 12 + 1


def parse_month_key(text, default=None):
    """قالب‌های 2026-10، 2026/10 و 10/2026 (میلادی) را می‌پذیرد؛ نامعتبر → default."""
    try:
        parts = tr_num(str(text)).strip().replace("-", "/").split("/")
        a, b = int(parts[0]), int(parts[1])
        y, m = (a, b) if a > 31 else (b, a)
        if len(parts) == 2 and 1900 <= y <= 2200 and 1 <= m <= 12:
            return y, m
    except Exception:
        pass
    return default


def pct_change(cur, prev):
    return None if not prev else round((cur - prev) / prev * 100)


def month_totals(year, month):
    start, end = month_bounds(year, month)
    rows, iss, tr, ex, canc, bt = range_summary(start, end)
    return {"year": year, "month": month, "key": month_key(year, month), "days": len(rows), "issued": iss,
            "transit": tr, "export": ex, "cancelled": canc, "avg": round(iss / len(rows), 1) if rows else 0,
            "branches": bt, "start": start, "end": end}


def month_compare(year, month):
    cur = month_totals(year, month)
    prev = month_totals(*shift_month(year, month, -1))
    return {"cur": cur, "prev": prev, "issued_pct": pct_change(cur["issued"], prev["issued"]),
            "avg_pct": pct_change(cur["avg"], prev["avg"])}


def months_with_data(limit=24):
    sums = {}
    for r in all_reports():
        d = _report_date(r)
        if d is not None:
            k = month_key(d.year, d.month)
            sums[k] = sums.get(k, 0) + total_issued(r)
    return [{"key": k, "total": v} for k, v in sorted(sums.items(), reverse=True)[:limit]]


def month_summary_text(year, month):
    cmp_ = month_compare(year, month)
    cur, prev = cmp_["cur"], cmp_["prev"]
    today = iran_today()
    state = "ماه جاری (تا امروز)" if (year, month) == (today.year, today.month) else "آرشیو ماه‌های گذشته"
    lines = [f"📈 آمار ماه میلادی {cur['key']}", f"☀️ شمسی: {shamsi(cur['start'])} تا {shamsi(cur['end'])}", f"📌 {state}", ""]
    if not cur["days"]:
        lines.append("❌ برای این ماه گزارشی ثبت نشده است.")
    else:
        pc = cmp_["issued_pct"]
        lines += [f"• روزهای ثبت‌شده: {cur['days']} | میانگین روزانه: {cur['avg']}",
                  f"✅ کل صادره: {cur['issued']}" + ("" if pc is None else f" ({pc:+d}٪ نسبت به ماه قبل)"),
                  f"❌ ابطالی: {cur['cancelled']}", f"🚛 ترانزیت: {cur['transit']} | 📦 صادرات: {cur['export']}"]
        if cur["branches"]:
            lines += ["", "🏢 تفکیک شعبه‌ها:"] + [f"▫️ {n}: {v} فقره" for n, v in cur["branches"].items()]
    if prev["days"]:
        lines.append(f"\n↩️ ماه قبل ({prev['key']}): {prev['issued']} فقره در {prev['days']} روز")
    return "\n".join(lines)


@flask_app.route("/api/months")
def api_months():
    if not session.get("logged_in"):
        return jsonify({"ok": False}), 401
    ym = parse_month_key(request.args.get("month", ""), default_month_tuple())
    cmp_ = month_compare(*ym)
    strip = lambda t: {k: v for k, v in t.items() if k not in ("start", "end")}
    return jsonify({"ok": True, "months": months_with_data(36), "selected": month_key(*ym),
                    "current": strip(cmp_["cur"]), "previous": strip(cmp_["prev"]),
                    "issued_change_pct": cmp_["issued_pct"], "avg_change_pct": cmp_["avg_pct"]})


async def send_month_files_tg(msg, year, month, menu):
    start, end = month_bounds(year, month)
    await msg.reply_text(month_summary_text(year, month), reply_markup=whatsapp_range_kb(start, end))
    rows, *_ = range_summary(start, end)
    if not rows:
        await msg.reply_text("هر ماه دیگری را هم می‌توانید از «📈 آمار ماهانه» ببینید.", reply_markup=menu)
        return
    files = []
    try:
        jpg = create_range_jpg(start, end)
        files.append(jpg)
        with open(jpg, "rb") as f:
            await msg.reply_photo(f)
        pdf = create_range_pdf(start, end)
        files.append(pdf)
        with open(pdf, "rb") as fp:
            await msg.reply_document(fp)
    except Exception:
        logger.exception("month files send failed")
    finally:
        for p in files:
            safe_remove(p)
    await msg.reply_text("✅ آمار ماه ارسال شد.", reply_markup=menu)



# ============================================================
# ورود به پنل با نام کاربری و رمز قابل‌تغییر (رمز به‌صورت هش‌شده ذخیره می‌شود)
# ============================================================
CRED_FILE = os.path.join(BASE_DIR, ".web_credentials.json")
LEGACY_PW_FILE = os.path.join(BASE_DIR, ".web_password")
_CRED_LOCK = threading.Lock()
_CREDS = {"data": None}
_DUMMY_HASH = {"v": None}
_USERNAME_RE = re.compile(r"^[A-Za-z0-9_.\-\u0600-\u06FF]{3,32}$")
ERR_BAD_CURRENT = "رمز فعلی اشتباه است."
_WEAK_PASSWORDS = {"123456", "12345678", "123456789", "password", "admin123", "11111111", "qwertyui", "12341234"}


def _write_creds(data):
    tmp = CRED_FILE + ".tmp"
    fd = os.open(tmp, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(fd, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False)
    os.replace(tmp, CRED_FILE)


def get_web_credentials():
    """اطلاعات ورود فعلی؛ در اولین اجرا از WEB_USERNAME/WEB_PASSWORD (یا رمز تصادفی) ساخته می‌شود."""
    with _CRED_LOCK:
        if _CREDS["data"] is None:
            data = None
            if os.getenv("RESET_WEB_CREDENTIALS", "0").strip() != "1":
                try:
                    with open(CRED_FILE, "r", encoding="utf-8") as fh:
                        data = json.load(fh)
                    if not all(k in data for k in ("username", "password_hash", "ver")):
                        data = None
                except Exception:
                    data = None
            if data is None:
                uname = (os.getenv("WEB_USERNAME", "") or "admin").strip()
                if not _USERNAME_RE.match(uname):
                    uname = "admin"
                data = {"username": uname, "password_hash": generate_password_hash(WEB_PASS),
                        "ver": int(time_module.time() * 1000), "updated_at": datetime.now(TEHRAN).isoformat(timespec="seconds")}
                try:
                    _write_creds(data)
                except Exception:
                    logger.exception("could not persist web credentials")
            _CREDS["data"] = data
        return dict(_CREDS["data"])


def verify_web_login(username, password):
    creds = get_web_credentials()
    user_ok = hmac.compare_digest(str(username).strip().lower().encode("utf-8"), creds["username"].lower().encode("utf-8"))
    if _DUMMY_HASH["v"] is None:
        _DUMMY_HASH["v"] = generate_password_hash(secrets.token_hex(8))
    # حتی با نام کاربری اشتباه هش بررسی می‌شود تا زمان پاسخ اطلاعاتی لو ندهد
    pw_ok = check_password_hash(creds["password_hash"] if user_ok else _DUMMY_HASH["v"], str(password))
    return bool(user_ok and pw_ok)


def validate_new_credentials(username=None, password=None, effective_username=""):
    if username is not None and not _USERNAME_RE.match(username):
        return "نام کاربری باید ۳ تا ۳۲ کاراکتر و فقط حروف، عدد و . _ - باشد."
    if password is not None:
        if len(password) < 8 or len(password) > 128:
            return "رمز عبور باید بین ۸ تا ۱۲۸ کاراکتر باشد."
        if password.lower() in _WEAK_PASSWORDS or password.lower() == effective_username.lower():
            return "این رمز بیش از حد ساده است؛ رمز قوی‌تری انتخاب کنید."
    return None


def _store_creds(username, password_hash):
    data = {"username": username, "password_hash": password_hash, "ver": int(time_module.time() * 1000),
            "updated_at": datetime.now(TEHRAN).isoformat(timespec="seconds")}
    with _CRED_LOCK:
        _write_creds(data)
        _CREDS["data"] = data
    return data


def change_web_credentials(current_password, new_username="", new_password=""):
    creds = get_web_credentials()
    if not check_password_hash(creds["password_hash"], str(current_password)):
        return False, ERR_BAD_CURRENT
    new_username, new_password = (new_username or "").strip(), new_password or ""
    if not new_username and not new_password:
        return False, "نام کاربری یا رمز جدید را وارد کنید."
    uname = new_username or creds["username"]
    err = validate_new_credentials(new_username or None, new_password or None, uname)
    if err:
        return False, err
    data = _store_creds(uname, generate_password_hash(new_password) if new_password else creds["password_hash"])
    if new_password:
        safe_remove(LEGACY_PW_FILE)  # فایل متنی رمز اولیه دیگر معتبر نیست
    return True, data


def reset_web_password():
    """رمز تصادفی جدید می‌سازد (برای فراموشی رمز، از ربات)."""
    creds = get_web_credentials()
    new_pw = secrets.token_urlsafe(10)
    _store_creds(creds["username"], generate_password_hash(new_pw))
    safe_remove(LEGACY_PW_FILE)
    return creds["username"], new_pw


@flask_app.before_request
def _session_version_guard():
    """با تغییر نام کاربری/رمز، نشست‌های قدیمی (سایر دستگاه‌ها) خودکار بسته می‌شوند."""
    if session.get("logged_in"):
        try:
            if session.get("cred_ver") != get_web_credentials()["ver"]:
                session.clear()
        except Exception:
            logger.exception("session version check failed")
    return None


@flask_app.route("/account/update", methods=["POST"])
def web_account_update():
    if not session.get("logged_in"):
        return redirect(url_for("web_login"))
    key = "acct:" + client_ip()
    now = time_module.time()
    attempts = [t for t in _LOGIN_ATTEMPTS.get(key, []) if now - t < 600]
    if len(attempts) >= 5:
        flash("تلاش‌های ناموفق زیاد است؛ ۱۰ دقیقه بعد دوباره امتحان کنید.", "danger")
        return redirect(url_for("web_dashboard"))
    new_pw = request.form.get("new_password", "")
    if new_pw and new_pw != request.form.get("confirm_password", ""):
        flash("❌ رمز جدید و تکرار آن یکسان نیست.", "danger")
        return redirect(url_for("web_dashboard"))
    ok, res = change_web_credentials(request.form.get("current_password", ""), request.form.get("new_username", ""), new_pw)
    if not ok:
        if res == ERR_BAD_CURRENT:  # فقط حدس رمز فعلی شمرده می‌شود، نه خطای اعتبارسنجی
            attempts.append(now)
            _LOGIN_ATTEMPTS[key] = attempts
        flash(f"❌ {res}", "danger")
        return redirect(url_for("web_dashboard"))
    _LOGIN_ATTEMPTS.pop(key, None)
    session["cred_ver"] = res["ver"]
    session["web_user"] = res["username"]
    changed = " و ".join(x for x, c in (("نام کاربری", request.form.get("new_username", "").strip()), ("رمز عبور", new_pw)) if c)
    log_action(ADMIN_ID, f"Web credentials changed: {changed}")
    notify_user_async(ADMIN_ID, f"🔐 {changed} پنل وب از داخل پنل تغییر کرد (IP: {client_ip()}). اگر کار شما نبوده، با /webreset رمز را بازنشانی کنید.")
    flash(f"✅ {changed} با موفقیت تغییر کرد. سایر دستگاه‌ها از پنل خارج شدند.", "success")
    return redirect(url_for("web_dashboard"))


async def webpass_cmd(u, c):
    """نمایش نام کاربری پنل؛ رمز فقط وقتی نشان داده می‌شود که هنوز همان رمز اولیه باشد."""
    if not u.effective_user or u.effective_user.id != ADMIN_ID:
        return
    creds = get_web_credentials()
    txt = f"👤 نام کاربری پنل: {creds['username']}\n"
    initial = None
    try:
        if os.path.exists(LEGACY_PW_FILE):
            with open(LEGACY_PW_FILE, "r", encoding="utf-8") as fh:
                cand = fh.read().strip()
            if cand and check_password_hash(creds["password_hash"], cand):
                initial = cand
    except Exception:
        initial = None
    if initial:
        txt += f"🔑 رمز اولیه: {initial}\n(بعد از تغییر رمز در پنل، دیگر نمایش داده نمی‌شود)"
    else:
        txt += "🔑 رمز ذخیره‌شده قابل نمایش نیست. برای دریافت رمز جدید /webreset را بزنید."
    msg = await u.message.reply_text(txt + f"\n🌐 آدرس: پورت {PORT} همین سرور\n⚠️ این پیام ۶۰ ثانیه بعد پاک می‌شود.")
    await _delete_later(c, msg)


async def webreset_cmd(u, c):
    if not u.effective_user or u.effective_user.id != ADMIN_ID:
        return
    username, pw = reset_web_password()
    log_action(ADMIN_ID, "Web password reset via bot")
    msg = await u.message.reply_text(f"🔐 رمز جدید پنل ساخته شد (رمز قبلی و نشست‌های قبلی باطل شدند)\n👤 نام کاربری: {username}\n🔑 رمز: {pw}\n\n"
                                     "بعد از ورود، از «حساب کاربری» رمز دلخواه خود را بگذارید.\n⚠️ این پیام ۶۰ ثانیه بعد پاک می‌شود.")
    await _delete_later(c, msg)


async def _delete_later(c, msg, delay=60):
    async def _del(ctx):
        try:
            await ctx.bot.delete_message(chat_id=msg.chat_id, message_id=msg.message_id)
        except Exception:
            pass
    if c.job_queue:
        c.job_queue.run_once(_del, delay)


def build_telegram_application():
    tg_app = Application.builder().token(TOKEN).connect_timeout(30.0).read_timeout(30.0).write_timeout(30.0).build()
    tg_app.add_error_handler(error_handler)
    tg_app.add_handler(TypeHandler(Update, access_gate), group=-1)

    if tg_app.job_queue:
        tg_app.job_queue.run_daily(nightly_reminder_job, time=time(hour=20, minute=30, tzinfo=TEHRAN))
        tg_app.job_queue.run_daily(nightly_backup_job, time=time(hour=23, minute=0, tzinfo=TEHRAN))
        tg_app.job_queue.run_daily(morning_gap_job, time=time(hour=9, minute=15, tzinfo=TEHRAN))
        tg_app.job_queue.run_daily(weekly_summary_job, time=time(hour=8, minute=0, tzinfo=TEHRAN))
        tg_app.job_queue.run_daily(monthly_summary_job, time=time(hour=8, minute=30, tzinfo=TEHRAN))
        tg_app.job_queue.run_repeating(cleanup_job, interval=1800, first=60)

    fallbacks = [
        CommandHandler("cancel", cancel),
        MessageHandler(filters.Regex(r"^❌ لغو$"), cancel),
        MessageHandler(filters.Regex(r"^🔙 بازگشت$"), back),
    ]

    report_conv = ConversationHandler(
        entry_points=[
            MessageHandler(filters.Regex(r"^📝 ثبت گزارش جدید$"), new_report),
            MessageHandler(filters.Regex(r"^⚡ ثبت فوق‌سریع$"), fast_report_start),
            MessageHandler(filters.Regex(r"^✏️ ویرایش گزارش$"), edit_start),
        ],
        states={
            FAST_REPORT_STATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, fast_report_receive)],
            DATE: [
                CallbackQueryHandler(date_quick_callback, pattern=r"^date_"),
                CallbackQueryHandler(duplicate_callback, pattern=r"^dup_"),
                MessageHandler(filters.TEXT & ~filters.COMMAND, get_date),
            ],
            EDIT_DATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, get_edit_date)],
            EDIT_CHOICE: [CallbackQueryHandler(edit_choice_callback, pattern=r"^ed_")],
            COUNT: [
                CallbackQueryHandler(count_callback, pattern=r"^bcount:"),
                MessageHandler(filters.TEXT & ~filters.COMMAND, get_count),
            ],
            NAME: [
                CallbackQueryHandler(name_callback, pattern=r"^bname:"),
                MessageHandler(filters.TEXT & ~filters.COMMAND, get_name),
            ],
            ISSUED: [MessageHandler(filters.TEXT & ~filters.COMMAND, get_issued)],
            START: [
                CallbackQueryHandler(auto_serial_callback, pattern=r"^autoserial:"),
                MessageHandler(filters.TEXT & ~filters.COMMAND, get_start),
            ],
            END: [
                CallbackQueryHandler(end_serial_callback, pattern=r"^endserial:"),
                MessageHandler(filters.TEXT & ~filters.COMMAND, get_end),
            ],
            CANCELLED: [
                CallbackQueryHandler(cancelled_zero_callback, pattern=r"^canc_zero$"),
                MessageHandler(filters.TEXT & ~filters.COMMAND, get_cancelled),
            ],
            CANCELLED_SERIAL: [MessageHandler(filters.TEXT & ~filters.COMMAND, get_cancelled_serial)],
            TRANSIT: [MessageHandler(filters.TEXT & ~filters.COMMAND, get_transit)],
            EXPORT: [MessageHandler(filters.TEXT & ~filters.COMMAND, get_export)],
            CONFIRM_SAVE: [CallbackQueryHandler(confirm_save_callback, pattern=r"^conf_")],
        },
        fallbacks=fallbacks,
        allow_reentry=True,
        per_message=False,
        conversation_timeout=900,
    )

    tg_app.add_handler(report_conv, group=0)
    tg_app.add_handler(ConversationHandler(entry_points=[MessageHandler(filters.Regex(r"^🔎 جستجوی گزارش$"), search_report)], states={SEARCH_DATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, get_search_date)]}, fallbacks=fallbacks, per_message=False, conversation_timeout=900), group=0)
    tg_app.add_handler(ConversationHandler(entry_points=[MessageHandler(filters.Regex(r"^🔢 جستجوی سریال$"), serial_search_start)], states={SERIAL_SEARCH_STATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, get_serial_search)]}, fallbacks=fallbacks, per_message=False, conversation_timeout=900), group=0)
    tg_app.add_handler(ConversationHandler(entry_points=[MessageHandler(filters.Regex(r"^📊 گزارش بازه‌ای$"), range_start)], states={RANGE_START: [CallbackQueryHandler(quick_range_callback, pattern=r"^qr_"), MessageHandler(filters.TEXT & ~filters.COMMAND, get_range_start)], RANGE_END: [MessageHandler(filters.TEXT & ~filters.COMMAND, get_range_end)]}, fallbacks=fallbacks, per_message=False, conversation_timeout=900), group=0)
    tg_app.add_handler(ConversationHandler(entry_points=[MessageHandler(filters.Regex(r"^📈 آمار ماهانه$"), month_start)], states={MONTH: [MessageHandler(filters.TEXT & ~filters.COMMAND, get_month)]}, fallbacks=fallbacks, per_message=False, conversation_timeout=900), group=0)
    tg_app.add_handler(ConversationHandler(entry_points=[MessageHandler(filters.Regex(r"^🗑 حذف گزارش$"), delete_start)], states={DELETE_DATE: [CallbackQueryHandler(delete_callback, pattern=r"^del_"), MessageHandler(filters.TEXT & ~filters.COMMAND, get_delete_date)]}, fallbacks=fallbacks, per_message=False, conversation_timeout=900), group=0)
    tg_app.add_handler(ConversationHandler(entry_points=[MessageHandler(filters.Regex(r"^📤 ارسال گزارش$"), send_report_start)], states={SEND_DATE: [CallbackQueryHandler(send_date_quick_callback, pattern=r"^sdate_"), MessageHandler(filters.TEXT & ~filters.COMMAND, get_send_date)], SEND_PLATFORM: [CallbackQueryHandler(send_platform_callback, pattern=r"^plat_")], SEND_USER: [CallbackQueryHandler(send_user_callback, pattern=r"^sendu:"), MessageHandler(filters.TEXT & ~filters.COMMAND, get_send_user)], SEND_WA_PHONE: [MessageHandler(filters.TEXT & ~filters.COMMAND, get_send_wa_phone)]}, fallbacks=fallbacks, per_message=False, conversation_timeout=900), group=0)
    tg_app.add_handler(ConversationHandler(entry_points=[MessageHandler(filters.Regex(r"^🔍 آمار شعبه$"), branch_stats_start)], states={BRANCH_SEARCH_STATE: [CallbackQueryHandler(branch_stats_callback, pattern=r"^bname:"), MessageHandler(filters.TEXT & ~filters.COMMAND, get_branch_stats)]}, fallbacks=fallbacks, per_message=False, conversation_timeout=900), group=0)
    tg_app.add_handler(ConversationHandler(entry_points=[MessageHandler(filters.Regex(r"^📢 پیام همگانی$"), broadcast_start)], states={BROADCAST_STATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, broadcast_send)]}, fallbacks=fallbacks, per_message=False, conversation_timeout=900), group=0)
    tg_app.add_handler(ConversationHandler(entry_points=[MessageHandler(filters.Regex(r"^📥 بازیابی دیتابیس$"), restore_db_start)], states={RESTORE_DB_STATE: [MessageHandler(filters.Document.ALL, restore_db_receive)]}, fallbacks=fallbacks, per_message=False, conversation_timeout=900), group=0)

    tg_app.add_handler(CommandHandler("start", start), group=1)
    tg_app.add_handler(CommandHandler("help", help_cmd), group=1)
    tg_app.add_handler(CommandHandler("analysis", smart_analysis), group=1)
    tg_app.add_handler(CommandHandler("health", health_cmd), group=1)
    tg_app.add_handler(CommandHandler("webpass", webpass_cmd), group=1)
    tg_app.add_handler(CommandHandler("webreset", webreset_cmd), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^🧠 تحلیل هوشمند$"), smart_analysis), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^🩺 سلامت سیستم$"), health_cmd), group=1)
    tg_app.add_handler(CallbackQueryHandler(access_callback, pattern=r"^acc:"), group=1)
    tg_app.add_handler(CommandHandler("employees", employees_admin_menu), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^👷 مدیریت کارمندان$"), employees_admin_menu), group=1)
    tg_app.add_handler(CallbackQueryHandler(employees_callback, pattern=r"^emp:"), group=1)
    tg_app.add_handler(CallbackQueryHandler(employees_action_callback, pattern=r"^empx:"), group=1)
    tg_app.add_handler(CommandHandler("id", my_id), group=1)
    tg_app.add_handler(CommandHandler("amar", employee_amar_start), group=1)
    tg_app.add_handler(CommandHandler("amarstatus", employee_amar_status), group=1)
    tg_app.add_handler(CommandHandler("employeehelp", employee_help), group=1)
    tg_app.add_handler(CommandHandler("branchstatus", admin_branch_status), group=1)
    tg_app.add_handler(CommandHandler("cancel", employee_amar_cancel), group=1)
    # Button-first shortcuts for employee commands; commands remain available for power users.
    tg_app.add_handler(MessageHandler(filters.Regex(r"^📝 ثبت آمار شعبه$"), employee_amar_start), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^📋 وضعیت گزارش‌های من$"), employee_amar_status), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^❓ راهنمای کارمند$"), employee_help), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^❌ لغو ثبت آمار$"), employee_amar_cancel), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^📊 وضعیت دو شعبه$"), admin_branch_status), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^📊 وضعیت شعبه من$"), admin_branch_status), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^🔐 کنترل روز کاری$"), admin_workday_control), group=1)
    tg_app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, employee_amar_receive), group=5)
    tg_app.add_handler(CallbackQueryHandler(employee_review_callback, pattern=r"^emp(?:approve|reject):"), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^💰 صورت‌حساب مالی$"), send_monthly_financial_files), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^📅 گزارش امروز$"), today), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^⚙️ مدیریت$"), management), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^📅 روزهای جامانده$"), show_missing_days), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^📊 داشبورد سیستم$"), system_dashboard), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^🔙 بازگشت$"), back), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^📋 لیست گزارش‌ها$"), list_reports), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^👥 مدیریت کاربران$"), user_management), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^💾 پشتیبان‌گیری$"), backup), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^📜 لاگ مدیران$"), show_audit_logs), group=1)
    tg_app.add_handler(CallbackQueryHandler(user_callback, pattern=r"^user:"), group=1)
    tg_app.add_handler(CallbackQueryHandler(user_action, pattern=r"^ua:"), group=1)
    tg_app.add_handler(MessageHandler(filters.Regex(r"^❌ لغو$"), cancel), group=1)
    return tg_app


def run_bot_in_background():
    global GLOBAL_BOT_APP, GLOBAL_LOOP
    if not TOKEN or TOKEN == "YOUR_BOT_TOKEN":
        return
    try:
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        GLOBAL_LOOP = loop
        tg_app = build_telegram_application()
        GLOBAL_BOT_APP = tg_app
        tg_app.run_polling(stop_signals=None, close_loop=False)
    except Exception as e:
        logger.error(f"Bot error: {e}")


init_db()
_startup_creds = get_web_credentials()
logger.info("پنل وب: نام کاربری «%s». تغییر از داخل پنل (حساب کاربری)؛ فراموشی رمز: /webreset در ربات.", _startup_creds["username"])
cleanup_generated_files()
if not os.getenv("FLASK_SECRET_KEY"):
    logger.info("A persistent Flask session key will be stored in .flask_secret_key beside this script.")

if __name__ == "__main__":
    web_thread = threading.Thread(target=run_flask_server, daemon=True)
    web_thread.start()
    run_bot_in_background()
else:
    bot_thread = threading.Thread(target=run_bot_in_background, daemon=True)
    bot_thread.start()